#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""parse_bank.py — עובר ושב exports -> work/normalized/<account-id>.csv (P3, F1/F2).

Purpose : for every `accounts[]` entry read its files (glob), dispatch on the `bank` id to the
          format spec in formats.py, assert the exact header, normalise every transaction to
          common.NORMALIZED_COLUMNS (amount_ils > 0 = money OUT), validate the running balance
          (F1 newest-first chain for 'bank_xlsx_a'; F2 reconstructed balance with sparse checkpoints
          for 'bank_xls_b') and list every distinct description with count and signed total (the
          table that drives merchant-rule writing).
Inputs  : tazrim.config.json (`accounts[]`: id, bank, label, owner, files, opening_balance?);
          the export files (format A .xlsx via openpyxl; format B .xls via xlrd or .xlsx via openpyxl).
          Source files are never modified.
Outputs : work/normalized/<account-id>.csv per account; exactly one JSON object on stdout:
          {"ok": true, "accounts": {<id>: {rows, files, date_from, date_to, inflows, outflows,
           net, balance: {...}, descriptions: [...], anomalies: [...]}}, "warnings": [...]}.
Exit    : 0 ok; 1 parse failure (header mismatch, unreadable file, no files); 2 config error.
"""
import datetime as _dt
import hashlib
import os
import sys
from collections import OrderedDict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (NORMALIZED_COLUMNS, emit, ensure_dirs, fail, log, parse_args,  # noqa: E402
                    project_path, resolve_files, stable_id, write_csv)
from formats import BANKS, reader_for  # noqa: E402

TOL = 0.005


# ----------------------------------------------------------------------------- cell helpers
def s(v):
    """Stringify a cell; None / blank-space -> ''; integral floats without '.0'."""
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return str(v).strip()


def is_blank(v):
    return s(v) == ""


def num(v):
    """Float or None (accepts '1,234.50' strings)."""
    if v is None or (isinstance(v, str) and not v.strip()):
        return None
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(str(v).replace(",", "").replace("₪", "").strip())
    except ValueError:
        return None


def to_date(v, datemode=0):
    """datetime / date / Excel serial / 'dd/mm/yyyy' / 'yyyy-mm-dd' -> datetime.date or None."""
    if v is None or v == "":
        return None
    if isinstance(v, _dt.datetime):
        return v.date()
    if isinstance(v, _dt.date):
        return v
    if isinstance(v, (int, float)):
        try:
            import xlrd
            return xlrd.xldate_as_datetime(float(v), datemode).date()
        except Exception:  # noqa: BLE001
            return None
    t = str(v).strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d.%m.%Y", "%d/%m/%y", "%d.%m.%y"):
        try:
            return _dt.datetime.strptime(t[:10], fmt).date()
        except ValueError:
            continue
    return None


def md5(path):
    h = hashlib.md5()
    with open(path, "rb") as fh:
        for chunk in iter(lambda: fh.read(1 << 16), b""):
            h.update(chunk)
    return h.hexdigest()


def load_grid(path, sheet_name, reader):
    """Return (rows, datemode): rows = list of lists of raw cell values (1-based sheet rows ->
    index 0), read with openpyxl or xlrd. Sheet resolved by name, else the first sheet."""
    if reader == "xlrd":
        import xlrd
        wb = xlrd.open_workbook(path)
        sh = wb.sheet_by_name(sheet_name) if sheet_name in wb.sheet_names() else wb.sheet_by_index(0)
        rows = [[sh.cell_value(r, c) for c in range(sh.ncols)] for r in range(sh.nrows)]
        return rows, wb.datemode
    import openpyxl
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb[sheet_name] if sheet_name in wb.sheetnames else wb.worksheets[0]
    rows = [list(r) for r in ws.iter_rows(values_only=True)]
    wb.close()
    return rows, 0


def header_map(row, expected, where):
    """Assert `row` (cell values) equals `expected` (positionally, blanks == ''); return
    {header name: column index} for the non-blank names."""
    got = [s(v) for v in row[:len(expected)]] + [""] * max(0, len(expected) - len(row))
    if got != list(expected):
        raise ValueError("%s: header mismatch.\n  expected: %s\n  got:      %s" % (where, expected, got))
    return {h: i for i, h in enumerate(got) if h}


# ----------------------------------------------------------------------------- format A (F1)
def parse_bank_xlsx_a(path, account, spec):
    """Newest-first statement with a balance on every row -> (rows, balance_report, anomalies)."""
    rows_raw, _ = load_grid(path, spec["sheet"], reader_for("bank_xlsx_a", path))
    hi = spec["header_row"] - 1
    if len(rows_raw) <= hi:
        raise ValueError("%s: sheet has no header row %d" % (path, spec["header_row"]))
    col = header_map(rows_raw[hi], spec["header"], os.path.basename(path))
    c = spec["columns"]
    fname = os.path.basename(path)
    out, chain, anomalies = [], [], []
    for idx in range(spec["data_from_row"] - 1, len(rows_raw)):
        vals = rows_raw[idx]
        row_ref = idx + 1
        if all(is_blank(v) for v in vals):
            continue

        def g(name):
            i = col.get(c[name])
            return vals[i] if i is not None and i < len(vals) else None

        d = to_date(g("date"))
        amt = num(g("amount"))
        if d is None or amt is None:
            anomalies.append("row %d: unparseable date/amount %r" % (row_ref, [s(v) for v in vals[:5]]))
            continue
        bal = num(g("balance"))
        details = []
        if not is_blank(g("reference")):
            details.append("אסמכתה: %s" % s(g("reference")))
        if not is_blank(g("channel")):
            details.append("ערוץ: %s" % s(g("channel")))
        if not is_blank(g("fee")):
            details.append("עמלה: %s" % s(g("fee")))
        vd = to_date(g("value_date"))
        if vd and vd != d:
            details.append("יום ערך: %s" % vd.isoformat())
        if bal is not None:
            details.append("יתרה: %s" % s(bal))
        signed = round(-amt, 2)                      # file +credit/-debit -> ours +out/-in
        out.append({
            "id": stable_id(account["id"], fname, row_ref), "source": account["label"], "card": "",
            "original_name": s(g("description")), "txn_date": d.isoformat(), "charge_date": d.isoformat(),
            "amount_ils": signed, "orig_currency": "ILS", "orig_amount": abs(signed),
            "details": " | ".join(details), "source_file": fname, "row_ref": row_ref,
        })
        if bal is not None:
            chain.append((row_ref, amt, bal))
    breaks = []
    for (r1, a1, b1), (r2, a2, b2) in zip(chain, chain[1:]):
        if abs(round(b1 - a1, 2) - round(b2, 2)) > TOL:
            breaks.append("rows %d->%d: %.2f - (%.2f) = %.2f != %.2f" % (r1, r2, b1, a1, b1 - a1, b2))
    implied_opening = round(chain[-1][2] - chain[-1][1], 2) if chain else None
    report = {
        "method": "F1", "order": "newest_first", "checks": max(0, len(chain) - 1),
        "breaks": len(breaks), "break_details": breaks, "ok": not breaks,
        "closing_balance": chain[0][2] if chain else None,
        "opening_balance": implied_opening, "opening_source": "implied (last balance - last amount)",
    }
    if account.get("opening_balance") is not None and implied_opening is not None:
        cfg_open = float(account["opening_balance"])
        report["opening_balance_config"] = cfg_open
        if abs(cfg_open - implied_opening) > TOL:
            anomalies.append("configured opening_balance %.2f != implied %.2f" % (cfg_open, implied_opening))
    return out, report, anomalies


# ----------------------------------------------------------------------------- format B (F2)
def parse_bank_xls_b(path, account, spec):
    """Oldest-first statement, separate credit/debit columns, balance only on the last row of a
    date group -> (rows, balance_report, anomalies)."""
    reader = reader_for("bank_xls_b", path)
    rows_raw, datemode = load_grid(path, spec["sheet"], reader)
    hi = spec["header_row"] - 1
    if len(rows_raw) <= hi:
        raise ValueError("%s: sheet has no header row %d" % (path, spec["header_row"]))
    col = header_map(rows_raw[hi], spec["header"], os.path.basename(path))
    c = spec["columns"]
    fname = os.path.basename(path)
    anomalies = []

    # opening balance: the 'יתרת פתיחה' row (its value in the balance column), else config, else inferred
    opening, opening_src = None, None
    start_idx = spec["data_from_row"] - 1
    orow = rows_raw[spec["opening_row"] - 1] if len(rows_raw) >= spec["opening_row"] else []
    if any(s(v) == spec["opening_marker"] for v in orow):
        opening = num(orow[col[c["balance"]]]) if col[c["balance"]] < len(orow) else None
        opening_src = "file row %d (%s)" % (spec["opening_row"], spec["opening_marker"])
    else:
        start_idx = spec["opening_row"] - 1         # no opening row: data starts one row earlier
    if account.get("opening_balance") is not None:
        cfg_open = float(account["opening_balance"])
        if opening is not None and abs(cfg_open - opening) > TOL:
            anomalies.append("configured opening_balance %.2f != file opening %.2f (file value used)" % (cfg_open, opening))
        if opening is None:
            opening, opening_src = cfg_open, "config accounts[].opening_balance"

    parsed = []   # (row_ref, date, credit, debit, balance_or_None, desc, ref, optype, vdate)
    for idx in range(start_idx, len(rows_raw)):
        vals = rows_raw[idx]
        row_ref = idx + 1
        if all(is_blank(v) for v in vals):
            continue

        def g(name):
            i = col.get(c[name])
            return vals[i] if i is not None and i < len(vals) else None

        d = to_date(g("date"), datemode)
        if d is None:
            anomalies.append("row %d: bad date %r" % (row_ref, s(g("date"))))
            continue
        credit = num(g("credit")) or 0.0
        debit = num(g("debit")) or 0.0
        if credit and debit:
            anomalies.append("row %d: both credit and debit set" % row_ref)
        if not credit and not debit:
            anomalies.append("row %d: zero amount" % row_ref)
        parsed.append((row_ref, d, credit, debit, num(g("balance")), s(g("description")),
                       s(g("reference")), s(g("op_type")), to_date(g("value_date"), datemode)))

    inferred = False
    if opening is None:
        # infer from the first checkpoint: opening = checkpoint balance - net movement up to it
        running = 0.0
        for (_r, _d, cr, db, bal, *_rest) in parsed:
            running += cr - db
            if bal is not None:
                opening, opening_src, inferred = round(bal - running, 2), "inferred from first balance checkpoint", True
                break
        if opening is None:
            opening, opening_src, inferred = 0.0, "unknown (no opening row, no config, no checkpoint) - assumed 0", True
            anomalies.append("opening balance unknown; assumed 0.00")

    out, mismatches, checkpoints = [], [], 0
    running = float(opening)
    for (row_ref, d, credit, debit, bal, desc, ref, optype, vd) in parsed:
        signed = round(debit - credit, 2)
        running = round(running + credit - debit, 2)
        details = []
        if ref:
            details.append("אסמכתא: %s" % ref)
        if optype:
            details.append("סוג פעולה: %s" % optype)
        if vd and vd != d:
            details.append("תאריך ערך: %s" % vd.isoformat())
        if bal is not None:
            details.append("יתרה: %s" % s(bal))
            checkpoints += 1
            if abs(bal - running) > TOL:
                mismatches.append("row %d: sheet balance %.2f vs reconstructed %.2f" % (row_ref, bal, running))
        out.append({
            "id": stable_id(account["id"], fname, row_ref), "source": account["label"], "card": "",
            "original_name": desc, "txn_date": d.isoformat(), "charge_date": d.isoformat(),
            "amount_ils": signed, "orig_currency": "ILS", "orig_amount": abs(signed),
            "details": " | ".join(details), "source_file": fname, "row_ref": row_ref,
        })
    report = {
        "method": "F2", "order": "oldest_first", "checks": checkpoints, "breaks": len(mismatches),
        "break_details": mismatches, "ok": not mismatches, "opening_balance": opening,
        "opening_source": opening_src, "opening_inferred": inferred, "closing_balance": running,
    }
    return out, report, anomalies


PARSERS = {"bank_xlsx_a": parse_bank_xlsx_a, "bank_xls_b": parse_bank_xls_b}


# ----------------------------------------------------------------------------- driver
def describe_rows(rows):
    """Distinct descriptions with count and signed total (+ = out), largest |total| first."""
    agg = OrderedDict()
    for r in rows:
        k = r["original_name"]
        cnt, tot = agg.get(k, (0, 0.0))
        agg[k] = (cnt + 1, round(tot + r["amount_ils"], 2))
    return [{"name": k, "count": v[0], "total": v[1]}
            for k, v in sorted(agg.items(), key=lambda kv: -abs(kv[1][1]))]


def parse_account(cfg, account):
    """Parse every file of one account -> (rows, summary). Raises ValueError on format errors."""
    bank = account["bank"]
    if bank not in PARSERS:
        raise ValueError("no parser for bank id %r (known: %s)" % (bank, ", ".join(sorted(PARSERS))))
    spec = BANKS[bank]
    files = resolve_files(cfg, account.get("files"))
    if not files:
        raise ValueError("account %s: no files match %r" % (account["id"], account.get("files")))
    rows, reports, anomalies, seen, skipped = [], {}, [], {}, []
    for path in files:
        digest = md5(path)
        if digest in seen:
            skipped.append("%s is byte-identical to %s - skipped" % (os.path.basename(path), seen[digest]))
            continue
        seen[digest] = os.path.basename(path)
        r, rep, an = PARSERS[bank](path, account, spec)
        rows += r
        reports[os.path.basename(path)] = rep
        anomalies += ["%s: %s" % (os.path.basename(path), a) for a in an]
        log("[%s] %s: %d rows, balance %s (%d checks, %d breaks)" % (
            account["id"], os.path.basename(path), len(r), rep["method"], rep["checks"], rep["breaks"]))
    inflow = round(sum(-r["amount_ils"] for r in rows if r["amount_ils"] < 0), 2)
    outflow = round(sum(r["amount_ils"] for r in rows if r["amount_ils"] > 0), 2)
    dates = sorted(r["txn_date"] for r in rows)
    summary = {
        "label": account["label"], "bank": bank, "files": [os.path.basename(f) for f in files],
        "rows": len(rows), "date_from": dates[0] if dates else None, "date_to": dates[-1] if dates else None,
        "inflows": inflow, "outflows": outflow, "net": round(inflow - outflow, 2),
        "balance": reports[list(reports)[0]] if len(reports) == 1 else reports,
        "balance_ok": all(rep["ok"] for rep in reports.values()),
        "descriptions": describe_rows(rows), "anomalies": anomalies, "skipped_files": skipped,
    }
    return rows, summary


def main(argv=None):
    args, cfg = parse_args("parse bank (עובר ושב) exports into work/normalized/<account-id>.csv", argv=argv)
    if not cfg["accounts"]:
        fail("no accounts[] configured", "add at least one account to tazrim.config.json")
    ensure_dirs(cfg)
    out_dir = project_path(cfg, "work.normalized")
    result, warnings = {}, list(cfg.warnings)
    for account in cfg["accounts"]:
        try:
            rows, summary = parse_account(cfg, account)
        except Exception as e:  # noqa: BLE001
            fail("account %s: %s" % (account["id"], e),
                 "check the file layout against references/bank-formats.md (formats.py '%s')" % account["bank"])
            return
        path = os.path.join(out_dir, "%s.csv" % account["id"])
        rows.sort(key=lambda r: (r["txn_date"], r["source_file"], int(r["row_ref"])))
        write_csv(path, rows, NORMALIZED_COLUMNS)
        summary["csv"] = path
        result[account["id"]] = summary
        for a in summary["anomalies"]:
            log("  ANOMALY %s: %s" % (account["id"], a))
        for w in summary["skipped_files"]:
            warnings.append("%s: %s" % (account["id"], w))
        if not summary["balance_ok"]:
            warnings.append("%s: balance chain has breaks - see accounts[%s].balance" % (account["id"], account["id"]))
    emit({"ok": True, "normalized_dir": out_dir, "accounts": result, "warnings": warnings})


if __name__ == "__main__":
    main()
