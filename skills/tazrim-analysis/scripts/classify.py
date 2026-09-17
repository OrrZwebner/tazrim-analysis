#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""classify.py — merge every normalized CSV, classify each row, write work/database.csv (P10).

Purpose : ordered classification (order matters; flags are recomputed last):
           1. concatenate work/normalized/*.csv (+ p2p CSVs from `p2p[]` unless level=overview);
              apply generated rules from config (card debits in the bank via
              cards[].bank_debit_pattern / benefit_programs[].bank_debit_pattern ->
              type "תשלום כרטיס אשראי"; cards[].fx_settlement.transfer_pattern -> internal FX
              transfer; p2p[].bank_withdrawal_pattern -> internal P2P withdrawal), then
              rules/merchant_rules.csv (case-insensitive substring; LONGEST pattern wins, ties in
              file order; columns pattern,name_clean,type,group,cat[,group2,cat2],note);
              unmatched bank inflows -> income "אחר / לא מזוהה", everything else -> expense
              "אחר / לא מזוהה" with rule_note "לא נמצא כלל סיווג";
           2. unmatched foreign-currency rows -> "הוצאות בחו"ל" (country from details);
           3. benefit-card loads -> internal transfer;
           4. P2P pairing (F6): card rows containing `card_marker` vs app outgoing rows
              (|Δamount| <= match_tolerance, |Δdays| <= match_days, nearest date, each app row
              once) -> card row "כפילות" with linked_id; unmatched card rows -> "<owner> – העברה
              לא מזוהה (BIT/PayBox)"; app incoming -> income; "Withdrawal" rows -> internal;
           5. canonicalise category names to the primary scheme (template if
              categories.from_template, else references/category-scheme-default.csv) and to the
              optional secondary scheme; report categories that are in no scheme;
           6. computed club-discount rows (source_file "מחושב") -> expense, linked to the debit;
           7. rules/user_labels.csv (id,user_text,user_cat) then rules/user_labels_interpreted.csv
              (id,type,group_tz,cat_tz,group_new,cat_new,trip,name_clean,note) overrides;
           8. estimates[] synthetic rows (source "ידני – הערכה"; optional funded_by_cat retypes
              the funding cash withdrawals as internal) and trips[] attribution by country word /
              currency / date range — both skipped at level overview;
           9. helper columns month, month_name, pay (always "<label> (<owner>)"), person,
              in_window, summed = in_window & type in {הוצאה, הכנסה}, amount (positive for
              both expenses and incomes; refunds stay negative; type carries the direction).
          Level gating: overview lumps unknowns silently; standard/deep add a "נדרש:" marker to
          rule_note so the row lands in the "לסיווג ידני" sheet.
Inputs  : tazrim.config.json; work/normalized/*.csv; the rules CSVs; the template workbook
          when categories.from_template is set.
Outputs : work/database.csv (common.DB_COLUMNS); one JSON object on stdout with counts by
          type, in-window totals, unknown count, P2P pairs, categories not in the scheme.
Exit    : 0 ok; 1 failure (no normalized CSVs / rules unreadable); 2 config error.
"""
import datetime as _dt
import glob
import os
import re
import sys
from collections import OrderedDict, defaultdict

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (CAT_ABROAD, CAT_BENEFIT_LOAD, CAT_CARD_DEBIT, CAT_CASH_FUNDING,  # noqa: E402
                    CAT_CLUB_DISCOUNT, CAT_FX_TRANSFER, CAT_P2P_INCOME, CAT_P2P_WITHDRAWAL,
                    DB_COLUMNS, GENERATED_CATEGORIES, GROUP_ABROAD, GROUP_CARD_DEBIT, GROUP_INTERNAL,
                    NOTE_NO_RULE, SHARED, SUMMED_TYPES, TYPE_CARD_DEBIT, TYPE_DUPLICATE, TYPE_EXPENSE, TYPE_INCOME, TYPE_INTERNAL, TYPE_REIMBURSEMENT,
                    UNKNOWN_CAT, UNKNOWN_GROUP_EXPENSE, UNKNOWN_GROUP_INCOME, USER_CAT_COL,
                    USER_TEXT_COL, emit, ensure_dirs, fail, flag, in_window, log, month_label,
                    months, owner_label, parse_args, project_path, read_csv, resolve_files,
                    scheme_path, skill_root, stable_id, write_csv)
from formats import ISSUERS  # noqa: E402
from fx_rates import COUNTRY_CCY, country_in  # noqa: E402

LOAD_FLAG = ISSUERS["club_docx"]["card_segment"]["load_flag"]
COMPUTED_FILE = "מחושב"
# category / group names of everything classify assigns without a rule live in common.GENERATED_CATEGORIES
DISCOUNT_NAME = CAT_CLUB_DISCOUNT
ESTIMATE_SOURCE = "ידני – הערכה"
ESTIMATE_FILE = "הערכה"
ABROAD_GROUP, ABROAD_CAT = GROUP_ABROAD, CAT_ABROAD
ABROAD_CATS = {ABROAD_CAT, "טיסות", "לינה", "חופשות בארץ"}
GROUP_CARD = GROUP_CARD_DEBIT
CAT_CARD, CAT_LOAD, CAT_FX, CAT_P2P_OUT = CAT_CARD_DEBIT, CAT_BENEFIT_LOAD, CAT_FX_TRANSFER, CAT_P2P_WITHDRAWAL
CAT_P2P_IN = CAT_P2P_INCOME
TRIP_UNASSIGNED = "לא משויך – לסיווג ידני"
NEEDED = "נדרש:"
CLASS_COLS = ["name_clean", "type", "group_tz", "cat_tz", "group_new", "cat_new", "rule_note"]
_P2P_WITHDRAW_RE = re.compile(r"withdrawal", re.IGNORECASE)


# ----------------------------------------------------------------------------- scheme
def _norm_key(text):
    return "".join(ch for ch in str(text) if ch not in " ,،״\"'׳`‏‎")


def _read_scheme_csv(path):
    """OrderedDict cat -> {'group','fixed'} from a group,cat,fixed CSV (first spelling wins)."""
    out = OrderedDict()
    for r in read_csv(path, skip_comments=True):
        c = (r.get("cat") or "").strip()
        if c and c not in out:
            out[c] = {"group": (r.get("group") or "").strip(),
                      "fixed": (r.get("fixed") or "").strip().lower() in ("yes", "y", "1", "true", "כן")}
    return out


def load_scheme(cfg):
    """(primary, secondary, source): each scheme is an OrderedDict cat -> {'group','fixed'}."""
    ft = cfg["categories"].get("from_template")
    primary = OrderedDict()
    if ft:
        import openpyxl
        from openpyxl.utils import column_index_from_string
        path = project_path(cfg, ft["file"])
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        ws = wb[ft["sheet"]] if ft["sheet"] in wb.sheetnames else wb.worksheets[0]
        gc, cc = column_index_from_string(ft["groups_col"]), column_index_from_string(ft["cats_col"])
        group = ""
        for r in range(int(ft["first_row"]), int(ft["last_row"]) + 1):
            g, c = ws.cell(r, gc).value, ws.cell(r, cc).value
            if g not in (None, ""):
                group = str(g).strip()
            if c not in (None, ""):
                primary[str(c).strip()] = {"group": group, "fixed": str(c).strip() in cfg.get("fixed_categories", [])}
        wb.close()
        source = "template:%s" % os.path.basename(path)
        n_tpl = len(primary)
        # the workbook's category lists are template + scheme CSV (build_excel.build_scheme), so the
        # classifier must know the groups of the CSV categories a user may pick in "לסיווג ידני"
        for cat, entry in _read_scheme_csv(scheme_path(cfg)).items():
            if cat not in primary:
                primary[cat] = entry
        log("scheme %s: %d template categories + %d from %s" % (source, n_tpl, len(primary) - n_tpl,
                                                                 os.path.basename(scheme_path(cfg))))
    else:
        path = scheme_path(cfg)
        primary = _read_scheme_csv(path)
        source = "default" if path.startswith(skill_root()) else path
    # generated categories (card debits, loads, FX, P2P, club discount, abroad, unknown) are part
    # of every primary scheme: append the missing ones in memory, once, with a single log line
    added = [c for c in GENERATED_CATEGORIES if c not in primary]
    for c in added:
        primary[c] = {"group": GENERATED_CATEGORIES[c], "fixed": False}
    if added:
        log("scheme %s: %d generated categories appended in memory (also listed in the workbook): %s"
            % (source, len(added), ", ".join(added)))
    secondary = None
    sec = cfg["categories"].get("secondary_scheme")
    if sec:
        secondary = OrderedDict()
        for r in read_csv(project_path(cfg, sec), skip_comments=True):
            if r.get("cat"):
                secondary[r["cat"].strip()] = {"group": r.get("group", "").strip(), "fixed": False}
    return primary, secondary, source


class Canon(object):
    """Canonicalise category text to the exact scheme spelling (spaces/quotes/commas stripped
    before matching); `group_of(cat)` returns the scheme group or None."""

    def __init__(self, scheme):
        self.scheme = scheme or OrderedDict()
        self.keys = {_norm_key(c): c for c in self.scheme}
        self.unknown = defaultdict(int)

    def cat(self, text):
        t = str(text or "").strip()
        return self.keys.get(_norm_key(t), t)

    def group_of(self, cat):
        e = self.scheme.get(cat)
        return e["group"] if e else None


# ----------------------------------------------------------------------------- rules
def load_rules(cfg):
    """Merchant rules sorted longest-pattern-first (stable). Falls back to the shipped starter
    with a warning when the project file is missing."""
    path = project_path(cfg, "rules.merchant_rules")
    warn = None
    if not os.path.isfile(path):
        path = os.path.join(skill_root(), "references", "merchant-rules-starter.csv")
        warn = "rules/merchant_rules.csv not found - using the shipped starter %s" % path
    rules = []
    for i, r in enumerate(read_csv(path, skip_comments=True)):
        pat = (r.get("pattern") or "").strip()
        if not pat:
            continue
        rules.append({
            "pattern": pat, "lower": pat.lower(), "name_clean": (r.get("name_clean") or "").strip() or pat,
            "type": (r.get("type") or TYPE_EXPENSE).strip(), "group": (r.get("group") or r.get("group_tz") or "").strip(),
            "cat": (r.get("cat") or r.get("cat_tz") or "").strip(),
            "group2": (r.get("group2") or r.get("group_new") or "").strip(), "cat2": (r.get("cat2") or r.get("cat_new") or "").strip(),
            "note": (r.get("note") or "").strip(), "order": i, "scope": None,
        })
    rules.sort(key=lambda r: (-len(r["pattern"]), r["order"]))
    return rules, path, warn


def generated_rules(cfg):
    """Rules derived from config, scoped to the settling account (`scope` = account id)."""
    out = []
    for c in cfg["cards"]:
        if c.get("bank_debit_pattern"):
            out.append({"pattern": c["bank_debit_pattern"], "lower": c["bank_debit_pattern"].lower(),
                        "name_clean": "תשלום כרטיס אשראי – %s" % c["label"], "type": TYPE_CARD_DEBIT,
                        "group": GROUP_CARD, "cat": CAT_CARD, "group2": GROUP_CARD, "cat2": CAT_CARD,
                        "note": "חיוב הכרטיס בבנק; ההוצאות נספרות משורות הפירוט (F13)", "scope": c["settles_from"], "order": -1})
        tp = (c.get("fx_settlement") or {}).get("transfer_pattern")
        if tp:
            out.append({"pattern": tp, "lower": tp.lower(), "name_clean": "העברה לחשבון מט\"ח – %s" % c["label"],
                        "type": TYPE_INTERNAL, "group": GROUP_INTERNAL, "cat": CAT_FX, "group2": GROUP_INTERNAL, "cat2": CAT_FX,
                        "note": "העברה פנימית לחשבון המט\"ח; רק עמלת החליפין היא הוצאה (F4)", "scope": c["settles_from"], "order": -1})
    for b in cfg["benefit_programs"]:
        pat = b.get("bank_debit_pattern") or ISSUERS["club_docx"]["bank_debit_pattern_default"]
        out.append({"pattern": pat, "lower": pat.lower(), "name_clean": "תשלום כרטיס אשראי – %s" % b["label"],
                    "type": TYPE_CARD_DEBIT, "group": GROUP_CARD, "cat": CAT_CARD, "group2": GROUP_CARD, "cat2": CAT_CARD,
                    "note": "חיוב מועדון ההטבות בבנק; ההוצאות נספרות בערך נקוב + שורת הנחה מחושבת (F7)", "scope": b["settles_from"], "order": -1})
    for p in cfg["p2p"]:
        bw = p.get("bank_withdrawal_pattern")
        if bw:
            out.append({"pattern": bw, "lower": bw.lower(), "name_clean": "משיכת יתרת %s לבנק" % p["label"],
                        "type": TYPE_INTERNAL, "group": GROUP_INTERNAL, "cat": CAT_P2P_OUT, "group2": GROUP_INTERNAL, "cat2": CAT_P2P_OUT,
                        "note": "משיכת יתרת P2P לחשבון; ההכנסות נספרות משורות האפליקציה", "scope": None, "order": -1})
    out.sort(key=lambda r: -len(r["pattern"]))
    return out


def match_rule(row, rules):
    name = (row["original_name"] or "").lower()
    for r in rules:
        if r.get("scope") and not (row["_section"] == "accounts" and row["_entry"] and row["_entry"]["id"] == r["scope"]):
            continue
        if r["lower"] in name:
            return r
    return None


# ----------------------------------------------------------------------------- rows / entries
def _entry_index(cfg):
    idx = {"card": {}, "last4": {}, "label": {}}
    for c in cfg["cards"]:
        idx["card"][c["label"]] = ("cards", c)
        idx["card"][c["id"]] = ("cards", c)
        if c.get("last4"):
            idx["last4"][str(c["last4"])] = ("cards", c)
    for section in ("accounts", "benefit_programs", "p2p"):
        for e in cfg[section]:
            idx["label"][e["label"]] = (section, e)
            idx["label"][e["id"]] = (section, e)
    return idx


def resolve_entry(idx, row):
    """(section, entry) for a normalized row: by card label/id, by card last4, then by source."""
    card = row.get("card") or ""
    if card in idx["card"]:
        return idx["card"][card]
    m = re.search(r"(\d{4})\s*$", card)
    if m and m.group(1) in idx["last4"]:
        return idx["last4"][m.group(1)]
    src = row.get("source") or ""
    if src in idx["label"]:
        return idx["label"][src]
    return None, None


def load_rows(cfg):
    """All normalized CSVs (+ p2p CSVs unless overview) as dicts with _section/_entry attached."""
    idx = _entry_index(cfg)
    rows, files = [], []
    norm = sorted(glob.glob(os.path.join(project_path(cfg, "work.normalized"), "*.csv")))
    if not norm:
        raise ValueError("no normalized CSVs under %s - run the parsers first" % project_path(cfg, "work.normalized"))
    for path in norm:
        files.append(path)
        for r in read_csv(path):
            r["amount_ils"] = float(r["amount_ils"] or 0)
            r["_section"], r["_entry"] = resolve_entry(idx, r)
            rows.append(r)
    if cfg.level != "overview":
        for p in cfg["p2p"]:
            for path in resolve_files(cfg, p.get("csv")):
                files.append(path)
                for r in read_csv(path):
                    r["amount_ils"] = float(r["amount_ils"] or 0)
                    r["_section"], r["_entry"] = "p2p", p
                    rows.append(r)
    return rows, files


def set_class(row, name_clean=None, type_=None, group=None, cat=None, group2=None, cat2=None, note=None):
    if name_clean is not None:
        row["name_clean"] = name_clean
    if type_ is not None:
        row["type"] = type_
    if group is not None:
        row["group_tz"] = group
    if cat is not None:
        row["cat_tz"] = cat
    if group2 is not None:
        row["group_new"] = group2
    if cat2 is not None:
        row["cat_new"] = cat2
    if note is not None:
        row["rule_note"] = note


def add_note(row, text):
    prev = row.get("rule_note") or ""
    row["rule_note"] = (prev + " | " + text).strip(" |") if prev else text


def is_unknown(row):
    return row.get("cat_tz") == UNKNOWN_CAT


# ----------------------------------------------------------------------------- steps
def classify_rows(cfg, rows, rules, gen_rules, canon, canon2):
    level = cfg.level
    ask = level != "overview"
    for row in rows:
        for c in CLASS_COLS + ["linked_id", "trip"]:
            row.setdefault(c, "")
        r = match_rule(row, gen_rules) or match_rule(row, rules)
        if r is None:
            if row["_section"] == "accounts" and row["amount_ils"] < 0:
                set_class(row, row["original_name"], TYPE_INCOME, UNKNOWN_GROUP_INCOME, UNKNOWN_CAT,
                          UNKNOWN_GROUP_INCOME, UNKNOWN_CAT, NOTE_NO_RULE)
            else:
                set_class(row, row["original_name"], TYPE_EXPENSE, UNKNOWN_GROUP_EXPENSE, UNKNOWN_CAT,
                          UNKNOWN_GROUP_EXPENSE, UNKNOWN_CAT, NOTE_NO_RULE)
            if ask:
                add_note(row, NEEDED + " סיווג - בחר קטגוריה מהרשימה או כתוב טקסט חופשי")
        else:
            set_class(row, r["name_clean"], r["type"], r["group"], r["cat"], r["group2"] or r["group"],
                      r["cat2"] or r["cat"], r["note"])
            if r["type"] == TYPE_INCOME and row["amount_ils"] > 0:
                add_note(row, NEEDED + " כיוון התנועה - הכלל מסווג כהכנסה אך הכסף יצא")
            elif r["type"] in (TYPE_EXPENSE,) and row["_section"] == "accounts" and row["amount_ils"] < 0 and "החזר" not in r["cat"]:
                add_note(row, "כסף נכנס בשורה שסווגה כהוצאה - נספר כזיכוי/החזר (סכום שלילי)")
        # step 2: unmatched foreign currency -> abroad
        if r is None and (row.get("orig_currency") or "ILS") not in ("ILS", "") and row["type"] == TYPE_EXPENSE:
            country = country_in(row.get("details", "")) or row["orig_currency"]
            set_class(row, 'הוצאה בחו"ל (%s) – %s' % (country, row["original_name"]), TYPE_EXPENSE,
                      ABROAD_GROUP, ABROAD_CAT, ABROAD_GROUP, ABROAD_CAT,
                      "סווג אוטומטית לפי מטבע/מדינה (%s)" % country)
        # step 3: benefit-card loads -> internal
        if (row.get("details") or "").startswith(LOAD_FLAG):
            set_class(row, "טעינת כרטיס הטבות – %s" % (row["source"] or ""), TYPE_INTERNAL, GROUP_INTERNAL, CAT_LOAD,
                      GROUP_INTERNAL, CAT_LOAD, "טעינה = העברה פנימית; הרכישות מהכרטיס נספרות בערך נקוב")
        # p2p app rows: withdrawals internal, incoming income (also at any level for csv rows present)
        if row["_section"] == "p2p":
            if _P2P_WITHDRAW_RE.search(row["original_name"] or ""):
                set_class(row, "משיכת יתרת %s לבנק" % row["_entry"]["label"], TYPE_INTERNAL, GROUP_INTERNAL, CAT_P2P_OUT,
                          GROUP_INTERNAL, CAT_P2P_OUT, "משיכה לחשבון הבנק - לא הוצאה")
            elif row["amount_ils"] < 0:
                note0 = (row.get("details") or "").split("|")[0].strip()
                set_class(row, "%s נכנס – %s%s" % (row["_entry"]["app"], row["original_name"], " (%s)" % note0 if note0 else ""),
                          TYPE_INCOME, UNKNOWN_GROUP_INCOME, CAT_P2P_IN, UNKNOWN_GROUP_INCOME, CAT_P2P_IN,
                          "תקבול P2P - נספר כהכנסה (החזר מחברים); שנה קטגוריה ב'לסיווג ידני' אם זו הכנסה אחרת")
            elif is_unknown(row):
                add_note(row, "שולם מיתרת ה-P2P (לא דרך כרטיס)")
    return rows


def pair_p2p(cfg, rows):
    """F6. Returns (pairs, unmatched_card_rows)."""
    pairs, unmatched = 0, 0
    if cfg.level == "overview" or not cfg["p2p"]:
        return pairs, unmatched
    app_out = [r for r in rows if r["_section"] == "p2p" and r["amount_ils"] > 0 and r["type"] != TYPE_INTERNAL]
    used = set()
    markers = {(p.get("card_marker") or "BIT").upper(): p for p in cfg["p2p"]}
    for cr in rows:
        if cr["_section"] != "cards":
            continue
        name_u = (cr["original_name"] or "").upper()
        hit_marker = next((m for m in markers if m in name_u), None)
        if not hit_marker:
            continue
        p = markers[hit_marker]
        tol, days = float(p.get("match_tolerance", 0.01)), int(p.get("match_days", 5))
        cd = _dt.date.fromisoformat(cr["txn_date"][:10])
        best = None
        for j, ar in enumerate(app_out):
            if j in used or abs(ar["amount_ils"] - cr["amount_ils"]) > tol:
                continue
            dd = abs((_dt.date.fromisoformat(ar["txn_date"][:10]) - cd).days)
            if dd <= days and (best is None or dd < best[0]):
                best = (dd, j)
        if best is not None:
            j = best[1]
            used.add(j)
            ar = app_out[j]
            set_class(cr, "מימון P2P → %s" % ar["name_clean"], TYPE_DUPLICATE, ar["group_tz"], ar["cat_tz"], ar["group_new"], ar["cat_new"],
                      "חיוב כרטיס שמימן תשלום %s %s (%s) – ההוצאה נספרת בשורת האפליקציה בתאריך התשלום" % (ar["_entry"]["app"], ar["id"], ar["original_name"]))
            cr["linked_id"] = ar["id"]
            ar["linked_id"] = cr["id"]
            add_note(ar, "נספר כאן; מומן בחיוב כרטיס %s" % cr["id"])
            pairs += 1
        else:
            owner = owner_label(cfg, (cr["_entry"] or {}).get("owner", SHARED))
            cat = "%s – העברה לא מזוהה (BIT/PayBox)" % owner
            set_class(cr, "תשלום P2P מכרטיס – %s" % cr["original_name"], TYPE_EXPENSE, UNKNOWN_GROUP_EXPENSE, cat,
                      UNKNOWN_GROUP_EXPENSE, cat, "לא נמצאה שורת אפליקציה תואמת | " + NEEDED + " למי ועל מה (או צילום מהאפליקציה)")
            unmatched += 1
    return pairs, unmatched


def canonicalise(rows, canon, canon2):
    """Canonicalise cat_tz/cat_new spellings to the schemes and take the group from the scheme.
    Generated categories keep the group the classifier set (UNKNOWN_CAT differs by type);
    the per-owner unmatched-P2P bucket is dynamic and never reported as missing."""
    not_in = defaultdict(int)
    for row in rows:
        c = canon.cat(row["cat_tz"])
        row["cat_tz"] = c
        g = canon.group_of(c)
        if c in GENERATED_CATEGORIES:
            pass
        elif g:
            row["group_tz"] = g
        elif "העברה לא מזוהה" not in c:
            not_in[c] += 1
        if canon2 is not None:
            c2 = canon2.cat(row["cat_new"])
            row["cat_new"] = c2
            g2 = canon2.group_of(c2)
            if g2:
                row["group_new"] = g2
    return dict(not_in)


def link_discount_rows(rows):
    by_id = {r["id"]: r for r in rows}
    n = 0
    for row in rows:
        if row.get("source_file") != COMPUTED_FILE:
            continue
        set_class(row, DISCOUNT_NAME, TYPE_EXPENSE, UNKNOWN_GROUP_EXPENSE, DISCOUNT_NAME, UNKNOWN_GROUP_EXPENSE, DISCOUNT_NAME,
                  "שורה מחושבת (F7): רכישות המועדון נספרות בערך נקוב; שורה זו מקזזת את ההנחה כך שהסה\"כ תואם למזומן שיצא")
        m = re.search(r"מזהה חיוב בנק:\s*([^\s|]+)", row.get("details") or "")
        if m:
            first = m.group(1).split(",")[0]
            if first in by_id:
                row["linked_id"] = first
        n += 1
    return n


def apply_user_labels(cfg, rows, canon, canon2):
    by_id = {r["id"]: r for r in rows}
    applied = {"user_labels": 0, "interpreted": 0, "unknown_ids": []}
    trip_words = {}
    for t in cfg.get("trips", []):
        for w in t.get("countries", []):
            trip_words[w] = t["label"]
    path = project_path(cfg, "rules.user_labels")
    if os.path.isfile(path):
        for r in read_csv(path, skip_comments=True):
            rid = (r.get("id") or "").strip()
            text = (r.get("user_text") or r.get(USER_TEXT_COL) or "").strip()
            cat = (r.get("user_cat") or r.get(USER_CAT_COL) or "").strip()
            row = by_id.get(rid)
            if row is None:
                if rid:
                    applied["unknown_ids"].append(rid)
                continue
            if not text and not cat:
                continue
            if cat:
                c = canon.cat(cat)
                row["cat_tz"] = c
                g = canon.group_of(c)
                if g:
                    row["group_tz"] = g
                if canon2 is not None:
                    c2 = canon2.cat(cat)
                    if canon2.group_of(c2):
                        row["cat_new"], row["group_new"] = c2, canon2.group_of(c2)
                else:
                    row["cat_new"], row["group_new"] = row["cat_tz"], row["group_tz"]
                if row["type"] not in SUMMED_TYPES and row["type"] != TYPE_DUPLICATE:
                    pass
            if text:
                row["name_clean"] = "%s – %s" % (row["original_name"], text)
                for w, label in trip_words.items():
                    if w in text:
                        row["trip"] = label
            row["rule_note"] = "סיווג ידני של המשתמש" + (": %s" % text if text else "")
            applied["user_labels"] += 1
    path = project_path(cfg, "rules.user_labels_interpreted")
    if os.path.isfile(path):
        for r in read_csv(path, skip_comments=True):
            rid = (r.get("id") or "").strip()
            row = by_id.get(rid)
            if row is None:
                if rid:
                    applied["unknown_ids"].append(rid)
                continue
            for k in ("type", "group_tz", "cat_tz", "group_new", "cat_new", "name_clean"):
                v = (r.get(k) or "").strip()
                if v:
                    row[k] = canon.cat(v) if k == "cat_tz" else v
            if (r.get("trip") or "").strip():
                row["trip"] = r["trip"].strip()
            prev = row.get("rule_note") or ""
            note = (r.get("note") or "").strip()
            row["rule_note"] = ("סיווג לפי המשתמש: %s" % note + (" | " + prev if prev else "")).strip(" |")
            applied["interpreted"] += 1
    return applied


def add_estimates(cfg, rows, canon):
    if cfg.level == "overview":
        return 0
    out = []
    month_list = months(cfg)
    for est in cfg.get("estimates", []):
        every = max(1, int(est.get("every_n_months", 1) or 1))
        day = int(est.get("day", 15) or 15)
        cat = canon.cat(est.get("cat") or UNKNOWN_CAT)
        group = canon.group_of(cat) or est.get("group") or UNKNOWN_GROUP_EXPENSE
        for k, ym in enumerate(month_list):
            if k % every:
                continue
            d = "%s-%02d" % (ym, day)
            out.append({
                "id": stable_id("EST", est.get("id") or est.get("name") or "est", ym), "source": ESTIMATE_SOURCE, "card": "",
                "original_name": "%s (הערכה)" % (est.get("name") or est.get("id") or ""), "txn_date": d, "charge_date": d,
                "amount_ils": float(est.get("amount", 0)), "orig_currency": "ILS", "orig_amount": float(est.get("amount", 0)),
                "details": est.get("note") or "", "source_file": ESTIMATE_FILE, "row_ref": ym,
                "_section": "estimates", "_entry": None,
                "name_clean": "%s – הערכה" % (est.get("name") or ""), "type": TYPE_EXPENSE, "group_tz": group, "cat_tz": cat,
                "group_new": group, "cat_new": cat, "rule_note": "שורה מוערכת שהוסיף המשתמש (לא מהבנק)", "linked_id": "", "trip": "",
            })
        fb = est.get("funded_by_cat")
        if fb:
            fb = canon.cat(fb)
            for row in rows:
                if row.get("cat_tz") == fb and row.get("type") == TYPE_EXPENSE:
                    set_class(row, row["name_clean"] + " (מממן הוצאות מוערכות)", TYPE_INTERNAL, GROUP_INTERNAL, CAT_CASH_FUNDING,
                              GROUP_INTERNAL, CAT_CASH_FUNDING, "משיכת מזומן שמממנת שורות הערכה - לא נספרת פעמיים")
    rows.extend(out)
    return len(out)


def attribute_trips(cfg, rows):
    if cfg.level == "overview" or not cfg.get("trips"):
        return {}
    trips = []
    for t in cfg["trips"]:
        ccys = {COUNTRY_CCY[c] for c in t.get("countries", []) if c in COUNTRY_CCY}
        trips.append((t, ccys))
    counts = defaultdict(int)
    for row in rows:
        if row.get("trip"):
            counts[row["trip"]] += 1
            continue
        if not (row.get("group_tz") == ABROAD_GROUP or row.get("cat_tz") in ABROAD_CATS):
            continue
        text = " ".join(str(row.get(k) or "") for k in ("details", "name_clean", "original_name"))
        label = None
        for t, ccys in trips:
            if any(c in text for c in t.get("countries", [])):
                label = t["label"]
                break
        if label is None:
            for t, ccys in trips:
                if row.get("orig_currency") in ccys and t["start"] <= row["txn_date"][:10] <= t["end"]:
                    label = t["label"]
                    break
        if label is None:
            for t, ccys in trips:
                if t["start"] <= row["txn_date"][:10] <= t["end"]:
                    label = t["label"]
                    break
        if label is None:
            label = TRIP_UNASSIGNED
            add_note(row, NEEDED + " לאיזו נסיעה/מדינה שייכת ההוצאה? (כתוב שם מדינה בטקסט החופשי)")
        row["trip"] = label
        counts[label] += 1
    return dict(counts)


def finalise(cfg, rows):
    """Helper columns, recomputed LAST (rows were appended / retyped above)."""
    for row in rows:
        cd = str(row.get("charge_date") or row.get("txn_date") or "")[:10]
        row["charge_date"] = cd
        row["month"] = cd[:7]
        row["month_name"] = month_label(cd[:7])
        e, section = row.get("_entry"), row.get("_section")
        if e:
            row["pay"] = "%s (%s)" % (e.get("label", e["id"]), owner_label(cfg, e.get("owner", SHARED)))
            row["person"] = owner_label(cfg, e.get("owner", SHARED))
        else:
            src = row.get("source") or (ESTIMATE_SOURCE if section == "estimates" else "")
            row["pay"] = "%s (%s)" % (src, owner_label(cfg, SHARED))
            row["person"] = owner_label(cfg, SHARED)
        iw = in_window(cd, cfg)
        row["in_window"] = flag(iw)
        row["summed"] = flag(iw and row["type"] in SUMMED_TYPES)
        a = float(row["amount_ils"])
        if row["type"] == TYPE_INCOME:
            amt = -a
        elif row["type"] == TYPE_REIMBURSEMENT:
            amt = abs(a)
        else:
            amt = a
        row["amount"] = "%.2f" % amt
        row["_amount"] = amt
        for c in DB_COLUMNS:
            row.setdefault(c, "")
    rows.sort(key=lambda r: (r["charge_date"], r["source"], r["id"]))
    return rows


def summarise(cfg, rows, extra):
    by_type = defaultdict(lambda: {"count": 0, "sum": 0.0})
    win = defaultdict(lambda: {"count": 0, "sum": 0.0})
    by_group = defaultdict(float)
    unknown = 0
    for r in rows:
        by_type[r["type"]]["count"] += 1
        by_type[r["type"]]["sum"] = round(by_type[r["type"]]["sum"] + r["_amount"], 2)
        if r["summed"] == "כן":
            win[r["type"]]["count"] += 1
            win[r["type"]]["sum"] = round(win[r["type"]]["sum"] + r["_amount"], 2)
            if r["type"] == TYPE_EXPENSE:
                by_group[r["group_tz"]] = round(by_group[r["group_tz"]] + r["_amount"], 2)
            if is_unknown(r) or NEEDED in (r.get("rule_note") or ""):
                unknown += 1
    income = win[TYPE_INCOME]["sum"]
    expense = win[TYPE_EXPENSE]["sum"]
    out = {
        "ok": True, "level": cfg.level, "rows": len(rows), "months": months(cfg),
        "by_type": dict(by_type), "in_window_by_type": dict(win),
        "window_totals": {"income": income, "expense": expense, "balance": round(income - expense, 2)},
        "expenses_by_group": OrderedDict(sorted(by_group.items(), key=lambda kv: -kv[1])),
        "unknown_in_window": unknown,
    }
    out.update(extra)
    return out


# ----------------------------------------------------------------------------- main
def run(cfg):
    ensure_dirs(cfg)
    rows, files = load_rows(cfg)
    if not rows:
        raise ValueError("no normalized rows found under %s - run the parsers first" % project_path(cfg, "work.normalized"))
    primary, secondary, scheme_src = load_scheme(cfg)
    canon, canon2 = Canon(primary), (Canon(secondary) if secondary else None)
    rules, rules_path, rules_warn = load_rules(cfg)
    gen = generated_rules(cfg)
    classify_rows(cfg, rows, rules, gen, canon, canon2)
    pairs, unmatched = pair_p2p(cfg, rows)
    n_disc = link_discount_rows(rows)
    not_in = canonicalise(rows, canon, canon2)
    labels = apply_user_labels(cfg, rows, canon, canon2)
    n_est = add_estimates(cfg, rows, canon)
    trips = attribute_trips(cfg, rows)
    finalise(cfg, rows)
    out_path = project_path(cfg, "work.database")
    write_csv(out_path, rows, DB_COLUMNS)
    warnings = list(cfg.warnings)
    if rules_warn:
        warnings.append(rules_warn)
    if not_in:
        warnings.append("categories used by rules/labels but absent from the primary scheme (%s) - they are appended to the workbook lists automatically: %s" % (scheme_src, ", ".join(sorted(not_in))))
    if labels["unknown_ids"]:
        warnings.append("user label ids not found in the database (ids drift after a re-parse?): %s" % ", ".join(labels["unknown_ids"][:20]))
    summary = summarise(cfg, rows, {
        "database": out_path, "inputs": [os.path.basename(f) for f in files], "scheme": scheme_src,
        "scheme_categories": len(primary), "secondary_scheme": bool(secondary), "rules": len(rules), "rules_file": rules_path,
        "generated_rules": len(gen), "p2p_pairs": pairs, "p2p_unmatched_card_rows": unmatched,
        "discount_rows": n_disc, "estimate_rows": n_est, "trips": trips, "user_labels_applied": labels,
        "categories_not_in_scheme": not_in, "warnings": warnings,
    })
    return summary


def main(argv=None):
    args, cfg = parse_args("classify normalized rows into work/database.csv", argv=argv)
    try:
        summary = run(cfg)
    except Exception as e:  # noqa: BLE001
        fail("classify: %s" % e, "run the parsers first; check rules/merchant_rules.csv and the category scheme")
        return
    for w in summary["warnings"]:
        log("WARNING: " + w)
    log("rows=%d | in-window income %.2f expense %.2f | unknown %d | p2p pairs %d" % (
        summary["rows"], summary["window_totals"]["income"], summary["window_totals"]["expense"],
        summary["unknown_in_window"], summary["p2p_pairs"]))
    emit(summary)


if __name__ == "__main__":
    main()
