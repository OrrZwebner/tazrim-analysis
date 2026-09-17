#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""make_fixtures.py — generate a complete SYNTHETIC tazrim project for tests.

Purpose : write a small project (config, bank/card/benefit/p2p inputs, BOI rate cache,
          template workbook, merchant rules) in the exact institutional layouts described in
          formats.py, plus `expected.json` with hand-derived expected values, so every parser
          and the classifier can be tested without any real data.
Inputs  : --out DIR (required). The standard --config/--project-dir/--level flags are accepted
          and ignored (no config is read).
Outputs : the files under DIR (listed in the JSON); exactly one JSON object on stdout:
          {"ok": true, "out": DIR, "files": [...], "expected": {...}}.
Exit    : 0 on success; 1 with {"ok": false, ...} on failure (e.g. missing openpyxl/python-docx).

Every merchant, person, amount and date below is invented. Expected values in EXPECTED are
typed as literals (derived by hand from the row lists in this file), never computed from them.
Python 3.8 compatible. Needs openpyxl and python-docx.
"""
import csv
import datetime as dt
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import NORMALIZED_COLUMNS, build_parser, emit, fail, log, write_csv, write_json  # noqa: E402
from formats import BANKS, BOI, ISSUERS  # noqa: E402

try:
    import openpyxl
    from docx import Document
except ImportError as e:  # pragma: no cover
    fail("missing dependency: %s" % e, "pip install -r requirements.txt")

D = dt.datetime  # datetime cells for openpyxl

# --------------------------------------------------------------------------- config
CONFIG = {
    "language": "he",
    "analysis_level": "standard",
    "window": {"start": "2025-01-01", "end": "2025-02-28", "month_rule": "charge_date"},
    "household": {
        "label": "משק בית לדוגמה",
        "people": [{"id": "p1", "label": "בן/בת זוג א'"}, {"id": "p2", "label": "בן/בת זוג ב'"}],
        "shared_label": "משותף (עו\"ש)",
    },
    "accounts": [
        {"id": "bank_a", "bank": "bank_xlsx_a", "label": "עו\"ש חשבון א'", "owner": "shared",
         "files": "inputs/bank/bank_a/*.xlsx"},
        {"id": "bank_b", "bank": "bank_xls_b", "label": "עו\"ש חשבון ב'", "owner": "p2",
         "files": "inputs/bank/bank_b/*.xls*", "opening_balance": 3000.00},
    ],
    "cards": [
        {"id": "card_1234", "issuer": "card_cycle_xlsx", "last4": "1234", "label": "ויזה 1234", "owner": "p1",
         "files": "inputs/cards/card_cycle/card_*.xlsx", "settles_from": "bank_a",
         "bank_debit_pattern": "חיוב לכרטיס ויזה 1234", "cycle_day": 10,
         "fx_settlement": {"currency": "USD"}},
        {"id": "card_5678", "issuer": "card_cycle_xlsx", "last4": "5678", "label": "ויזה 5678", "owner": "p2",
         "files": "inputs/cards/card_cycle/card_*.xlsx", "settles_from": "bank_a",
         "bank_debit_pattern": "חיוב לכרטיס ויזה 5678", "cycle_day": 10,
         "fx_settlement": {"currency": "EUR"}},
        {"id": "card_9012", "issuer": "card_blocks_xls", "last4": "9012", "label": "מאסטרקארד 9012",
         "owner": "p2", "files": "inputs/cards/card_9012/9012_*.xls*", "settles_from": "bank_b",
         "bank_debit_pattern": "9012 - כרטיס אשראי", "cycle_day": 2},
    ],
    "benefit_programs": [
        {"id": "club", "issuer": "club_docx", "label": "מועדון הטבות", "owner": "p2",
         "files": "inputs/benefits/club/*.docx", "settles_from": "bank_b",
         "bank_debit_pattern": "חיוב מועדון הטבות", "default_debit_day": 2},
    ],
    "p2p": [
        {"id": "bit_p1", "app": "BIT", "owner": "p1", "csv": "inputs/p2p/bit_p1.csv",
         "screenshots": None, "card_marker": "BIT", "match_days": 5, "match_tolerance": 0.01},
    ],
    "categories": {
        "scheme": "default",
        "from_template": {"file": "inputs/template/תזרים.xlsx", "sheet": "הכנסות - הוצאות",
                          "groups_col": "A", "cats_col": "B", "first_row": 2, "last_row": 7},
        "secondary_scheme": None,
    },
    "rules": {"merchant_rules": "rules/merchant_rules.csv", "user_labels": "rules/user_labels.csv",
              "user_labels_interpreted": "rules/user_labels_interpreted.csv",
              "fx_overrides": "rules/fx_overrides.csv"},
    "fixed_categories": ["שכר דירה", "ועד בית"],
    "fx": {"rates_dir": "work/boi_rates", "source": "boi_sdmx", "manual_rates": None},
    "verify": {"excel_recalc": "never"},
    "dashboard": {"chartjs": "cdn"},
    "thresholds": {"highlight_expense_avg": 3000, "highlight_income_avg": 12000},
}

# --------------------------------------------------------------------------- bank: format A (.xlsx)
# Newest first; file sign +credit/-debit; balance on every row (chain: bal[i]-amt[i] == bal[i+1]).
BANK_A_ROWS = [  # (date, value date, description, amount, balance, reference, fee, channel)
    (D(2025, 2, 10), D(2025, 2, 10), "חיוב לכרטיס ויזה 5678", -55.00, 9165.00, 1007, None, "אינטרנט"),
    (D(2025, 2, 10), D(2025, 2, 10), "חיוב לכרטיס ויזה 1234", -120.00, 9220.00, 1006, None, "אינטרנט"),
    (D(2025, 1, 20), D(2025, 1, 20), "העברה מהחשבון", -1000.00, 9340.00, 1005, None, "אינטרנט"),
    (D(2025, 1, 15), D(2025, 1, 15), "עמלת ניהול חשבון", -10.00, 10340.00, 1004, None, "סניף"),
    (D(2025, 1, 10), D(2025, 1, 10), "חיוב לכרטיס ויזה 1234", -650.00, 10350.00, 1003, None, "אינטרנט"),
    (D(2025, 1, 2), D(2025, 1, 2), "שכר דירה לדוגמה", -2000.00, 11000.00, 1002, None, "הוראת קבע"),
    (D(2025, 1, 1), D(2025, 1, 1), "משכורת לדוגמה", 8000.00, 13000.00, 1001, None, "זיכוי"),
]

# --------------------------------------------------------------------------- bank: format B (.xlsx twin of the .xls layout)
# Oldest first; separate credit/debit; balance only on the last row of a date group.
BANK_B_OPENING = 3000.00
BANK_B_ROWS = [  # (balance or None, value date, credit, debit, description, reference, op type, date)
    (None, D(2025, 2, 2), None, 300.00, "9012 - כרטיס אשראי", 2001, "חיוב", D(2025, 2, 2)),
    (2520.00, D(2025, 2, 2), None, 180.00, "חיוב מועדון הטבות", 2002, "חיוב", D(2025, 2, 2)),
]

# --------------------------------------------------------------------------- cards: cycle file (all cards)
CYCLE_HEADER = ISSUERS["card_cycle_xlsx"]["header"]
CYCLE_FILES = {
    # card_0225 = cycle charged 10/01/2025. Row 2 has NO פירוט cell -> charge date shifts to col 7.
    "card_0225.xlsx": {
        "row7": "4 עסקאות לחיוב בחודש ינואר בכל הכרטיסים",
        "row8": 'עסקאות באשראי - סה"כ חיוב:  + ₪650 + $100',
        "rows": [
            ["ויזה 1234", "סופר לדוגמה", "15/12/2024", 400, "מנפיק לדוגמה", "ישראל", " רגילה", "10/01/2025", 400, "--"],
            ["ויזה 1234", "חנות בגדים לדוגמה", "03/01/2025", 250, "מנפיק לדוגמה", "ישראל", "10/01/2025", 250, "--"],
            ["ויזה 1234", "EXAMPLE STORE US", "05/01/2025", 100, "מנפיק לדוגמה", 'חו"ל', " ארצות הברית", "12/01/2025", 100, "--"],
            ["ויזה 5678", "הוראת קבע לדוגמה", "01/01/2025", 55, "מנפיק לדוגמה", "ישראל", "הוראת קבע", "--", 55, "--"],
        ],
    },
    # card_0325 = cycle charged 10/02/2025: the standing order now charged + one BIT-funded row.
    "card_0325.xlsx": {
        "row7": "2 עסקאות לחיוב בחודש פברואר בכל הכרטיסים",
        "row8": 'עסקאות באשראי - סה"כ חיוב:  + ₪175',
        "rows": [
            ["ויזה 5678", "הוראת קבע לדוגמה", "01/01/2025", 55, "מנפיק לדוגמה", "ישראל", "הוראת קבע", "10/02/2025", 55, "--"],
            ["ויזה 1234", "BIT העברה לדוגמה", "20/01/2025", 120, "מנפיק לדוגמה", "ישראל", " רגילה", "10/02/2025", 120, "--"],
        ],
    },
}

# --------------------------------------------------------------------------- cards: per-card statement (detail .xlsx)
BLOCKS_HEADER = ISSUERS["card_blocks_xls"]["detail_xlsx"]["header"]
ISR_TITLE = ["פירוט עסקאות", "מאסטרקארד 9012", "לחיוב ב-02.02", "שנת 2025"]
ISR_ROWS = [
    ["15.01.25", "סופר לדוגמה", 200, "₪", 200, "₪", "000001", ""],
    ["20.01.25", "מסעדה לדוגמה", 100, "₪", 100, "₪", "000002", ""],
]
BLOCKS_TOTAL = ['סה"כ לחיוב החודש בכרטיס בש"ח', 300]

# --------------------------------------------------------------------------- benefits docx
BHZ = ISSUERS["club_docx"]
DOCX_PARAGRAPHS = (
    ["פירוט רכישות", ""]
    + BHZ["benefits_segment"]["header_sequence"]
    + ["מועדון לדוגמה", "20.01.25", "מוצר לדוגמה", "הטבה לדוגמה", "1", "₪30", "מומש", "", ""]
    + ["", "כרטיס נטען", ""]
    + BHZ["card_segment"]["header_sequence"]
    + ["15/01/25", "1001", "", "טעינה", "", "טעינה", "₪200"]
    + ["25/01/25", "1002", "", "סופר לדוגמה", "סניף לדוגמה", "חיוב", "₪-70"]
    + ["", "סוף הרשימה"]
)

# --------------------------------------------------------------------------- p2p csv (normalized schema)
P2P_ROWS = [
    {"id": "BIT-000001", "source": "BIT בן/בת זוג א'", "card": "BIT", "original_name": "חבר לדוגמה",
     "txn_date": "2025-01-21", "charge_date": "2025-01-21", "amount_ils": "120.00",
     "orig_currency": "ILS", "orig_amount": "120.00", "details": "ארוחה משותפת | הושלם | יוצא",
     "source_file": "bit_p1.csv", "row_ref": 2},
    {"id": "BIT-000002", "source": "BIT בן/בת זוג א'", "card": "BIT", "original_name": "חברה לדוגמה",
     "txn_date": "2025-02-05", "charge_date": "2025-02-05", "amount_ils": "-60.00",
     "orig_currency": "ILS", "orig_amount": "60.00", "details": "החזר | הושלם | נכנס",
     "source_file": "bit_p1.csv", "row_ref": 3},
]

# --------------------------------------------------------------------------- BOI cache (SDMX shape)
BOI_RATES = {
    "USD": [("2025-01-10", 3.45), ("2025-01-12", 3.50), ("2025-01-13", 3.52), ("2025-02-10", 3.55)],
    "EUR": [("2025-01-10", 3.78), ("2025-01-12", 3.80), ("2025-01-13", 3.81), ("2025-02-10", 3.85)],
}

# --------------------------------------------------------------------------- template + rules
TEMPLATE_SHEET = "הכנסות - הוצאות"
TEMPLATE_ROWS = [("מזון", "סופרמרקט"), ("מזון", "מסעדות"), ("דיור", "שכר דירה"),
                 ("דיור", "ועד בית"), ("תחבורה", "דלק"), ("תחבורה", "תחבורה ציבורית")]
RULES_COLUMNS = ["pattern", "name_clean", "type", "group", "cat", "group2", "cat2", "note"]
RULES = [
    ("סופר לדוגמה", "סופר לדוגמה", "הוצאה", "מזון", "סופרמרקט", "מזון", "סופרמרקט", ""),
    ("מסעדה לדוגמה", "מסעדה לדוגמה", "הוצאה", "מזון", "מסעדות", "מזון", "מסעדות", ""),
    ("שכר דירה לדוגמה", "שכר דירה", "הוצאה", "דיור", "שכר דירה", "דיור", "שכר דירה", ""),
    ("משכורת לדוגמה", "משכורת", "הכנסה", "הכנסות", "משכורת", "הכנסות", "משכורת", ""),
]

# --------------------------------------------------------------------------- expected (hand-derived literals)
EXPECTED = {
    "window": {"start": "2025-01-01", "end": "2025-02-28", "months": ["2025-01", "2025-02"]},
    "bank_a": {
        "rows": 7, "inflows": 8000.00, "outflows": 3835.00, "net": 4165.00,
        "opening_balance_implied": 5000.00, "closing_balance": 9165.00, "balance_breaks": 0,
        "card_debits": {"2025-01-10|1234": 650.00, "2025-02-10|1234": 120.00, "2025-02-10|5678": 55.00},
    },
    "bank_b": {
        "rows": 2, "opening_balance": 3000.00, "closing_balance": 2520.00,
        "inflows": 0.00, "outflows": 480.00, "balance_checkpoints": 1, "balance_mismatches": 0,
        "card_debits": {"2025-02-02|9012": 300.00}, "max_debits": {"2025-02-02": 180.00},
    },
    "card_cycle": {
        "files": {"card_0225.xlsx": {"cycle_date": "2025-01-10", "parsed": 4, "kept": 3},
                  "card_0325.xlsx": {"cycle_date": "2025-02-10", "parsed": 2, "kept": 2}},
        "kept_rows": 5,
        "uncharged_dropped": 1, "uncharged_kept": 0, "cross_file_duplicates": 0,
        "reconcile": [
            {"cycle": "2025-01-10", "card": "ויזה 1234", "ils_sum": 650.00, "bank": 650.00, "diff": 0.00,
             "fx_sum": 100.00, "fx_ccy": "USD"},
            {"cycle": "2025-02-10", "card": "ויזה 1234", "ils_sum": 120.00, "bank": 120.00, "diff": 0.00},
            {"cycle": "2025-02-10", "card": "ויזה 5678", "ils_sum": 55.00, "bank": 55.00, "diff": 0.00},
        ],
        "fx_row": {"merchant": "EXAMPLE STORE US", "charge_date": "2025-01-12", "charge_usd": 100.00,
                   "boi_rate": 3.50, "amount_ils": 350.00, "orig_currency": "USD"},
        "total_amount_ils": 1175.00,
        "column_shift_row": {"file": "card_0225.xlsx", "merchant": "חנות בגדים לדוגמה", "charge_date": "2025-01-10",
                             "charge_amount": 250.00},
    },
    "card_blocks": {
        "rows": 2, "card": "9012", "charge_date": "2025-02-02", "block_total": 300.00,
        "bank_debit": 300.00, "diff": 0.00,
    },
    "club": {
        "benefits": 1, "loads": 1, "purchases": 1,
        "benefits_total": 30.00, "loads_total": 200.00, "purchases_total": 70.00,
        "activity_month": "2025-01", "bank_debit_date": "2025-02-02", "bank_debit": 180.00,
        "loads_plus_benefits": 230.00, "discount": 50.00, "discount_row_amount": -50.00,
    },
    "p2p": {
        "outgoing": [{"amount": 120.00, "date": "2025-01-21", "matches_card_row": {"file": "card_0325.xlsx",
                      "merchant": "BIT העברה לדוגמה", "txn_date": "2025-01-20", "delta_days": 1}}],
        "incoming": [{"amount": -60.00, "date": "2025-02-05", "type": "הכנסה"}],
        "pairs": 1,
    },
    "classify": {
        "note": "assumes the 4 fixture rules only; unmatched rows default per P10 step 1",
        "unmatched_rows_to_manual_sheet": ["חנות בגדים לדוגמה", "הוראת קבע לדוגמה", "העברה מהחשבון",
                                           "עמלת ניהול חשבון", "EXAMPLE STORE US (→ abroad by currency)"],
        "card_debit_rows_non_summed": 5,
        "duplicate_rows": 1,
        "internal_rows": 1,
        "income_in_window": 8060.00,
        "expense_in_window_if_all_unknowns_are_expenses": 4535.00,
        "expense_breakdown": {"שכר דירה": 2000.00, "עמלה": 10.00, "העברה מהחשבון": 1000.00,
                              "סופר לדוגמה": 600.00, "חנות בגדים לדוגמה": 250.00, "EXAMPLE STORE US": 350.00,
                              "הוראת קבע לדוגמה": 55.00, "BIT app row": 120.00, "מסעדה לדוגמה": 100.00,
                              "club benefit": 30.00, "club purchase": 70.00, "הנחת מועדון": -50.00},
    },
    "template": {"sheet": TEMPLATE_SHEET, "groups": 3, "categories": 6, "range": "A2:B7"},
    "rules": {"patterns": 4},
}


# --------------------------------------------------------------------------- writers
def _mk(path):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    return path


def write_bank_a(path):
    spec = BANKS["bank_xlsx_a"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.cell(1, 1, "עובר ושב — לקוח לדוגמה")
    for c, h in enumerate(spec["header"], 1):
        ws.cell(spec["header_row"], c, h)
    for r, row in enumerate(BANK_A_ROWS, spec["data_from_row"]):
        for c, v in enumerate(row, 1):
            if v is not None:
                ws.cell(r, c, v)
    wb.save(_mk(path))


def write_bank_b(path):
    spec = BANKS["bank_xls_b"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.cell(1, 2, "תנועות בחשבון — לקוח לדוגמה")
    for c, h in enumerate(spec["header"], 1):
        if h:
            ws.cell(spec["header_row"], c, h)
    ws.cell(spec["opening_row"], 2, BANK_B_OPENING)
    ws.cell(spec["opening_row"], 6, spec["opening_marker"])
    for r, row in enumerate(BANK_B_ROWS, spec["data_from_row"]):
        for c, v in enumerate(row, 2):
            if v is not None:
                ws.cell(r, c, v)
    wb.save(_mk(path))


def write_card_cycle(path, fspec):
    spec = ISSUERS["card_cycle_xlsx"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    ws.cell(spec["title_row"], 1, "פירוט עסקאות - כל הכרטיסים הבנקאיים ")
    ws.cell(7, 1, fspec["row7"])
    ws.cell(8, 1, fspec["row8"])
    for c, h in enumerate(CYCLE_HEADER, 1):
        ws.cell(spec["header_row"], c, h)
    r = spec["data_from_row"]
    for row in fspec["rows"]:
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        r += 1
    ws.cell(r, 1, "הודעה: הנתונים לדוגמה בלבד")
    wb.save(_mk(path))


def write_card_blocks(path):
    spec = ISSUERS["card_blocks_xls"]["detail_xlsx"]
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = spec["sheet"]
    for r, t in enumerate(ISR_TITLE, 1):
        ws.cell(r, 1, t)
    hrow = 6
    for c, h in enumerate(BLOCKS_HEADER, 1):
        ws.cell(hrow, c, h)
    r = hrow + 1
    for row in ISR_ROWS:
        for c, v in enumerate(row, 1):
            ws.cell(r, c, v)
        r += 1
    ws.cell(r, 1, BLOCKS_TOTAL[0])
    ws.cell(r, 2, BLOCKS_TOTAL[1])
    wb.save(_mk(path))


def write_docx(path):
    doc = Document()
    for p in DOCX_PARAGRAPHS:
        doc.add_paragraph(p)
    doc.save(_mk(path))


def write_boi(rates_dir):
    out = []
    for ccy, rows in BOI_RATES.items():
        path = _mk(os.path.join(rates_dir, BOI["cache_file_template"].format(ccy=ccy)))
        with open(path, "w", encoding="utf-8", newline="") as fh:
            w = csv.writer(fh)
            w.writerow(BOI["csv_columns"])
            for day, rate in rows:
                w.writerow(["RER_%s_ILS" % ccy, "D", ccy, "ILS", "ILS", "OF00", "BOI_MRKT", "V", "F",
                            "Y", 0, "", day, rate, "YP"])
        out.append(path)
    return out


def write_template(path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = TEMPLATE_SHEET
    ws.sheet_view.rightToLeft = True
    ws["A1"], ws["B1"], ws["C1"] = "קבוצה", "קטגוריה", "ממוצע חודשי"
    for r, (g, c) in enumerate(TEMPLATE_ROWS, 2):
        ws.cell(r, 1, g)
        ws.cell(r, 2, c)
    wb.save(_mk(path))


def write_rules(path):
    _mk(path)
    with open(path, "w", encoding="utf-8-sig", newline="") as fh:
        fh.write("# synthetic merchant rules — pattern is a case-insensitive substring, first match wins\n")
        w = csv.writer(fh)
        w.writerow(RULES_COLUMNS)
        for r in RULES:
            w.writerow(r)


def generate(out):
    out = os.path.abspath(out)
    files = []

    cfg_path = os.path.join(out, "tazrim.config.json")
    write_json(cfg_path, CONFIG)
    files.append(cfg_path)

    p = os.path.join(out, "inputs", "bank", "bank_a", "עובר ושב.xlsx")
    write_bank_a(p); files.append(p)
    p = os.path.join(out, "inputs", "bank", "bank_b", "עובר ושב.xlsx")
    write_bank_b(p); files.append(p)
    for name, fspec in CYCLE_FILES.items():
        p = os.path.join(out, "inputs", "cards", "card_cycle", name)
        write_card_cycle(p, fspec); files.append(p)
    p = os.path.join(out, "inputs", "cards", "card_9012", "9012_0225.xlsx")
    write_card_blocks(p); files.append(p)
    p = os.path.join(out, "inputs", "benefits", "club", "פירוט.docx")
    write_docx(p); files.append(p)
    p = _mk(os.path.join(out, "inputs", "p2p", "bit_p1.csv"))
    write_csv(p, P2P_ROWS, NORMALIZED_COLUMNS); files.append(p)
    files += write_boi(os.path.join(out, "work", "boi_rates"))
    p = os.path.join(out, "inputs", "template", "תזרים.xlsx")
    write_template(p); files.append(p)
    p = os.path.join(out, "rules", "merchant_rules.csv")
    write_rules(p); files.append(p)
    p = os.path.join(out, "expected.json")
    write_json(p, EXPECTED); files.append(p)
    for d in ("work/normalized", "outputs", "notes"):
        os.makedirs(os.path.join(out, d), exist_ok=True)
    return out, files


def main(argv=None):
    def extra(p):
        p.add_argument("--out", required=True, help="target directory for the synthetic project")
    args = build_parser(__doc__.splitlines()[0], extra).parse_args(argv)
    try:
        out, files = generate(args.out)
    except Exception as e:  # noqa: BLE001
        fail("fixture generation failed: %s" % e, "check --out is writable and openpyxl/python-docx are installed")
        return
    log("wrote %d files under %s" % (len(files), out))
    emit({"ok": True, "out": out, "config": os.path.join(out, "tazrim.config.json"),
          "files": [os.path.relpath(f, out) for f in files], "expected": EXPECTED})


if __name__ == "__main__":
    main()
