#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""apply_user_labels.py — the re-label loop (P18, mode B).

Purpose : read the user's entries in the "לסיווג ידני" sheet of the current workbook (columns
          found by header name: מזהה, "למילוי משתמש (טקסט חופשי)", "קטגוריה (בחירה מהרשימה)"),
          merge them into rules/user_labels.csv (id, user_text, user_cat, timestamp — the sheet
          wins over an older entry for the same id; earlier ids that no longer appear in the sheet
          are kept, so a label survives every rebuild), then run the downstream steps for the
          level as subprocesses of the sibling scripts, stopping at the first failure:
            overview : classify → build_excel → verify → make_summary → make_figures → make_report_html
            standard / deep : classify → build_excel → add_transfers_sheet → verify → make_summary
                              → make_figures → make_report_html → build_dashboard
          A sibling script that does not exist yet is reported with status "missing" and skipped.
Inputs  : tazrim.config.json, outputs/תזרים.xlsx (saved by the user after editing the sheet).
Outputs : rules/user_labels.csv; one JSON object on stdout {ok, labels: {read, merged, file},
          steps: [{name, status: ok|fail|missing|skipped, seconds, result}], stopped_at}.
Exit    : 0 when every executed step succeeded; 1 when a step failed or the workbook/sheet is
          missing; 2 on a config error.

Remind the user to close the workbook WITHOUT saving before reopening the rebuilt file (Excel's
in-memory copy would overwrite the rebuild, FM25). Python 3.8 compatible. Needs openpyxl.
"""
import datetime as _dt
import json
import os
import subprocess
import sys
import time

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (  # noqa: E402
    DB_HEADERS_HE, SHEET_MANUAL, USER_CAT_COL, USER_TEXT_COL, emit, fail, log, parse_args,
    project_path, read_csv, write_csv,
)

try:
    import openpyxl
except ImportError as _e:  # pragma: no cover
    fail("missing dependency: %s" % _e, "pip install -r requirements.txt (openpyxl)")

HERE = os.path.dirname(os.path.abspath(__file__))
LABEL_COLUMNS = ["id", "user_text", "user_cat", "timestamp"]
STEPS_BY_LEVEL = {
    "overview": ["classify", "build_excel", "verify", "make_summary", "make_figures", "make_report_html"],
    "standard": ["classify", "build_excel", "add_transfers_sheet", "verify", "make_summary", "make_figures",
                 "make_report_html", "build_dashboard"],
}
STEPS_BY_LEVEL["deep"] = list(STEPS_BY_LEVEL["standard"])


def read_sheet_labels(xlsx):
    """[(id, user_text, user_cat)] from the לסיווג ידני sheet; header row located by 'מזהה'."""
    wb = openpyxl.load_workbook(xlsx, read_only=True, data_only=True)
    try:
        if SHEET_MANUAL not in wb.sheetnames:
            fail("sheet %r not found in %s" % (SHEET_MANUAL, xlsx),
                 "the sheet exists from level standard; rebuild with build_excel.py --level standard|deep")
        hdrs, rows = None, []
        for row in wb[SHEET_MANUAL].iter_rows(values_only=True):
            if hdrs is None:
                if row and row[0] == DB_HEADERS_HE["id"]:
                    hdrs = list(row)
                continue
            if row and row[0]:
                d = dict(zip(hdrs, row))
                t = str(d.get(USER_TEXT_COL) or "").strip()
                c = str(d.get(USER_CAT_COL) or "").strip()
                if t or c:
                    rows.append((str(row[0]), t, c))
        if hdrs is None:
            fail("header row with %r not found in sheet %r" % (DB_HEADERS_HE["id"], SHEET_MANUAL), "rebuild the workbook")
        for col in (USER_TEXT_COL, USER_CAT_COL):
            if col not in hdrs:
                fail("column %r missing in sheet %r" % (col, SHEET_MANUAL), "rebuild the workbook with build_excel.py")
        return rows
    finally:
        wb.close()


def merge_labels(path, sheet_rows):
    """Existing CSV rows + sheet rows (sheet wins per id). Returns the merged list."""
    existing = {}
    if os.path.isfile(path):
        for r in read_csv(path):
            if r.get("id"):
                existing[r["id"]] = {k: r.get(k, "") for k in LABEL_COLUMNS}
    now = _dt.datetime.now().isoformat(timespec="seconds")
    for rid, t, c in sheet_rows:
        prev = existing.get(rid)
        if prev and prev["user_text"] == t and prev["user_cat"] == c:
            continue
        existing[rid] = {"id": rid, "user_text": t, "user_cat": c, "timestamp": now}
    return [existing[k] for k in sorted(existing)]


def run_step(name, cfg, level):
    script = os.path.join(HERE, name + ".py")
    if not os.path.isfile(script):
        return {"name": name, "status": "missing", "seconds": 0.0, "result": {"error": "script not found: %s" % script}}
    cmd = [sys.executable, script, "--config", cfg.path, "--project-dir", cfg.project_dir, "--level", level]
    t0 = time.time()
    try:
        r = subprocess.run(cmd, capture_output=True, text=True, timeout=1800)
    except subprocess.TimeoutExpired:
        return {"name": name, "status": "fail", "seconds": round(time.time() - t0, 1), "result": {"error": "timeout"}}
    last = [ln for ln in r.stdout.strip().splitlines() if ln.strip()]
    try:
        result = json.loads(last[-1]) if last else {}
    except ValueError:
        result = {"stdout_tail": r.stdout[-500:]}
    if r.stderr.strip():
        log("---- %s stderr (tail):\n%s" % (name, r.stderr.strip()[-1500:]))
    ok = r.returncode == 0 and result.get("ok", True) is not False
    if not ok and "error" not in result:
        result["error"] = "exit code %d" % r.returncode
    return {"name": name, "status": "ok" if ok else "fail", "seconds": round(time.time() - t0, 1), "result": result}


def main(argv=None):
    def extra(p):
        p.add_argument("--workbook", default=None, help="workbook to read (default: config outputs.workbook)")
        p.add_argument("--no-run", action="store_true", help="only update rules/user_labels.csv; do not run the pipeline")
    args, cfg = parse_args(__doc__.splitlines()[0], extra, argv)
    xlsx = args.workbook or project_path(cfg, "outputs.workbook")
    if not os.path.isfile(xlsx):
        fail("workbook not found: %s" % xlsx, "run build_excel.py first")
    sheet_rows = read_sheet_labels(xlsx)
    labels_path = project_path(cfg, "rules.user_labels")
    merged = merge_labels(labels_path, sheet_rows)
    write_csv(labels_path, merged, LABEL_COLUMNS)
    log("%d labels read from the sheet; %d in %s" % (len(sheet_rows), len(merged), labels_path))
    steps, stopped = [], None
    if not args.no_run:
        for name in STEPS_BY_LEVEL[cfg.level]:
            st = run_step(name, cfg, cfg.level)
            steps.append(st)
            log("== %s: %s (%.1fs)" % (name, st["status"], st["seconds"]))
            if st["status"] == "fail":
                stopped = name
                break
    ok = stopped is None
    emit({"ok": ok, "labels": {"read": len(sheet_rows), "merged": len(merged), "file": labels_path},
          "steps": steps, "stopped_at": stopped,
          "reminder": "סגור את הקובץ באקסל בלי לשמור לפני שפותחים את הקובץ המעודכן | close the workbook WITHOUT saving before reopening"})
    if not ok:
        sys.exit(1)


if __name__ == "__main__":
    main()
