# -*- coding: utf-8 -*-
"""End-to-end tests of the whole chain on the synthetic fixture (make_fixtures.py).

Mode A (first run, level standard, verify.excel_recalc = never): every step ok, database rows,
workbook sheet set, cash-reconciliation residual == 0, dashboard row count, report sections,
per-card / per-person breakdowns against expected.json.
Mode B (re-label): a user text + category typed into "לסיווג ידני" -> rules/user_labels.csv,
the database row, and the rebuilt workbook.
Mode C (add a month): a third card-cycle file -> existing ids unchanged, new rows present.
Levels overview and deep: the run succeeds with the level's sheet set.
Run: python3 -m unittest discover -s scripts/tests -v   (about 40 s; needs pandas/openpyxl/matplotlib)
"""
import csv
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

import common  # noqa: E402
import make_fixtures  # noqa: E402

try:
    import openpyxl
    import pandas as pd
    import matplotlib  # noqa: F401
    from docx import Document  # noqa: F401
except ImportError:  # pragma: no cover
    openpyxl = pd = None

CYCLE_SHEET = common.SHEET_MANUAL


def run(script, cfg_path, *extra):
    """Run a sibling script; return (returncode, last JSON object on stdout, stderr)."""
    cmd = [sys.executable, os.path.join(SCRIPTS, script), "--config", cfg_path] + list(extra)
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=600)
    lines = [ln for ln in r.stdout.strip().splitlines() if ln.strip()]
    try:
        out = json.loads(lines[-1]) if lines else {}
    except ValueError:
        out = {"stdout_tail": r.stdout[-500:]}
    return r.returncode, out, r.stderr


def read_db(project):
    return pd.read_csv(os.path.join(project, "work", "database.csv"), encoding="utf-8-sig", dtype=str,
                       keep_default_na=False)


@unittest.skipIf(openpyxl is None or pd is None, "pandas / openpyxl / matplotlib / python-docx not installed")
class TestPipelineEndToEnd(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="tz_e2e_")
        cls.project, _ = make_fixtures.generate(os.path.join(cls.tmp, "proj"))
        cls.cfg_path = os.path.join(cls.project, "tazrim.config.json")
        with open(cls.cfg_path, encoding="utf-8") as fh:
            raw = json.load(fh)
        raw["verify"] = {"excel_recalc": "never"}
        with open(cls.cfg_path, "w", encoding="utf-8") as fh:
            json.dump(raw, fh, ensure_ascii=False, indent=1)
        cls.expected = make_fixtures.EXPECTED
        cls.rc, cls.res, cls.err = run("run_pipeline.py", cls.cfg_path, "--mode", "A", "--level", "standard")

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def steps(self, res):
        return {s["name"]: s for s in res["steps"]}

    # ------------------------------------------------------------------ mode A
    def test_01_mode_a_every_step_ok(self):
        self.assertEqual(self.rc, 0, self.err[-2000:])
        self.assertTrue(self.res["ok"])
        st = self.steps(self.res)
        self.assertEqual(sorted(st), sorted(["parse_bank", "parse_card_cycle", "parse_card_blocks", "parse_benefits", "classify",
                                             "build_excel", "add_transfers_sheet", "verify", "make_summary",
                                             "make_figures", "make_report_html", "build_dashboard"]))
        for name, s in st.items():
            self.assertEqual(s["status"], "ok", "%s: %s" % (name, s["result"]))
        self.assertIsNone(self.res["stopped_at"])

    def test_02_database_rows(self):
        db = read_db(self.project)
        self.assertEqual(list(db.columns), common.DB_COLUMNS)
        E = self.expected
        # bank rows + kept card rows + per-card statement + docx (benefits+loads+purchases+discount row) + p2p
        n_expected = (E["bank_a"]["rows"] + E["bank_b"]["rows"] + E["card_cycle"]["kept_rows"] + E["card_blocks"]["rows"]
                      + E["club"]["benefits"] + E["club"]["loads"] + E["club"]["purchases"] + 1
                      + len(E["p2p"]["outgoing"]) + len(E["p2p"]["incoming"]))
        self.assertEqual(len(db), n_expected)
        self.assertEqual(self.steps(self.res)["classify"]["result"]["rows"], n_expected)
        amt = pd.to_numeric(db["amount"])
        summed = db["summed"] == common.YES
        self.assertAlmostEqual(float(amt[summed & (db["type"] == common.TYPE_EXPENSE)].sum()),
                               E["classify"]["expense_in_window_if_all_unknowns_are_expenses"], 2)
        self.assertAlmostEqual(float(amt[summed & (db["type"] == common.TYPE_INCOME)].sum()), E["classify"]["income_in_window"], 2)
        self.assertEqual(int((db["type"] == common.TYPE_CARD_DEBIT).sum()), E["classify"]["card_debit_rows_non_summed"])
        self.assertEqual(int((db["type"] == common.TYPE_DUPLICATE).sum()), E["classify"]["duplicate_rows"])
        self.assertEqual(int((db["type"] == common.TYPE_INTERNAL).sum()), E["classify"]["internal_rows"])
        self.assertTrue((db.loc[db["type"] == common.TYPE_CARD_DEBIT, "summed"] == common.NO).all())

    def test_03_workbook_sheet_set(self):
        wb = openpyxl.load_workbook(os.path.join(self.project, "outputs", "תזרים.xlsx"), read_only=True)
        try:
            ordered = list(wb.sheetnames)
            hdr_exp = [c.value for c in next(wb[common.SHEET_EXP].iter_rows(min_row=2, max_row=2))]
        finally:
            wb.close()
        names = set(ordered)
        for s in common.SHEETS_BY_LEVEL["standard"]:
            self.assertIn(s, names)
        self.assertNotIn(common.SHEET_PROPOSED, names)
        self.assertNotIn(make_fixtures.TEMPLATE_SHEET, names)       # v0.1.1: template tab is not copied
        self.assertEqual(ordered[0], common.SHEET_EXP)              # הוצאות is the first tab
        self.assertEqual(ordered.index(common.SHEET_TRANSFERS), ordered.index(common.SHEET_ANALYSIS) + 1)
        self.assertFalse(any("תקציב" in str(h) for h in hdr_exp))   # v0.1.1: no budget columns
        self.assertEqual(self.steps(self.res)["verify"]["result"]["summary"]["fail"], 0)

    def test_04_reconciliation_residual_zero(self):
        with open(os.path.join(self.project, "work", "summary.json"), encoding="utf-8") as fh:
            S = json.load(fh)
        rc = S["recon"]
        E = self.expected
        self.assertEqual(rc["bank_change_source"], "balances")
        self.assertAlmostEqual(rc["per_account"]["bank_a"]["balance_before"], E["bank_a"]["opening_balance_implied"], 2)
        self.assertAlmostEqual(rc["per_account"]["bank_a"]["balance_end"], E["bank_a"]["closing_balance"], 2)
        self.assertAlmostEqual(rc["per_account"]["bank_b"]["balance_before"], E["bank_b"]["opening_balance"], 2)
        self.assertAlmostEqual(rc["per_account"]["bank_b"]["balance_end"], E["bank_b"]["closing_balance"], 2)
        self.assertAlmostEqual(rc["bank_change"], E["bank_a"]["net"] - E["bank_b"]["outflows"], 2)
        self.assertAlmostEqual(rc["residual"], 0.0, delta=0.01)
        self.assertTrue(rc["residual_ok"])
        self.assertEqual(rc["per_account"]["bank_a"]["chain_breaks"], 0)
        self.assertEqual(rc["per_account"]["bank_b"]["chain_breaks"], 0)

    def test_05_breakdowns_match_expected(self):
        with open(os.path.join(self.project, "work", "summary.json"), encoding="utf-8") as fh:
            S = json.load(fh)
        E = self.expected
        xb = E["classify"]["expense_breakdown"]
        cards = {k: v["total"] for k, v in S["expenses_by_card"].items()}
        cfg = common.load_config(self.cfg_path)
        by_id = {c["id"]: c["label"] for c in cfg["cards"] + cfg["benefit_programs"]}
        # card 1234: super + clothes + FX row (its BIT-funded row is a duplicate, not an expense)
        self.assertAlmostEqual(cards[by_id["card_1234"]], E["card_cycle"]["total_amount_ils"] - xb["הוראת קבע לדוגמה"] - xb["BIT app row"], 2)
        self.assertAlmostEqual(cards[by_id["card_5678"]], xb["הוראת קבע לדוגמה"], 2)
        self.assertAlmostEqual(cards[by_id["card_9012"]], E["card_blocks"]["block_total"], 2)
        self.assertAlmostEqual(cards[by_id["club"]], xb["club benefit"] + xb["club purchase"] + xb["הנחת מועדון"], 2)
        people = {common.owner_label(cfg, p): p for p in ("p1", "p2", common.SHARED)}
        per = {people[k]: v["total"] for k, v in S["expenses_by_person"].items()}
        self.assertAlmostEqual(per["p1"], cards[by_id["card_1234"]] + xb["BIT app row"], 2)
        self.assertAlmostEqual(per["p2"], cards[by_id["card_5678"]] + cards[by_id["card_9012"]] + cards[by_id["club"]], 2)
        self.assertAlmostEqual(per[common.SHARED], xb["שכר דירה"] + xb["עמלה"] + xb["העברה מהחשבון"], 2)
        inc = {people[k]: v["total"] for k, v in S["income_by_person"].items()}
        self.assertAlmostEqual(inc[common.SHARED], E["bank_a"]["inflows"], 2)
        self.assertAlmostEqual(inc["p1"], -E["p2p"]["incoming"][0]["amount"], 2)
        self.assertAlmostEqual(sum(per.values()), S["total_expenses"], 2)
        self.assertTrue(all(v["n_rows"] > 0 for v in S["expenses_by_card"].values()))
        self.assertAlmostEqual(S["card_debits"]["table"][0]["diff"], 0.0, 2)

    def test_06_dashboard_and_report(self):
        st = self.steps(self.res)
        db = read_db(self.project)
        self.assertEqual(st["build_dashboard"]["result"]["rows"], len(db))
        self.assertTrue(os.path.isfile(os.path.join(self.project, "outputs", "dashboard.html")))
        self.assertGreaterEqual(st["make_report_html"]["result"]["sections"], 6)
        self.assertAlmostEqual(st["make_report_html"]["result"]["residual"], 0.0, delta=0.01)
        with open(os.path.join(self.project, "work", "report.html"), encoding="utf-8") as fh:
            html = fh.read()
        self.assertIn("<h2", html)
        self.assertGreaterEqual(len(st["make_figures"]["result"]["figures"]), 3)

    # ------------------------------------------------------------------ mode B
    def test_07_mode_b_relabel(self):
        db = read_db(self.project)
        row = db[db["original_name"] == "חנות בגדים לדוגמה"].iloc[0]
        rid = row["id"]
        self.assertEqual(row["cat_tz"], common.UNKNOWN_CAT)
        xlsx = os.path.join(self.project, "outputs", "תזרים.xlsx")
        wb = openpyxl.load_workbook(xlsx)
        ws = wb[CYCLE_SHEET]
        hdr_row, hdr = None, None
        for r in ws.iter_rows(min_row=1, max_row=10):
            if r[0].value == common.DB_HEADERS_HE["id"]:
                hdr_row, hdr = r[0].row, [c.value for c in r]
                break
        self.assertIsNotNone(hdr_row)
        c_text, c_cat = hdr.index(common.USER_TEXT_COL) + 1, hdr.index(common.USER_CAT_COL) + 1
        target = None
        for r in range(hdr_row + 1, ws.max_row + 1):
            if ws.cell(r, 1).value == rid:
                target = r
                break
        self.assertIsNotNone(target, "row %s not offered in %s" % (rid, CYCLE_SHEET))
        user_text, user_cat = "מעיל חורף לדוגמה", "ביגוד והנעלה"
        ws.cell(target, c_text, user_text)
        ws.cell(target, c_cat, user_cat)
        wb.save(xlsx)
        wb.close()
        rc, res, err = run("apply_user_labels.py", self.cfg_path, "--level", "standard")
        self.assertEqual(rc, 0, err[-2000:])
        self.assertTrue(res["ok"], res)
        self.assertEqual(res["labels"]["read"], 1)
        for s in res["steps"]:
            self.assertEqual(s["status"], "ok", "%s: %s" % (s["name"], s["result"]))
        with open(os.path.join(self.project, "rules", "user_labels.csv"), encoding="utf-8-sig") as fh:
            labels = list(csv.DictReader(fh))
        self.assertEqual([(l["id"], l["user_text"], l["user_cat"]) for l in labels], [(rid, user_text, user_cat)])
        db2 = read_db(self.project)
        row2 = db2[db2["id"] == rid].iloc[0]
        self.assertEqual(row2["cat_tz"], user_cat)
        self.assertEqual(row2["group_tz"], "ביגוד וטיפוח")          # group taken from the scheme CSV
        self.assertIn(user_text, row2["name_clean"])
        self.assertIn(user_text, row2["rule_note"])
        # the rebuilt workbook carries the text (database sheet) and the category in the expense sheet
        wb = openpyxl.load_workbook(xlsx, read_only=True)
        try:
            found_text = any(user_text in str(c) for r in wb[common.SHEET_DB].iter_rows(values_only=True) for c in r if c)
            found_cat = any(c == user_cat for r in wb[common.SHEET_EXP].iter_rows(values_only=True) for c in r)
        finally:
            wb.close()
        self.assertTrue(found_text)
        self.assertTrue(found_cat)
        # the summary was rebuilt and the identity still closes (a relabel changes no cash)
        with open(os.path.join(self.project, "work", "summary.json"), encoding="utf-8") as fh:
            S = json.load(fh)
        self.assertAlmostEqual(S["recon"]["residual"], 0.0, delta=0.01)
        self.assertIn("ביגוד וטיפוח / ביגוד והנעלה", S["expenses_by_cat"])

    # ------------------------------------------------------------------ mode C
    def test_08_mode_c_add_cycle_file(self):
        before = read_db(self.project)
        ids_before = set(before["id"])
        cycle_dir = os.path.join(self.project, "inputs", "cards", "card_cycle")
        # third cycle: card_0425 = charged 10/03/2025 (after the window: rows land with in_window = לא)
        third = {
            "row7": "2 עסקאות לחיוב בחודש מרץ בכל הכרטיסים",
            "row8": 'עסקאות באשראי - סה"כ חיוב:  + ₪330',
            "rows": [
                ["ויזה 1234", "סופר לדוגמה", "02/03/2025", 210, "מנפיק לדוגמה", "ישראל", " רגילה", "10/03/2025", 210, "--"],
                ["ויזה 5678", "מסעדה לדוגמה", "05/03/2025", 120, "מנפיק לדוגמה", "ישראל", " רגילה", "10/03/2025", 120, "--"],
            ],
        }
        make_fixtures.write_card_cycle(os.path.join(cycle_dir, "card_0425.xlsx"), third)
        rc, res, err = run("run_pipeline.py", self.cfg_path, "--mode", "C", "--level", "standard")
        self.assertEqual(rc, 0, err[-2000:])
        for s in res["steps"]:
            self.assertEqual(s["status"], "ok", "%s: %s" % (s["name"], s["result"]))
        after = read_db(self.project)
        self.assertTrue(ids_before <= set(after["id"]))                          # stable ids (F: id drift)
        new = after[~after["id"].isin(ids_before)]
        self.assertEqual(len(new), 2)
        self.assertEqual(set(new["source_file"]), {"card_0425.xlsx"})
        self.assertTrue((new["in_window"] == common.NO).all())
        self.assertTrue((new["summed"] == common.NO).all())
        self.assertEqual(set(new["charge_date"]), {"2025-03-10"})
        # untouched rows keep their classification, including the mode-B label
        merged = before.merge(after, on="id", suffixes=("_a", "_b"))
        self.assertEqual(len(merged), len(before))
        for col in ("cat_tz", "type", "amount", "name_clean"):
            self.assertTrue((merged[col + "_a"] == merged[col + "_b"]).all(), col)
        with open(os.path.join(self.project, "work", "summary.json"), encoding="utf-8") as fh:
            S = json.load(fh)
        self.assertAlmostEqual(S["recon"]["residual"], 0.0, delta=0.01)

    # ------------------------------------------------------------------ levels
    def test_09_levels_overview_and_deep(self):
        for level, absent in (("overview", [common.SHEET_MANUAL, common.SHEET_PROPOSED]), ("deep", [])):
            rc, res, err = run("run_pipeline.py", self.cfg_path, "--mode", "A", "--level", level)
            self.assertEqual(rc, 0, "%s: %s" % (level, err[-2000:]))
            for s in res["steps"]:
                self.assertIn(s["status"], ("ok", "skipped"), "%s/%s: %s" % (level, s["name"], s["result"]))
            wb = openpyxl.load_workbook(os.path.join(self.project, "outputs", "תזרים.xlsx"), read_only=True)
            try:
                names = set(wb.sheetnames)
            finally:
                wb.close()
            for s in common.SHEETS_BY_LEVEL[level]:
                self.assertIn(s, names, level)
            for s in absent:
                self.assertNotIn(s, names, level)


if __name__ == "__main__":
    unittest.main()
