# -*- coding: utf-8 -*-
"""Tests for make_fixtures.py: the synthetic project generates, every listed file exists and
opens with the right library, the layouts match formats.py, the config loads without warnings,
and the hand-derived expected values are internally consistent with the written files.
Run: python3 -m unittest discover -s scripts/tests -v
"""
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

import common  # noqa: E402
import formats  # noqa: E402
import make_fixtures  # noqa: E402

try:
    import openpyxl
    from docx import Document
except ImportError:  # pragma: no cover
    openpyxl = Document = None


@unittest.skipIf(openpyxl is None, "openpyxl / python-docx not installed")
class TestFixtures(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.out = tempfile.mkdtemp(prefix="tz_fix_")
        cls.root, cls.files = make_fixtures.generate(cls.out)
        cls.expected = make_fixtures.EXPECTED

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.out, ignore_errors=True)

    def test_every_listed_file_exists(self):
        self.assertGreaterEqual(len(self.files), 12)
        for f in self.files:
            self.assertTrue(os.path.isfile(f), f)
            self.assertGreater(os.path.getsize(f), 0, f)

    def test_config_loads_without_warnings(self):
        cfg = common.load_config(os.path.join(self.out, "tazrim.config.json"))
        self.assertEqual(cfg.warnings, [], cfg.warnings)
        self.assertEqual(common.months(cfg), self.expected["window"]["months"])
        self.assertEqual(len(cfg["accounts"]), 2)
        self.assertEqual(len(cfg["cards"]), 3)
        self.assertEqual(len(common.resolve_files(cfg, cfg["cards"][0]["files"])), 2)

    def test_bank_xlsx_a_layout_and_balance_chain(self):
        spec = formats.BANKS["bank_xlsx_a"]
        wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/bank/bank_a/עובר ושב.xlsx"), data_only=True)
        ws = wb[spec["sheet"]]
        header = [ws.cell(spec["header_row"], c).value for c in range(1, len(spec["header"]) + 1)]
        self.assertEqual(header, spec["header"])
        rows = []
        for r in range(spec["data_from_row"], ws.max_row + 1):
            v = [ws.cell(r, c).value for c in range(1, 9)]
            if v[0] is None:
                break
            rows.append(v)
        self.assertEqual(len(rows), self.expected["bank_a"]["rows"])
        # F1 chain on the file's own numbers (newest first)
        for a, b in zip(rows, rows[1:]):
            self.assertAlmostEqual(a[4] - a[3], b[4], places=2)
        self.assertAlmostEqual(rows[0][4], self.expected["bank_a"]["closing_balance"])
        self.assertAlmostEqual(rows[-1][4] - rows[-1][3], self.expected["bank_a"]["opening_balance_implied"])
        inflow = sum(v[3] for v in rows if v[3] > 0)
        outflow = -sum(v[3] for v in rows if v[3] < 0)
        self.assertAlmostEqual(inflow, self.expected["bank_a"]["inflows"])
        self.assertAlmostEqual(outflow, self.expected["bank_a"]["outflows"])

    def test_bank_xls_b_layout_and_reconstructed_balance(self):
        spec = formats.BANKS["bank_xls_b"]
        wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/bank/bank_b/עובר ושב.xlsx"), data_only=True)
        ws = wb[spec["sheet"]]
        header = [ws.cell(spec["header_row"], c).value or "" for c in range(1, len(spec["header"]) + 1)]
        self.assertEqual(header, spec["header"])
        self.assertEqual(ws.cell(spec["opening_row"], 6).value, spec["opening_marker"])
        opening = ws.cell(spec["opening_row"], 2).value
        self.assertAlmostEqual(opening, self.expected["bank_b"]["opening_balance"])
        running, checkpoints, n = opening, 0, 0
        for r in range(spec["data_from_row"], ws.max_row + 1):
            v = [ws.cell(r, c).value for c in range(1, 10)]
            if v[8] is None:
                break
            n += 1
            running += (v[3] or 0) - (v[4] or 0)
            if v[1] is not None:
                checkpoints += 1
                self.assertAlmostEqual(v[1], running)
        self.assertEqual(n, self.expected["bank_b"]["rows"])
        self.assertEqual(checkpoints, self.expected["bank_b"]["balance_checkpoints"])
        self.assertAlmostEqual(running, self.expected["bank_b"]["closing_balance"])

    def test_cal_layout(self):
        spec = formats.ISSUERS["card_cycle_xlsx"]
        for name, exp in self.expected["card_cycle"]["files"].items():
            wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/cards/card_cycle", name), data_only=True)
            ws = wb[spec["sheet"]]
            header = [ws.cell(spec["header_row"], c).value for c in range(1, len(spec["header"]) + 1)]
            self.assertEqual(header, spec["header"])
            self.assertRegex(ws.cell(7, 1).value, spec["cycle_count_pattern"])
            self.assertIn(spec["cycle_total_marker"], ws.cell(8, 1).value)
            n = 0
            for r in range(spec["data_from_row"], ws.max_row + 1):
                a = ws.cell(r, 1).value
                if a is None or str(a).startswith(spec["stop_marker"]):
                    break
                n += 1
            self.assertEqual(n, exp["parsed"], name)
        # the column-shift row: charge date sits in col 7 (no פירוט cell)
        wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/cards/card_cycle/card_0225.xlsx"), data_only=True)
        ws = wb[spec["sheet"]]
        shift = self.expected["card_cycle"]["column_shift_row"]
        row = [ws.cell(11, c).value for c in range(1, 11)]
        self.assertEqual(row[1], shift["merchant"])
        self.assertEqual(row[6], "10/01/2025")
        self.assertEqual(row[7], shift["charge_amount"])
        # the uncharged standing order carries '--' as charge date
        row = [ws.cell(13, c).value for c in range(1, 11)]
        self.assertEqual(row[7], spec["uncharged_marker"])

    def test_card_blocks_new_layout(self):
        spec = formats.ISSUERS["card_blocks_xls"]["detail_xlsx"]
        wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/cards/card_9012/9012_0225.xlsx"), data_only=True)
        ws = wb[spec["sheet"]]
        cells = [ws.cell(r, c).value for r in range(1, 15) for c in range(1, 9)]
        texts = [str(v) for v in cells if v is not None]
        self.assertTrue(any(t.startswith(spec["charge_cell_prefix"]) for t in texts))
        hrow = next(r for r in range(1, 15) if ws.cell(r, 1).value == spec["header_marker"])
        header = [ws.cell(hrow, c).value for c in range(1, len(spec["header"]) + 1)]
        self.assertEqual(header, spec["header"])
        total = 0.0
        n = 0
        for r in range(hrow + 1, ws.max_row + 1):
            a = ws.cell(r, 1).value
            if str(a).startswith(spec["total_marker"]):
                self.assertAlmostEqual(ws.cell(r, 2).value, self.expected["card_blocks"]["block_total"])
                break
            n += 1
            total += ws.cell(r, 5).value
        self.assertEqual(n, self.expected["card_blocks"]["rows"])
        self.assertAlmostEqual(total, self.expected["card_blocks"]["block_total"])

    def test_docx_record_shapes(self):
        spec = formats.ISSUERS["club_docx"]
        ps = [p.text.strip() for p in Document(os.path.join(self.out, "inputs/benefits/club/פירוט.docx")).paragraphs]
        bh = spec["benefits_segment"]["header_sequence"]
        i = next(k for k in range(len(ps)) if ps[k:k + len(bh)] == bh) + len(bh)
        self.assertIsNone(spec["benefits_segment"]["record_start_marker"])   # no club name baked into the spec
        self.assertTrue(ps[i] and not ps[i][0].isdigit())                    # paragraph 0 = club label
        self.assertRegex(ps[i + 1], r"^\d{2}\.\d{2}\.\d{2}$")                # paragraph 1 = purchase date
        rec = ps[i:i + spec["benefits_segment"]["record_paragraphs"]]
        self.assertEqual(len(rec), 9)
        self.assertIn("30", rec[5])
        ch = spec["card_segment"]["header_sequence"]
        j = next(k for k in range(len(ps)) if ps[k:k + len(ch)] == ch) + len(ch)
        load = ps[j:j + 7]
        purchase = ps[j + 7:j + 14]
        self.assertEqual(load[5], "טעינה")
        self.assertIn("200", load[6])
        self.assertEqual(purchase[5], "חיוב")
        self.assertIn("-70", purchase[6])
        e = self.expected["club"]
        self.assertAlmostEqual(e["loads_total"] + e["benefits_total"] - e["bank_debit"], e["discount"])

    def test_p2p_csv_schema(self):
        rows = common.read_csv(os.path.join(self.out, "inputs/p2p/bit_p1.csv"))
        self.assertEqual(list(rows[0].keys()), common.NORMALIZED_COLUMNS)
        self.assertEqual(len(rows), 2)
        self.assertEqual(float(rows[0]["amount_ils"]), 120.0)
        self.assertEqual(float(rows[1]["amount_ils"]), -60.0)

    def test_boi_cache_shape(self):
        p = os.path.join(self.out, "work/boi_rates/RER_USD_ILS.csv")
        rows = common.read_csv(p)
        self.assertEqual(list(rows[0].keys()), formats.BOI["csv_columns"])
        rate = {r["TIME_PERIOD"]: float(r["OBS_VALUE"]) for r in rows}
        self.assertEqual(rate["2025-01-12"], self.expected["card_cycle"]["fx_row"]["boi_rate"])
        self.assertAlmostEqual(self.expected["card_cycle"]["fx_row"]["charge_usd"] * rate["2025-01-12"],
                               self.expected["card_cycle"]["fx_row"]["amount_ils"])
        self.assertTrue(os.path.isfile(os.path.join(self.out, "work/boi_rates/RER_EUR_ILS.csv")))

    def test_template_and_rules(self):
        wb = openpyxl.load_workbook(os.path.join(self.out, "inputs/template/תזרים.xlsx"))
        ws = wb[self.expected["template"]["sheet"]]
        cats = [ws.cell(r, 2).value for r in range(2, 8)]
        groups = {ws.cell(r, 1).value for r in range(2, 8)}
        self.assertEqual(len(cats), self.expected["template"]["categories"])
        self.assertEqual(len(groups), self.expected["template"]["groups"])
        rules = common.read_csv(os.path.join(self.out, "rules/merchant_rules.csv"), skip_comments=True)
        self.assertEqual(len(rules), self.expected["rules"]["patterns"])
        self.assertEqual(list(rules[0].keys()), make_fixtures.RULES_COLUMNS)
        for r in rules:
            self.assertNotIn(",", r["cat"])  # FM5

    def test_expected_json_written(self):
        with open(os.path.join(self.out, "expected.json"), encoding="utf-8") as fh:
            exp = json.load(fh)
        self.assertEqual(exp["card_cycle"]["total_amount_ils"], 1175.0)
        self.assertEqual(exp["classify"]["income_in_window"], 8060.0)

    def test_cli_emits_single_json(self):
        out = tempfile.mkdtemp(prefix="tz_cli_")
        try:
            proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, "make_fixtures.py"), "--out", out],
                                  stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=120)
            self.assertEqual(proc.returncode, 0, proc.stderr.decode("utf-8", "replace"))
            lines = [l for l in proc.stdout.decode("utf-8").splitlines() if l.strip()]
            self.assertEqual(len(lines), 1)
            obj = json.loads(lines[0])
            self.assertTrue(obj["ok"])
            self.assertIn("expected", obj)
            self.assertTrue(all(os.path.isfile(os.path.join(out, f)) for f in obj["files"]))
        finally:
            shutil.rmtree(out, ignore_errors=True)


if __name__ == "__main__":
    unittest.main()
