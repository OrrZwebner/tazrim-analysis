# -*- coding: utf-8 -*-
"""Tests for make_figures.py, make_report_html.py, build_dashboard.py and render_pdf.sh on the
synthetic project of synth_db_outputs.py (summary produced by make_summary.build_summary).
Run: python3 -m unittest discover -s scripts/tests -v
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, HERE)

import common  # noqa: E402
import synth_db_outputs as synth  # noqa: E402
import make_report_html  # noqa: E402
import build_dashboard  # noqa: E402

try:
    import pandas  # noqa: F401
    import make_summary
except ImportError:  # pragma: no cover
    make_summary = None
try:
    import matplotlib  # noqa: F401
    import make_figures
except ImportError:  # pragma: no cover
    make_figures = None
try:
    import openpyxl
except ImportError:  # pragma: no cover
    openpyxl = None

CDNJS = "https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"


def _read(path):
    with open(path, encoding="utf-8") as fh:
        return fh.read()


def _script_of(html):
    return html.split("<script>")[-1].split("</script>")[0]


def _top_level_consts(js):
    names = []
    for m in re.findall(r"^const\s+(.+?);\s*$", js, re.M):
        for part in re.split(r",\s*(?=[A-Za-z_$][\w$]*\s*=)", m):
            names.append(part.split("=")[0].strip())
    return names


@unittest.skipIf(make_summary is None, "pandas not installed")
class TestOutputs(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix="tz_out_")
        cls.cfg_path, cls.E = synth.make_project(cls.tmp)
        cls.cfg = common.load_config(cls.cfg_path)
        cls.S = make_summary.build_summary(cls.cfg)
        common.write_json(common.project_path(cls.cfg, "work.summary"), cls.S)
        cls.fig_dir = common.project_path(cls.cfg, "work.figures")
        cls.figs = None
        if make_figures is not None:
            cls.figs, cls.font = make_figures.make_figures(cls.cfg)

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    # ------------------------------------------------------------------ figures
    @unittest.skipIf(make_figures is None, "matplotlib not installed")
    def test_figures_count_per_level(self):
        self.assertEqual(len(self.figs), 14)
        for p in self.figs:
            self.assertTrue(os.path.isfile(p) and os.path.getsize(p) > 1000, p)
        for level, n in (("overview", 5), ("standard", 11), ("deep", 14)):
            out = os.path.join(self.tmp, "figs_" + level)
            written, _ = make_figures.make_figures(self.cfg, out_dir=out, level=level)
            self.assertEqual(len(written), n, level)
            self.assertEqual(len(written), len(make_figures.FIGURES[level]))

    @unittest.skipIf(make_figures is None, "matplotlib not installed")
    def test_bidi_fallback_keeps_numbers_logical(self):
        heb = make_figures.heb
        self.assertEqual(heb("ABC only"), "ABC only")
        out = heb("ממוצע 1,370 ₪")
        self.assertIn("1,370", out)          # digits stay in logical order
        self.assertTrue(out.endswith("עצוממ") or make_figures.BIDI == "python-bidi")

    # ------------------------------------------------------------------ report
    def test_report_deep_structure(self):
        res = make_report_html.make_report(self.cfg)
        html = _read(res["html"])
        self.assertIn('dir="rtl"', html)
        self.assertIn('lang="he"', html)
        self.assertIn("@page { size: A4", html)
        h2 = re.findall(r"<h2>(\d)\. ([^<]+)</h2>", html)
        self.assertEqual([int(n) for n, _ in h2], list(range(1, 9)))
        self.assertEqual([t for _, t in h2], make_report_html.SECTION_TITLES)
        self.assertEqual(res["sections"], 8)
        self.assertTrue(res["findings_present"])
        self.assertIn("ממצא לדוגמה 1", html)
        self.assertIn("פעולה לדוגמה", html)
        self.assertIn("פער נתונים לדוגמה", html)
        self.assertNotIn("/Users/", html)
        self.assertNotIn(self.tmp, html)
        # numbers come from the summary with thousands separators and ₪
        self.assertIn("14,085", html)
        self.assertIn("40,080", html)
        self.assertIn("₪", html)
        # reconciliation table closes to the residual and lists the components
        self.assertIn("שארית", html)
        self.assertIn(make_report_html.RECON_LABELS["fx_expenses"], html)
        self.assertIn(make_report_html.RECON_LABELS["club_balance_change"], html)
        # "ממוצע ללא אפס" column present in the category table
        self.assertIn("ממוצע ללא אפס", html)
        if self.figs:
            self.assertEqual(len(res["figures_embedded"]), 14)
            self.assertEqual(res["figures_missing"], [])
            self.assertGreaterEqual(html.count("data:image/png;base64,"), 14)

    def test_report_overview_has_four_sections(self):
        out = os.path.join(self.tmp, "work", "report_overview.html")
        res = make_report_html.make_report(self.cfg, out_path=out, level="overview")
        html = _read(out)
        h2 = re.findall(r"<h2>(\d)\. ", html)
        self.assertEqual(h2, ["1", "2", "3", "4"])
        self.assertEqual(res["sections"], 4)
        self.assertNotIn("<h2>5.", html)
        if self.figs:
            self.assertEqual(sorted(res["figures_embedded"]), sorted(f for f in res["figures_embedded"]))
            self.assertEqual(len(res["figures_embedded"]), 5)

    def test_report_placeholders_without_findings(self):
        out = os.path.join(self.tmp, "work", "report_nofindings.html")
        res = make_report_html.make_report(self.cfg, findings_path=os.path.join(self.tmp, "nope.json"), out_path=out)
        html = _read(out)
        self.assertFalse(res["findings_present"])
        self.assertIn(make_report_html.FINDINGS_MISSING, html)
        self.assertGreaterEqual(html.count('class="placeholder"'), 3)  # findings, actions, gaps, recommendations

    def test_report_missing_figures_render_placeholders(self):
        out = os.path.join(self.tmp, "work", "report_nofigs.html")
        res = make_report_html.make_report(self.cfg, fig_dir=os.path.join(self.tmp, "no_figs"), out_path=out)
        html = _read(out)
        self.assertEqual(len(res["figures_missing"]), 14)
        self.assertIn("איור", html)
        self.assertNotIn("data:image/png", html)

    def test_report_cli_json(self):
        p = subprocess.run([sys.executable, os.path.join(SCRIPTS, "make_report_html.py"), "--config", self.cfg_path,
                            "--out", os.path.join(self.tmp, "work", "report_cli.html")],
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 0, p.stderr)
        d = json.loads(p.stdout.strip().splitlines()[-1])
        self.assertTrue(d["ok"] and d["residual_ok"])

    def test_report_strict_exits_4_on_residual(self):
        bad = dict(self.S)
        bad["recon"] = dict(self.S["recon"], residual=1234.0, residual_ok=False)
        sp = os.path.join(self.tmp, "work", "summary_bad.json")
        common.write_json(sp, bad)
        p = subprocess.run([sys.executable, os.path.join(SCRIPTS, "make_report_html.py"), "--config", self.cfg_path,
                            "--summary", sp, "--out", os.path.join(self.tmp, "work", "report_bad.html"), "--strict"],
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 4)
        d = json.loads(p.stdout.strip().splitlines()[-1])
        self.assertFalse(d["ok"])
        html = _read(os.path.join(self.tmp, "work", "report_bad.html"))
        self.assertIn("ההתאמה אינה נסגרת", html)  # residual shown, never hidden

    # ------------------------------------------------------------------ dashboard
    def test_dashboard_from_csv(self):
        res = build_dashboard.build_dashboard(self.cfg)
        html = _read(res["dashboard"])
        self.assertEqual(res["source"], "database.csv")
        self.assertEqual(res["rows"], self.E["n_rows_db"])
        self.assertIn('dir="rtl"', html)
        self.assertIn('lang="he"', html)
        for r in synth.db_rows():
            self.assertIn('"id":"%s"' % r["id"], html)
        for tab in ("סקירה", "עסקאות", "העברות, BIT ו-PAYBOX", "הכנסות", "לסיווג", "עזרה"):
            self.assertIn(tab, html)
        self.assertIn(CDNJS, html)
        self.assertEqual(res["chartjs"], "cdn")
        self.assertNotIn("/Users/", html)
        self.assertNotIn(self.tmp, html)
        self.assertIn("משק בית לדוגמה", html)          # title from config
        self.assertIn('"people":["%s"' % synth.P1, html)  # people order from config
        self.assertIn("איך משתמשים (3 שורות)", html)
        js = _script_of(html)
        names = _top_level_consts(js)
        self.assertGreater(len(names), 20)
        self.assertEqual(len(names), len(set(names)), "duplicate top-level const: %s" % [n for n in names if names.count(n) > 1])
        self.assertIn("build_dashboard.py", html)         # refresh command in the footer
        self.assertNotIn("<\\/script", js.replace("<\\/", "").replace("</", ""))  # embedded JSON escaped
        node = shutil.which("node")
        if node:
            jsp = os.path.join(self.tmp, "dash.js")
            with open(jsp, "w", encoding="utf-8") as fh:
                fh.write(js)
            p = subprocess.run([node, "--check", jsp], stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
            self.assertEqual(p.returncode, 0, p.stderr)

    def test_dashboard_overview_tabs(self):
        out = os.path.join(self.tmp, "outputs", "dashboard_overview.html")
        res = build_dashboard.build_dashboard(self.cfg, out_path=out, level="overview")
        self.assertEqual(res["tabs"], ["overview", "explore", "income", "help"])

    def test_dashboard_transfers_fallback_direction(self):
        data, _order, n_tr, _src = build_dashboard.load_data(self.cfg)
        by = {d["id"]: d for d in data}
        self.assertEqual(by["T-007"]["tr_section"], 1)
        self.assertEqual(by["T-007"]["tr_dir"], "יוצא")
        self.assertEqual(by["T-045"]["tr_dir"], "נכנס")     # negative internal = money in (F17)
        self.assertEqual(by["T-030"]["tr_section"], 2)       # P2P app row
        self.assertEqual(by["T-014"]["tr_section"], 2)       # card row with the P2P marker
        self.assertEqual(by["T-011"]["tr_section"], 0)
        self.assertGreater(n_tr[1], 0)

    @unittest.skipIf(openpyxl is None, "openpyxl not installed")
    def test_dashboard_from_workbook_with_transfers_sheet(self):
        wb_path = common.project_path(self.cfg, "outputs.workbook")
        os.makedirs(os.path.dirname(wb_path), exist_ok=True)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = common.SHEET_DB
        cols = [c for c in common.DB_COLUMNS if c != "trip"]  # trip left out -> filled from CSV
        ws.append([common.DB_HEADERS_HE[c] for c in cols])
        for r in synth.db_rows():
            ws.append([r[c] for c in cols])
        tr = wb.create_sheet(common.SHEET_TRANSFERS)
        tr.append(["1. העברות בנקאיות"])
        tr.append([common.DB_HEADERS_HE["id"], "כיוון", "מוטב / צד שני", "הערה מצילום BIT / PAYBOX"])
        tr.append(["T-007", "יוצא", "בנק לדוגמה", ""])
        tr.append(["2. תשלומי אפליקציה"])
        tr.append([common.DB_HEADERS_HE["id"], "כיוון", "מוטב / צד שני", "הערה מצילום BIT / PAYBOX"])
        tr.append(["T-030", "יוצא", "מוטב לדוגמה", "הערה לדוגמה"])
        wb.save(wb_path)
        try:
            out = os.path.join(self.tmp, "outputs", "dashboard_wb.html")
            res = build_dashboard.build_dashboard(self.cfg, out_path=out)
            self.assertEqual(res["source"], os.path.basename(wb_path))
            self.assertEqual(res["rows"], self.E["n_rows_db"])
            html = _read(out)
            self.assertIn("מוטב לדוגמה", html)
            self.assertIn('"trip":"נסיעה לדוגמה"', html)   # column missing from the sheet, taken from the CSV
            self.assertEqual(res["transfers"], {1: 1, 2: 1, 3: 0})
        finally:
            os.remove(wb_path)

    def test_dashboard_cli(self):
        p = subprocess.run([sys.executable, os.path.join(SCRIPTS, "build_dashboard.py"), "--config", self.cfg_path,
                            "--out", os.path.join(self.tmp, "outputs", "dashboard_cli.html")],
                           stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 0, p.stderr)
        lines = [l for l in p.stdout.strip().splitlines() if l.strip()]
        self.assertEqual(len(lines), 1)
        self.assertTrue(json.loads(lines[0])["ok"])

    # ------------------------------------------------------------------ render_pdf.sh
    def test_render_pdf_usage_and_missing_input(self):
        sh = os.path.join(SCRIPTS, "render_pdf.sh")
        p = subprocess.run(["bash", sh, "-h"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 0)
        self.assertIn("Usage", p.stdout)
        p = subprocess.run(["bash", sh], stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 2)
        p = subprocess.run(["bash", sh, os.path.join(self.tmp, "missing.html")], stdout=subprocess.PIPE,
                           stderr=subprocess.PIPE, universal_newlines=True)
        self.assertEqual(p.returncode, 1)
        self.assertFalse(json.loads(p.stdout)["ok"])
        p = subprocess.run(["bash", sh, "--probe"], stdout=subprocess.PIPE, stderr=subprocess.PIPE, universal_newlines=True)
        self.assertIn(p.returncode, (0, 3))
        d = json.loads(p.stdout)
        self.assertEqual(d["ok"], p.returncode == 0)
        if p.returncode == 3:
            self.assertIn("engines_tried", d)


class TestPrivacyOfGenerators(unittest.TestCase):
    """The generators must not carry absolute user paths or hard-coded data literals."""

    def test_no_user_paths_in_sources(self):
        for f in ("make_summary.py", "make_figures.py", "make_report_html.py", "build_dashboard.py", "render_pdf.sh",
                  os.path.join("tests", "synth_db_outputs.py")):
            src = _read(os.path.join(SCRIPTS, f))
            self.assertNotIn("/Users/", src, f)


if __name__ == "__main__":
    unittest.main()
