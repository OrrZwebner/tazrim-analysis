# -*- coding: utf-8 -*-
"""Tests for parse_bank.py, parse_card_cycle.py, parse_card_blocks.py, parse_benefits.py on the synthetic
fixture (make_fixtures.py). Every expected value is the hand-typed literal from expected.json.
Run: python3 -m unittest discover -s scripts/tests -v
"""
import contextlib
import io
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)

import common  # noqa: E402
import make_fixtures  # noqa: E402

try:
    import openpyxl  # noqa: F401
    import docx  # noqa: F401
    HAVE_DEPS = True
except ImportError:  # pragma: no cover
    HAVE_DEPS = False


def run_main(module_name, argv):
    """Run <module>.main(argv) in-process, capturing the single JSON object it emits."""
    mod = __import__(module_name)
    buf = io.StringIO()
    err = io.StringIO()
    code = 0
    with contextlib.redirect_stdout(buf), contextlib.redirect_stderr(err):
        try:
            mod.main(argv)
        except SystemExit as e:
            code = e.code or 0
    lines = [l for l in buf.getvalue().splitlines() if l.strip()]
    assert len(lines) == 1, "expected exactly one JSON line on stdout, got %r" % lines
    return code, json.loads(lines[0]), err.getvalue()


def run_cli(script, argv):
    proc = subprocess.run([sys.executable, os.path.join(SCRIPTS, script)] + argv,
                          stdout=subprocess.PIPE, stderr=subprocess.PIPE, timeout=180)
    lines = [l for l in proc.stdout.decode("utf-8").splitlines() if l.strip()]
    return proc.returncode, (json.loads(lines[0]) if len(lines) == 1 else None), proc.stderr.decode("utf-8", "replace")


def read_norm(root, name):
    return common.read_csv(os.path.join(root, "work", "normalized", name + ".csv"))


@unittest.skipUnless(HAVE_DEPS, "openpyxl / python-docx not installed")
class TestParsers(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = tempfile.mkdtemp(prefix="tz_parse_")
        make_fixtures.generate(cls.root)
        cls.cfg_path = os.path.join(cls.root, "tazrim.config.json")
        with open(os.path.join(cls.root, "expected.json"), encoding="utf-8") as fh:
            cls.exp = json.load(fh)
        cls.bank = run_main("parse_bank", ["--config", cls.cfg_path])
        cls.cycle = run_main("parse_card_cycle", ["--config", cls.cfg_path, "--offline"])
        cls.isr = run_main("parse_card_blocks", ["--config", cls.cfg_path])
        cls.bhz = run_main("parse_benefits", ["--config", cls.cfg_path])

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.root, ignore_errors=True)

    # ------------------------------------------------------------------ bank
    def test_bank_xlsx_a_f1(self):
        code, out, _ = self.bank
        self.assertEqual(code, 0)
        self.assertTrue(out["ok"])
        d, e = out["accounts"]["bank_a"], self.exp["bank_a"]
        self.assertEqual(d["rows"], e["rows"])
        self.assertAlmostEqual(d["inflows"], e["inflows"])
        self.assertAlmostEqual(d["outflows"], e["outflows"])
        self.assertAlmostEqual(d["net"], e["net"])
        self.assertEqual(d["balance"]["breaks"], e["balance_breaks"])
        self.assertTrue(d["balance"]["ok"])
        self.assertAlmostEqual(d["balance"]["opening_balance"], e["opening_balance_implied"])
        self.assertAlmostEqual(d["balance"]["closing_balance"], e["closing_balance"])
        rows = read_norm(self.root, "bank_a")
        self.assertEqual(list(rows[0].keys()), common.NORMALIZED_COLUMNS)
        self.assertEqual(len(rows), e["rows"])
        by_name = {r["original_name"]: r for r in rows}
        self.assertAlmostEqual(float(by_name["משכורת לדוגמה"]["amount_ils"]), -8000.0)   # money in = negative
        self.assertAlmostEqual(float(by_name["שכר דירה לדוגמה"]["amount_ils"]), 2000.0)
        debits = {(r["charge_date"], r["original_name"][-4:]): float(r["amount_ils"]) for r in rows if "חיוב לכרטיס" in r["original_name"]}
        for k, v in e["card_debits"].items():
            d_, l4 = k.split("|")
            self.assertAlmostEqual(debits[(d_, l4)], v)

    def test_bank_xls_b_f2(self):
        _, out, _ = self.bank
        d, e = out["accounts"]["bank_b"], self.exp["bank_b"]
        self.assertEqual(d["rows"], e["rows"])
        self.assertAlmostEqual(d["balance"]["opening_balance"], e["opening_balance"])
        self.assertAlmostEqual(d["balance"]["closing_balance"], e["closing_balance"])
        self.assertEqual(d["balance"]["checks"], e["balance_checkpoints"])
        self.assertEqual(d["balance"]["breaks"], e["balance_mismatches"])
        self.assertFalse(d["balance"]["opening_inferred"])
        self.assertAlmostEqual(d["outflows"], e["outflows"])
        rows = read_norm(self.root, "bank_b")
        self.assertAlmostEqual(float(next(r for r in rows if r["original_name"].startswith("9012"))["amount_ils"]), 300.0)

    def test_bank_opening_inferred_when_absent(self):
        """No 'יתרת פתיחה' row and no config value -> opening inferred from the first checkpoint."""
        import parse_bank
        from formats import BANKS
        spec = BANKS["bank_xls_b"]
        path = os.path.join(self.root, "inputs", "bank", "bank_b", "עובר ושב.xlsx")
        wb = openpyxl.load_workbook(path)
        ws = wb[spec["sheet"]]
        ws.delete_rows(spec["opening_row"])
        tmp = os.path.join(self.root, "bank_b_noopen.xlsx")
        wb.save(tmp)
        acc = {"id": "bank_b", "bank": "bank_xls_b", "label": "x", "owner": "p2", "files": tmp, "opening_balance": None}
        rows, rep, an = parse_bank.parse_bank_xls_b(tmp, acc, spec)
        self.assertEqual(len(rows), self.exp["bank_b"]["rows"])
        self.assertTrue(rep["opening_inferred"])
        self.assertAlmostEqual(rep["opening_balance"], self.exp["bank_b"]["opening_balance"])
        self.assertAlmostEqual(rep["closing_balance"], self.exp["bank_b"]["closing_balance"])
        self.assertEqual(rep["breaks"], 0)

    def test_bank_header_mismatch_fails_loudly(self):
        import parse_bank
        from formats import BANKS
        path = os.path.join(self.root, "inputs", "bank", "bank_a", "עובר ושב.xlsx")
        wb = openpyxl.load_workbook(path)
        wb[BANKS["bank_xlsx_a"]["sheet"]].cell(BANKS["bank_xlsx_a"]["header_row"], 3, "עמודה אחרת")
        tmp = os.path.join(self.root, "bad_header.xlsx")
        wb.save(tmp)
        with self.assertRaises(ValueError):
            parse_bank.parse_bank_xlsx_a(tmp, {"id": "x", "label": "x"}, BANKS["bank_xlsx_a"])

    def test_bank_ids_stable_when_file_added(self):
        """Adding a second export to an account leaves the first file's ids unchanged (FM6)."""
        before = {r["id"]: r["original_name"] for r in read_norm(self.root, "bank_a")}
        extra = os.path.join(self.root, "inputs", "bank", "bank_a", "עובר ושב 2.xlsx")
        wb = openpyxl.load_workbook(os.path.join(self.root, "inputs", "bank", "bank_a", "עובר ושב.xlsx"))
        ws = wb.worksheets[0]
        ws.cell(9, 3, "תנועה נוספת לדוגמה")     # a different description -> not byte-identical
        wb.save(extra)
        try:
            code, out, _ = run_main("parse_bank", ["--config", self.cfg_path])
            self.assertEqual(code, 0)
            after = {r["id"]: r["original_name"] for r in read_norm(self.root, "bank_a")}
            self.assertEqual(out["accounts"]["bank_a"]["rows"], 14)
            for k, v in before.items():
                self.assertEqual(after.get(k), v)
        finally:
            os.remove(extra)
            run_main("parse_bank", ["--config", self.cfg_path])

    # ------------------------------------------------------------------ card cycle
    def test_cal_counts_and_carry_over(self):
        code, out, _ = self.cycle
        self.assertEqual(code, 0)
        r, e = out["report"], self.exp["card_cycle"]
        for f, ef in e["files"].items():
            self.assertEqual(r["per_file"][f]["parsed"], ef["parsed"], f)
            self.assertEqual(r["per_file"][f]["kept"], ef["kept"], f)
        self.assertEqual(r["kept"], e["kept_rows"])
        self.assertEqual(r["uncharged_dropped"], e["uncharged_dropped"])
        self.assertEqual(r["uncharged_kept"], e["uncharged_kept"])
        self.assertEqual(r["cross_file_duplicates"], e["cross_file_duplicates"])
        self.assertAlmostEqual(r["total_amount_ils"], e["total_amount_ils"])
        self.assertIn(e["files"]["card_0225.xlsx"]["cycle_date"], [x["cycle"] for x in r["reconcile"]])

    def test_cal_fx_row_boi_rate(self):
        _, out, _ = self.cycle
        e = self.exp["card_cycle"]["fx_row"]
        fx = out["report"]["fx_rows"]
        self.assertEqual(len(fx), 1)
        self.assertEqual(fx[0]["merchant"], e["merchant"])
        self.assertEqual(fx[0]["charge_date"], e["charge_date"])
        self.assertAlmostEqual(fx[0]["charge_amount"], e["charge_usd"])
        self.assertAlmostEqual(fx[0]["rate"]["rate"], e["boi_rate"])
        self.assertAlmostEqual(fx[0]["amount_ils"], e["amount_ils"])     # 100 USD x 3.50 = 350.00
        self.assertEqual(fx[0]["orig_currency"], e["orig_currency"])
        rows = read_norm(self.root, "card_1234")
        row = next(x for x in rows if x["original_name"] == e["merchant"])
        self.assertAlmostEqual(float(row["amount_ils"]), 350.0)
        self.assertIn("המרה משוערת", row["details"])

    def test_cal_column_shift_row(self):
        e = self.exp["card_cycle"]["column_shift_row"]
        rows = read_norm(self.root, "card_1234")
        row = next(x for x in rows if x["original_name"] == e["merchant"])
        self.assertEqual(row["charge_date"], e["charge_date"])
        self.assertAlmostEqual(float(row["amount_ils"]), e["charge_amount"])

    def test_cal_reconciliation_against_bank(self):
        _, out, _ = self.cycle
        got = {(x["cycle"], x["card"]): x for x in out["report"]["reconcile"]}
        for e in self.exp["card_cycle"]["reconcile"]:
            g = got[(e["cycle"], e["card"])]
            self.assertAlmostEqual(g["ils_sum"], e["ils_sum"])
            self.assertAlmostEqual(g["bank"], e["bank"])
            self.assertAlmostEqual(g["diff"], e["diff"])
            if "fx_sum" in e:
                self.assertAlmostEqual(g["fx_sum"], e["fx_sum"])
                self.assertEqual(g["fx_ccy"], e["fx_ccy"])

    def test_cal_standing_order_kept_when_next_file_missing(self):
        """F14 second branch: without the later file the uncharged row is carried to the next cycle."""
        import parse_card_cycle
        from fx_rates import RateBook
        alt = tempfile.mkdtemp(prefix="tz_cal_")
        try:
            make_fixtures.generate(alt)
            os.remove(os.path.join(alt, "inputs", "cards", "card_cycle", "card_0325.xlsx"))
            cfg = common.load_config(os.path.join(alt, "tazrim.config.json"))
            cards = [c for c in cfg["cards"] if c["issuer"] == "card_cycle_xlsx"]
            by_card, rep = parse_card_cycle.parse_card_cycle(cfg, cards, RateBook(cfg, allow_network=False))
            self.assertEqual(rep["uncharged_dropped"], 0)
            self.assertEqual(rep["uncharged_kept"], 1)
            row = by_card["card_5678"][0]
            self.assertEqual(row["charge_date"], "2025-02-10")
            self.assertTrue(any("משוער למחזור החיוב הבא" in d for d in row["details"]))
        finally:
            shutil.rmtree(alt, ignore_errors=True)

    def test_cal_byte_identical_file_skipped(self):
        import parse_card_cycle
        from fx_rates import RateBook
        alt = tempfile.mkdtemp(prefix="tz_cal_")
        try:
            make_fixtures.generate(alt)
            src = os.path.join(alt, "inputs", "cards", "card_cycle", "card_0325.xlsx")
            shutil.copyfile(src, os.path.join(alt, "inputs", "cards", "card_cycle", "card_0425.xlsx"))
            cfg = common.load_config(os.path.join(alt, "tazrim.config.json"))
            cards = [c for c in cfg["cards"] if c["issuer"] == "card_cycle_xlsx"]
            _, rep = parse_card_cycle.parse_card_cycle(cfg, cards, RateBook(cfg, allow_network=False))
            self.assertEqual(len(rep["skipped_files"]), 1)
            self.assertEqual(rep["kept"], self.exp["card_cycle"]["kept_rows"])
        finally:
            shutil.rmtree(alt, ignore_errors=True)

    def test_cal_manual_rate_fallback_and_no_rate(self):
        import parse_card_cycle
        from fx_rates import RateBook
        alt = tempfile.mkdtemp(prefix="tz_cal_")
        try:
            make_fixtures.generate(alt)
            os.remove(os.path.join(alt, "work", "boi_rates", "RER_USD_ILS.csv"))
            cfg_path = os.path.join(alt, "tazrim.config.json")
            with open(cfg_path, encoding="utf-8") as fh:
                cfg = json.load(fh)
            cfg["fx"]["manual_rates"] = "rules/manual_rates.csv"
            with open(cfg_path, "w", encoding="utf-8") as fh:
                json.dump(cfg, fh, ensure_ascii=False)
            with open(os.path.join(alt, "rules", "manual_rates.csv"), "w", encoding="utf-8") as fh:
                fh.write("date,ccy,rate\n2025-01-11,USD,4.00\n")
            cfg = common.load_config(cfg_path)
            cards = [c for c in cfg["cards"] if c["issuer"] == "card_cycle_xlsx"]
            book = RateBook(cfg, allow_network=False)
            by_card, rep = parse_card_cycle.parse_card_cycle(cfg, cards, book)
            fx = rep["fx_rows"][0]
            self.assertEqual(fx["rate"]["source"], "manual")
            self.assertAlmostEqual(fx["amount_ils"], 400.0)        # 100 USD x 4.00 (last rate <= 10 days back)
            row = next(r for r in by_card["card_1234"] if r["fx_ccy"])
            self.assertTrue(any("שער ידני" in d for d in row["details"]))
            # no rate anywhere: parse still succeeds, amount left unconverted, problem reported
            os.remove(os.path.join(alt, "rules", "manual_rates.csv"))
            book = RateBook(cfg, allow_network=False)
            by_card, rep = parse_card_cycle.parse_card_cycle(cfg, cards, book)
            self.assertEqual(rep["fx_rows"][0]["rate"]["source"], "none")
            self.assertAlmostEqual(rep["fx_rows"][0]["amount_ils"], 100.0)
            self.assertTrue(book.problems)
        finally:
            shutil.rmtree(alt, ignore_errors=True)

    def test_cal_fx_override_file(self):
        """rules/fx_overrides.csv forces a row typed ישראל onto the FX account."""
        import parse_card_cycle
        from fx_rates import RateBook
        alt = tempfile.mkdtemp(prefix="tz_cal_")
        try:
            make_fixtures.generate(alt)
            with open(os.path.join(alt, "rules", "fx_overrides.csv"), "w", encoding="utf-8") as fh:
                fh.write("source_file,row_ref,currency,reason\ncard_0225.xlsx,11,USD,reconciled against the header\n")
            cfg = common.load_config(os.path.join(alt, "tazrim.config.json"))
            cards = [c for c in cfg["cards"] if c["issuer"] == "card_cycle_xlsx"]
            _, rep = parse_card_cycle.parse_card_cycle(cfg, cards, RateBook(cfg, allow_network=False))
            self.assertEqual(len(rep["fx_rows"]), 2)
            ov = next(x for x in rep["fx_rows"] if x["row_ref"] == 11)
            self.assertAlmostEqual(ov["amount_ils"], 250.0 * 3.45)   # rate of 2025-01-10
        finally:
            shutil.rmtree(alt, ignore_errors=True)

    # ------------------------------------------------------------------ card blocks
    def test_card_blocks_new_layout(self):
        code, out, _ = self.isr
        self.assertEqual(code, 0)
        r, e = out["report"], self.exp["card_blocks"]
        self.assertEqual(r["rows"], e["rows"])
        self.assertTrue(r["self_check_ok"])
        blk = r["blocks"][0]
        self.assertEqual(blk["last4"], e["card"])
        self.assertEqual(blk["charge_date"], e["charge_date"])
        self.assertAlmostEqual(blk["total"], e["block_total"])
        rec = r["reconcile"][0]
        self.assertAlmostEqual(rec["statement"], e["block_total"])
        self.assertAlmostEqual(rec["bank"], e["bank_debit"])
        self.assertAlmostEqual(rec["diff"], e["diff"])
        self.assertEqual(r["missing_cycles"], [])
        rows = read_norm(self.root, "card_9012")
        self.assertEqual(len(rows), e["rows"])
        self.assertTrue(all(x["charge_date"] == e["charge_date"] for x in rows))
        self.assertAlmostEqual(sum(float(x["amount_ils"]) for x in rows), e["block_total"])

    def test_card_blocks_legacy_blocks_and_credit_sign(self):
        """Legacy 3-block layout (written here as xlsx with the same cells): block totals
        self-check, FX block currency labels, F24 credit sign."""
        import datetime as dt
        import parse_card_blocks
        from formats import ISSUERS
        leg = ISSUERS["card_blocks_xls"]["blocks_xls"]
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = leg["sheet"]
        rows = [
            [None, "כרטיס:9012 - כרטיס אשראי חודש החיוב: 02/02/2025"],
            [None, "עסקאות בשקלים חיוב בתאריך 02/02/2025"],
            [None] + leg["header_ils"],
            [None, dt.datetime(2025, 1, 5), "חנות לדוגמה", 150, 150, ""],
            [None, dt.datetime(2025, 1, 6), "זיכוי לדוגמה", 50, -50, "זיכוי"],
            [None, 'סה"כ', None, None, 100],
            [None, 'עסקאות במט"ח חיוב בתאריך 02/02/2025'],
            [None] + leg["header_fx"],
            [None, dt.datetime(2025, 1, 7), "EXAMPLE ABROAD", 20, 'דולר ארה"ב', 72, 'ש"ח', ""],
            [None, 'סה"כ', None, None, None, 72],
        ]
        for r, vals in enumerate(rows, 1):
            for c, v in enumerate(vals, 1):
                if v is not None:
                    ws.cell(r, c, v)
        path = os.path.join(self.root, "legacy_9012.xlsx")
        wb.save(path)
        grid, datemode = parse_card_blocks.load_grid(path, leg["sheet"], "openpyxl")
        self.assertTrue(parse_card_blocks.is_legacy_layout(grid))
        recs, blocks = parse_card_blocks.parse_legacy(grid, datemode, "legacy_9012.xlsx", {"9012": {}})
        self.assertEqual(len(recs), 3)
        self.assertEqual(len(blocks), 2)
        self.assertTrue(all(b["ok"] for b in blocks))
        credit = next(r for r in recs if r["original_name"] == "זיכוי לדוגמה")
        self.assertAlmostEqual(credit["amount_ils"], -50.0)
        self.assertAlmostEqual(credit["orig_amount"], -50.0)            # F24
        fx = next(r for r in recs if r["kind"] == "FX")
        self.assertEqual(fx["orig_currency"], "USD")
        self.assertAlmostEqual(fx["amount_ils"], 72.0)
        self.assertAlmostEqual(fx["orig_amount"], 20.0)

    def test_card_blocks_missing_cycle_reported_or_reconstructed(self):
        import parse_card_blocks
        alt = tempfile.mkdtemp(prefix="tz_isr_")
        try:
            make_fixtures.generate(alt)
            cfg_path = os.path.join(alt, "tazrim.config.json")
            run_main("parse_bank", ["--config", cfg_path])
            # add a bank debit for the card on a date without a statement file
            p = os.path.join(alt, "work", "normalized", "bank_b.csv")
            rows = common.read_csv(p)
            rows.append(dict(rows[0], id="bank_b-test01", charge_date="2025-01-02", txn_date="2025-01-02", amount_ils="123.00"))
            common.write_csv(p, rows, common.NORMALIZED_COLUMNS)
            cfg = common.load_config(cfg_path)
            cards = [c for c in cfg["cards"] if c["issuer"] == "card_blocks_xls"]
            recs, rep = parse_card_blocks.parse_card_blocks(cfg, cards, set(), False)
            self.assertEqual(len(rep["missing_cycles"]), 1)
            self.assertEqual(rep["synthetic_rows"], [])
            self.assertEqual(len(recs), self.exp["card_blocks"]["rows"])
            recs, rep = parse_card_blocks.parse_card_blocks(cfg, cards, set(), True)
            self.assertEqual(len(rep["synthetic_rows"]), 1)
            syn = next(r for r in recs if r["source_file"] == "SYNTHETIC")
            self.assertAlmostEqual(syn["amount_ils"], 123.0)
            self.assertIn("SYNTHETIC", syn["details"])
            # --ignore skips the file entirely
            recs, rep = parse_card_blocks.parse_card_blocks(cfg, cards, {"9012_0225.xlsx"}, False)
            self.assertEqual(len(recs), 0)
            self.assertEqual(len(rep["skipped"]), 1)
        finally:
            shutil.rmtree(alt, ignore_errors=True)

    # ------------------------------------------------------------------ benefits club
    def test_benefits_records_and_discount(self):
        code, out, _ = self.bhz
        self.assertEqual(code, 0)
        p, e = out["programs"]["club"], self.exp["club"]
        self.assertEqual(p["counts"]["benefits"], e["benefits"])
        self.assertEqual(p["counts"]["loads"], e["loads"])
        self.assertEqual(p["counts"]["purchases"], e["purchases"])
        self.assertAlmostEqual(p["totals"]["benefit"], e["benefits_total"])
        self.assertAlmostEqual(p["totals"]["load"], e["loads_total"])
        self.assertAlmostEqual(p["totals"]["purchase"], e["purchases_total"])
        rec = next(x for x in p["reconcile"] if x["debit_month"] == "2025-02")
        self.assertEqual(rec["bank_debit_date"], e["bank_debit_date"])
        self.assertAlmostEqual(rec["bank"], e["bank_debit"])
        self.assertAlmostEqual(rec["face_value"], e["loads_plus_benefits"])
        self.assertAlmostEqual(rec["discount"], e["discount"])
        self.assertEqual(len(p["discount_rows"]), 1)
        self.assertAlmostEqual(p["discount_rows"][0]["amount_ils"], e["discount_row_amount"])
        rows = read_norm(self.root, "club")
        self.assertEqual(len(rows), 4)
        load = next(r for r in rows if r["details"].startswith("טעינה – לא נסכם"))
        self.assertAlmostEqual(float(load["amount_ils"]), e["loads_total"])
        self.assertTrue(all(r["charge_date"] == e["bank_debit_date"] for r in rows))
        disc = next(r for r in rows if r["source_file"] == "מחושב")
        self.assertAlmostEqual(float(disc["amount_ils"]), -50.0)

    # ------------------------------------------------------------------ CLI contract
    def test_cli_single_json_and_help(self):
        for script, extra in (("parse_bank.py", []), ("parse_card_cycle.py", ["--offline"]),
                              ("parse_card_blocks.py", []), ("parse_benefits.py", [])):
            code, obj, err = run_cli(script, ["--config", self.cfg_path] + extra)
            self.assertEqual(code, 0, (script, err))
            self.assertTrue(obj and obj["ok"], script)
            code, _, _ = run_cli(script, ["--help"])
            self.assertEqual(code, 0, script)
        code, obj, _ = run_cli("parse_bank.py", ["--config", os.path.join(self.root, "missing.json")])
        self.assertEqual(code, 2)
        self.assertFalse(obj["ok"])
        self.assertIn("hint", obj)


if __name__ == "__main__":
    unittest.main()
