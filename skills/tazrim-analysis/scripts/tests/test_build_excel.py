# -*- coding: utf-8 -*-
"""Tests for build_excel.py, add_transfers_sheet.py, apply_user_labels.py and run_pipeline.py on the
synthetic database of synth_db.py (2-month window, hand-computed expectations).
Run: python3 -m unittest discover -s scripts/tests -v
"""
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import unittest

HERE = os.path.dirname(os.path.abspath(__file__))
SCRIPTS = os.path.dirname(HERE)
sys.path.insert(0, SCRIPTS)
sys.path.insert(0, HERE)

import common  # noqa: E402
import synth_db  # noqa: E402
from synth_db import EXPECTED  # noqa: E402

try:
    import openpyxl
    from openpyxl.worksheet.formula import ArrayFormula
except ImportError:  # pragma: no cover
    openpyxl = None

BUILD = os.path.join(SCRIPTS, "build_excel.py")
TRANSFERS = os.path.join(SCRIPTS, "add_transfers_sheet.py")
APPLY = os.path.join(SCRIPTS, "apply_user_labels.py")
PIPE = os.path.join(SCRIPTS, "run_pipeline.py")


def run(script, *args, **kw):
    """Run a script; return (returncode, parsed last-line JSON, stderr)."""
    r = subprocess.run([sys.executable, script] + list(args), capture_output=True, text=True, timeout=kw.get("timeout", 300))
    lines = [ln for ln in r.stdout.strip().splitlines() if ln.strip()]
    payload = json.loads(lines[-1]) if lines else None
    if len(lines) != 1 and kw.get("strict", True):
        raise AssertionError("%s printed %d stdout lines (expected exactly one JSON object): %r" % (script, len(lines), r.stdout[:300]))
    return r.returncode, payload, r.stderr


def header_map(ws, row=2):
    return {ws.cell(row, c).value: c for c in range(1, ws.max_column + 1) if ws.cell(row, c).value}


def formula_text(v):
    return v.text if hasattr(v, "text") else v


@unittest.skipIf(openpyxl is None, "openpyxl not installed")
class TestBuildExcel(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.root = tempfile.mkdtemp(prefix="tz_build_")
        cls.cfg_path, cls.db_path = synth_db.write_project(cls.root)
        cls.results, cls.books = {}, {}
        for lv in common.LEVELS:
            out = os.path.join(cls.root, "outputs", "tz_%s.xlsx" % lv)
            rc, res, err = run(BUILD, "--config", cls.cfg_path, "--level", lv, "--out", out)
            assert rc == 0 and res and res["ok"], (rc, res, err[-800:])
            cls.results[lv] = res
            cls.books[lv] = out

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.root, ignore_errors=True)

    def wb(self, level="standard"):
        return openpyxl.load_workbook(self.books[level])

    # ---- sheet set / months / names -------------------------------------------------------
    def test_sheet_set_per_level(self):
        for lv in common.LEVELS:
            names = set(openpyxl.load_workbook(self.books[lv], read_only=True).sheetnames)
            expected = set(common.SHEETS_BY_LEVEL[lv]) - {common.SHEET_TRANSFERS}   # transfers sheet is a separate step
            self.assertEqual(names, expected, lv)
        self.assertNotIn(common.SHEET_PIVOT, self.results["overview"]["sheets"])
        self.assertIn(common.SHEET_PROPOSED, self.results["deep"]["sheets"])
        self.assertNotIn(common.SHEET_PROPOSED, self.results["standard"]["sheets"])

    def test_sheet_order_and_hidden(self):
        wb = self.wb("deep")
        order = [s for s in common.SHEET_ORDER if s in wb.sheetnames]
        self.assertEqual(wb.sheetnames, order)
        self.assertEqual(wb[common.SHEET_LISTS].sheet_state, "hidden")
        for ws in wb.worksheets:
            self.assertTrue(ws.sheet_view.rightToLeft, ws.title)

    def test_two_month_columns(self):
        self.assertEqual(self.results["standard"]["n_months"], 2)
        ws = self.wb()[common.SHEET_EXP]
        hm = header_map(ws)
        for m in EXPECTED["months"]:
            self.assertIn(common.month_label(m), hm)
        self.assertIn('סה"כ 2 חודשים', hm)
        self.assertEqual(hm["ממוצע ללא אפס"], hm["ממוצע חודשי"] + 1)      # always adjacent
        # v0.1.1: no budget columns anywhere; fixed column order after the month columns
        tail = ['סה"כ 2 חודשים', "ממוצע חודשי", "ממוצע ללא אפס", "מקסימום", "חודש מקס", "מינימום", "חודש מין",
                "חציון", "סטיית תקן", "חודשים עם חיוב", "הערות"]
        for sheet in (common.SHEET_EXP, common.SHEET_INC):
            ws2 = self.wb()[sheet]
            hdrs = [ws2.cell(2, c).value for c in range(1, ws2.max_column + 1)]
            self.assertFalse(any("תקציב" in str(h) for h in hdrs), sheet)
            self.assertEqual(hdrs, ["קבוצה", "תת-קטגוריה"] + [common.month_label(m) for m in EXPECTED["months"]] + tail, sheet)

    def test_named_ranges(self):
        wb = self.wb()
        names = set(wb.defined_names.keys())
        for nm in ("DBname", "DBamt", "DBmonth", "DBtype", "DBgroup", "DBcat", "DBpay", "DBperson", "DBsummed", "DBin", "DBmname"):
            self.assertIn(nm, names)
        ref = wb.defined_names["DBamt"].attr_text
        self.assertIn("database", ref)
        self.assertTrue(ref.endswith("$%d" % (EXPECTED["n_rows"] + 1)), ref)
        self.assertEqual(wb[common.SHEET_DB].tables["DB"].ref, "A1:%s%d" % (openpyxl.utils.get_column_letter(len(common.DB_COLUMNS)), EXPECTED["n_rows"] + 1))

    # ---- statistic sheet structure --------------------------------------------------------
    def total_row(self, ws, label):
        rows = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 2).value == label]
        self.assertEqual(len(rows), 1, label)
        return rows[0]

    def test_total_row_references_database(self):
        wb = self.wb()
        we, pv = wb[common.SHEET_EXP], wb[common.SHEET_HELPER]
        tr = self.total_row(we, 'סה"כ הוצאות שוטפות (נסכם ישירות מ-database)')
        hm = header_map(we)
        f = formula_text(we.cell(tr, hm[common.month_label("2025-01")]).value)
        m = re.match(r"^='עזר_חודשי'!([A-Z]+)(\d+)$", f)
        self.assertTrue(m, f)
        helper = formula_text(pv.cell(int(m.group(2)), 5).value)
        self.assertTrue(helper.startswith("=SUMIFS(DBamt,DBtype,"), helper)
        self.assertIn('DBsummed,"כן"', helper)
        self.assertIn("DBmonth", helper)
        tot = formula_text(we.cell(tr, hm['סה"כ 2 חודשים']).value)
        self.assertTrue(tot.startswith("=SUM("), tot)
        # the check row below the total excludes the subtotal rows and shows ✓/✗
        chk = formula_text(we.cell(tr + 1, hm['סה"כ 2 חודשים']).value)
        self.assertTrue(chk.startswith("=SUM(") and "-" in chk, chk)
        self.assertIn('"✓ תואם לסה""כ"', formula_text(we.cell(tr + 1, hm["הערות"]).value))
        ti = self.total_row(wb[common.SHEET_INC], 'סה"כ הכנסות (נסכם ישירות מ-database)')
        self.assertGreater(ti, 3)

    def test_group_subtotals_and_sorting(self):
        we = self.wb()[common.SHEET_EXP]
        tr = self.total_row(we, 'סה"כ הוצאות שוטפות (נסכם ישירות מ-database)')
        col2 = [we.cell(r, 2).value for r in range(3, tr)]
        subtotals = [v for v in col2 if str(v).startswith('סה"כ ')]
        self.assertEqual(subtotals, ['סה"כ %s' % g for g in EXPECTED["group_order"]])
        cats = [v for v in col2 if not str(v).startswith('סה"כ ')]
        self.assertEqual(cats[:2], ["שכר דירה", "ועד בית"])            # דיור first (largest total), by average desc
        self.assertEqual(cats[2:4], ["סופרמרקט", "מסעדות"])
        self.assertEqual(set(cats), set(EXPECTED["cat_totals"]))
        hm = header_map(we)
        sub_row = 3 + col2.index('סה"כ דיור')
        f = formula_text(we.cell(sub_row, hm[common.month_label("2025-01")]).value)
        self.assertRegex(f, r"^=[A-Z]+\d+\+[A-Z]+\d+$")                 # '+' of the member rows
        self.assertEqual(we.cell(sub_row, 2).fill.fgColor.rgb[-6:], "FCE4D6")     # distinct subtotal fill (col 1 is merged per group)

    def test_zero_categories_dropped(self):
        we = self.wb()[common.SHEET_EXP]
        col2 = {we.cell(r, 2).value for r in range(1, we.max_row + 1)}
        for c in EXPECTED["dropped_categories"]:
            self.assertNotIn(c, col2)
        self.assertEqual(self.results["standard"]["dropped_categories"], EXPECTED["dropped_categories"])
        self.assertIn("גז", self.wb()[common.SHEET_LISTS]["J"] and [c.value for c in self.wb()[common.SHEET_LISTS]["J"]])  # still offered in the dropdown

    def test_stat_formulas(self):
        we = self.wb()[common.SHEET_EXP]
        hm = header_map(we)
        r = next(r for r in range(3, we.max_row) if we.cell(r, 2).value == "ועד בית")
        self.assertEqual(formula_text(we.cell(r, hm["ממוצע ללא אפס"]).value), '=IFERROR(AVERAGEIF(C%d:D%d,"<>0"),0)' % (r, r))
        self.assertIn("_xlfn.STDEV.S(", formula_text(we.cell(r, hm["סטיית תקן"]).value))
        self.assertTrue(formula_text(we.cell(r, hm["חודש מקס"]).value).startswith("=INDEX($C$2:$D$2,MATCH("))
        self.assertIsNotNone(we.cell(r, hm["מקסימום"]).comment)
        self.assertIn(common.month_label("2025-01"), we.cell(r, hm["מקסימום"]).comment.text)
        self.assertEqual(we.cell(r, hm["ממוצע ללא אפס"]).number_format, "#,##0;[Red]-#,##0")
        # red-bold highlight rule above the threshold, colour scales present
        rules = [rule for rng in we.conditional_formatting for rule in rng.rules]
        self.assertTrue(any(r_.type == "expression" and "3000" in r_.formula[0] for r_ in rules))
        self.assertTrue(any(r_.type == "colorScale" for r_ in rules))

    def test_non_summed_blocks(self):
        we = self.wb()[common.SHEET_EXP]
        col2 = [we.cell(r, 2).value for r in range(1, we.max_row + 1)]
        for c in ("קופת גמל", "העברה בין חשבונות", "ויזה 1234", "נסיעה אקדמית"):
            self.assertIn(c, col2)
        bal = next(r for r in range(1, we.max_row + 1) if str(we.cell(r, 2).value or "").startswith("מאזן פתוח"))
        f = formula_text(we.cell(bal, 3).value)
        self.assertIn('DBtype,"הוצאה בהחזר"', f)
        self.assertIn('DBtype,"החזר הוצאה"', f)

    def test_incomes_sheet(self):
        wi = self.wb()[common.SHEET_INC]
        ti = self.total_row(wi, 'סה"כ הכנסות (נסכם ישירות מ-database)')
        col2 = [wi.cell(r, 2).value for r in range(3, ti)]
        self.assertIn('סה"כ הכנסות %s' % synth_db.P1, col2)
        self.assertEqual(col2[0], "משכורת א'")                                # largest group first
        self.assertTrue(str(wi.cell(ti + 2, 2).value).startswith("מאזן חודשי"))

    # ---- interactive sheets ----------------------------------------------------------------
    def test_pivot_dynamic_arrays(self):
        wd = self.wb()[common.SHEET_PIVOT]
        self.assertIsInstance(wd["F4"].value, ArrayFormula)
        self.assertIn("SUMPRODUCT", wd["F4"].value.text)
        self.assertEqual(wd["C6"].value, "שכר דירה")                            # top category by average
        t = wd["B24"].value
        self.assertIsInstance(t, ArrayFormula)
        self.assertEqual(t.ref, "B24:F%d" % (24 + 250 - 1))
        for token in ("_xlfn.LET(", "_xlpm.c,", "_xlfn._xlws.FILTER(", "_xlfn.UNIQUE(", "_xlfn.SORTBY(", "_xlfn.SEQUENCE(250)"):
            self.assertIn(token, t.text)
        self.assertEqual(wd["H24"].value.ref, "H24:M%d" % (24 + 500 - 1))
        self.assertIn("F4/2", wd["F7"].value)                                    # ÷ N months
        link = formula_text(wd["G24"].value)
        self.assertIn("HYPERLINK(", link)
        self.assertIn("A_keys", link)
        dvs = [dv.formula1 for dv in wd.data_validations.dataValidation]
        self.assertIn("=L_months", dvs)
        self.assertTrue(any("INDIRECT(" in f for f in dvs))

    def test_hyperlinks_target_tall_ranges(self):
        wb = self.wb()
        we = wb[common.SHEET_EXP]
        r = next(r for r in range(3, we.max_row) if we.cell(r, 2).value == "סופרמרקט")
        target = we.cell(r, 2).hyperlink.target
        m = re.match(r"#'פירוט לפי קטגוריה'!A(\d+):N(\d+)$", target)
        self.assertTrue(m, target)
        self.assertEqual(int(m.group(2)) - int(m.group(1)), 45)
        wc = wb[common.SHEET_BY_CAT]
        anchor = int(m.group(1))
        self.assertEqual(wc.cell(anchor, 1).value, "מזון / סופרמרקט")
        back = wc.cell(anchor, 14).hyperlink.target
        self.assertTrue(back.startswith("#'הוצאות'!B%d:" % r), back)
        merchant = wc.cell(anchor + 2, 1)
        self.assertEqual(merchant.value, "סופר לדוגמה")
        self.assertTrue(merchant.hyperlink.target.startswith("#'פירוט עסקאות'!A"))
        self.assertTrue(formula_text(wc.cell(anchor + 2, 2).value).startswith("=SUMIFS(DBamt,DBname,$A%d,DBcat,\"סופרמרקט\"" % (anchor + 2)))

    def test_txns_and_variable_sheets(self):
        wb = self.wb()
        wt = wb[common.SHEET_TXNS]
        refs = [formula_text(c.value) for row in wt.iter_rows(min_row=3) for c in row if isinstance(formula_text(c.value), str) and formula_text(c.value).startswith("='database'!")]
        self.assertGreater(len(refs), 30 * 10 - 1)                             # 30 summed expense rows × 10 columns
        wv = wb[common.SHEET_VARIABLE]
        names = [wv.cell(r, 2).value for r in range(4, wv.max_row) if isinstance(wv.cell(r, 1).value, int)]
        self.assertNotIn("שכר דירה", names)                                     # fixed categories excluded
        self.assertEqual(names[0], "סופר לדוגמה")                               # sorted by total desc
        self.assertTrue(wv.auto_filter.ref)

    def test_analysis_charts(self):
        wa = self.wb("deep")[common.SHEET_ANALYSIS]
        self.assertGreaterEqual(len(wa._charts), 12)
        titles = [wa.cell(r, 1).value for r in range(1, wa.max_row + 1) if isinstance(wa.cell(r, 1).value, str)]
        self.assertTrue(any(t.startswith("ו-nz.") for t in titles))          # no-zero twin of the top-merchant chart
        self.assertTrue(any(t.startswith("י.") for t in titles))             # secondary scheme block (deep)
        first = next(r for r in range(1, wa.max_row) if str(wa.cell(r, 1).value or "").startswith("א."))
        self.assertEqual(wa.cell(first + 2, 1).value, common.month_label("2025-01"))

    # ---- לסיווג ידני ---------------------------------------------------------------------
    def test_manual_sheet(self):
        wm = self.wb()[common.SHEET_MANUAL]
        hdrs = [wm.cell(3, c).value for c in range(1, wm.max_column + 1)]
        self.assertEqual(hdrs[:len(common.DB_COLUMNS)], [common.DB_HEADERS_HE[k] for k in common.DB_COLUMNS])
        self.assertEqual(hdrs[-3:], ["נדרש:", common.USER_TEXT_COL, common.USER_CAT_COL])
        ids = [wm.cell(r, 1).value for r in range(4, wm.max_row + 1) if wm.cell(r, 1).value]
        self.assertEqual(set(ids), set(EXPECTED["manual_rows"]))
        self.assertEqual(self.results["standard"]["manual_rows"], len(EXPECTED["manual_rows"]))
        need = {wm.cell(r, 1).value: wm.cell(r, len(common.DB_COLUMNS) + 1).value for r in range(4, wm.max_row + 1)}
        self.assertTrue(need["SYN-UNK"].startswith("נדרש: בחירת קטגוריה"))
        self.assertEqual(need["SYN-PBX"], "נדרש: אשר את הקטגוריה או בחר אחרת")
        self.assertIn("ManualTable", wm.tables)
        self.assertTrue(any(dv.formula1 == "=L_allcats" for dv in wm.data_validations.dataValidation))

    def test_labels_survive_rebuild_and_no_budget_columns(self):
        path = os.path.join(self.root, "outputs", "roundtrip.xlsx")
        rc, res, err = run(BUILD, "--config", self.cfg_path, "--level", "standard", "--out", path)
        self.assertEqual(rc, 0, err[-500:])
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames[0], common.SHEET_EXP)                       # no template tab
        we, wm = wb[common.SHEET_EXP], wb[common.SHEET_MANUAL]
        for ws in wb.worksheets:
            for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 3)):
                for cell in row:
                    self.assertNotIn("תקציב", str(cell.value), "%s!%s" % (ws.title, cell.coordinate))
        self.assertNotIn("preserved_budgets", res)
        hdrs = [wm.cell(3, c).value for c in range(1, wm.max_column + 1)]
        ut, uc = hdrs.index(common.USER_TEXT_COL) + 1, hdrs.index(common.USER_CAT_COL) + 1
        rm = next(r for r in range(4, wm.max_row + 1) if wm.cell(r, 1).value == "SYN-UNK")
        wm.cell(rm, ut, "חנות בגדים לדוגמה")
        wm.cell(rm, uc, "סופרמרקט")
        wb.save(path)
        rc, res, err = run(BUILD, "--config", self.cfg_path, "--level", "standard", "--out", path)
        self.assertEqual(rc, 0, err[-500:])
        self.assertEqual(res["preserved_labels"], 1)
        wb = openpyxl.load_workbook(path)
        self.assertEqual(wb.sheetnames[0], common.SHEET_EXP)
        we, wm = wb[common.SHEET_EXP], wb[common.SHEET_MANUAL]
        self.assertFalse(any("תקציב" in str(h) for h in header_map(we)))
        rm = next(r for r in range(4, wm.max_row + 1) if wm.cell(r, 1).value == "SYN-UNK")
        self.assertEqual(wm.cell(rm, ut).value, "חנות בגדים לדוגמה")
        self.assertEqual(wm.cell(rm, uc).value, "סופרמרקט")

    # ---- error paths / variants ----------------------------------------------------------
    def test_comma_in_category_fails(self):
        root = tempfile.mkdtemp(prefix="tz_comma_")
        try:
            cfg_path, db_path = synth_db.write_project(root)
            rows = synth_db.rows()
            rows[0]["cat_tz"] = "מזון, מכולת"
            common.write_csv(db_path, rows, common.DB_COLUMNS)
            rc, res, err = run(BUILD, "--config", cfg_path, "--level", "standard")
            self.assertEqual(rc, 1)
            self.assertFalse(res["ok"])
            self.assertIn("comma", res["error"])
            self.assertIn("מזון, מכולת", res["error"])
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_missing_database_fails(self):
        root = tempfile.mkdtemp(prefix="tz_nodb_")
        try:
            cfg_path, db_path = synth_db.write_project(root)
            os.remove(db_path)
            rc, res, err = run(BUILD, "--config", cfg_path)
            self.assertEqual(rc, 1)
            self.assertIn("database not found", res["error"])
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_three_month_window(self):
        root = tempfile.mkdtemp(prefix="tz_n3_")
        try:
            cfg_path, _ = synth_db.write_project(root)
            cfg = synth_db.config()
            cfg["window"]["end"] = "2025-03-31"
            common.write_json(cfg_path, cfg)
            rc, res, err = run(BUILD, "--config", cfg_path, "--level", "overview")
            self.assertEqual(rc, 0, err[-500:])
            self.assertEqual(res["n_months"], 3)
            we = openpyxl.load_workbook(res["workbook"])[common.SHEET_EXP]
            hm = header_map(we)
            self.assertIn(common.month_label("2025-03"), hm)
            self.assertIn('סה"כ 3 חודשים', hm)
            r = next(r for r in range(3, we.max_row) if we.cell(r, 2).value == "דלק")
            self.assertEqual(formula_text(we.cell(r, hm['סה"כ 3 חודשים']).value), "=SUM(C%d:E%d)" % (r, r))
            self.assertEqual(hm["ממוצע חודשי"], hm['סה"כ 3 חודשים'] + 1)
        finally:
            shutil.rmtree(root, ignore_errors=True)

    def test_template_mode(self):
        root = tempfile.mkdtemp(prefix="tz_tpl_")
        try:
            cfg_path, _ = synth_db.write_project(root)
            tpl = os.path.join(root, "inputs", "template", "תזרים.xlsx")
            os.makedirs(os.path.dirname(tpl))
            twb = openpyxl.Workbook()
            ws = twb.active
            ws.title = "הכנסות - הוצאות"
            ws["A1"], ws["B1"], ws["C1"] = "קבוצה", "קטגוריה", "ממוצע חודשי"
            for r, (g, c) in enumerate([("מזון", "סופרמרקט"), ("מזון", "מסעדות"), ("דיור", "שכר דירה"), ("דיור", "גז")], 2):
                ws.cell(r, 1, g)
                ws.cell(r, 2, c)
            ws["C4"] = 4321
            twb.save(tpl)
            tpl_mtime = os.path.getmtime(tpl)
            cfg = synth_db.config()
            # link_avg_col / manual_cells are legacy keys (v0.1.0): accepted and ignored with a warning
            cfg["categories"]["from_template"] = {"file": "inputs/template/תזרים.xlsx", "sheet": "הכנסות - הוצאות", "groups_col": "A",
                                                  "cats_col": "B", "first_row": 2, "last_row": 5, "link_avg_col": "C", "manual_cells": ["C4"]}
            common.write_json(cfg_path, cfg)
            rc, res, err = run(BUILD, "--config", cfg_path, "--level", "standard")
            self.assertEqual(rc, 0, err[-500:])
            self.assertIn("link_avg_col is ignored", err)
            self.assertEqual(res["template"]["template_categories"], 4)
            wb = openpyxl.load_workbook(res["workbook"])
            self.assertEqual(wb.sheetnames[0], common.SHEET_EXP)                     # workbook built from scratch
            self.assertNotIn("הכנסות - הוצאות", wb.sheetnames)                       # template tab not copied
            self.assertEqual(wb.sheetnames, [s for s in common.SHEET_ORDER if s in wb.sheetnames])
            we = wb[common.SHEET_EXP]
            cats = [we.cell(r, 2).value for r in range(3, we.max_row + 1)]
            self.assertIn("סופרמרקט", cats)                                          # template category with data
            self.assertIn("גז", res["dropped_categories"])                           # template category without data
            self.assertTrue(os.path.isfile(tpl))                                     # template untouched
            self.assertEqual(os.path.getmtime(tpl), tpl_mtime)
            self.assertEqual(openpyxl.load_workbook(tpl)["הכנסות - הוצאות"]["C4"].value, 4321)
        finally:
            shutil.rmtree(root, ignore_errors=True)

    # ---- transfers sheet -------------------------------------------------------------------
    def test_transfers_sheet(self):
        rc, res, err = run(TRANSFERS, "--config", self.cfg_path, "--level", "overview", "--workbook", self.books["overview"])
        self.assertEqual(rc, 0)
        self.assertTrue(res["skipped"])
        rc, res, err = run(TRANSFERS, "--config", self.cfg_path, "--level", "standard", "--workbook", self.books["standard"])
        self.assertEqual(rc, 0, err[-500:])
        self.assertEqual(res["sections"]["bank"], {"rows": 7, "out": 12650.0, "in": 300.0})   # 12,400 + 250 (gemel deposit 1,750)
        self.assertEqual(res["sections"]["p2p"], {"rows": 3, "out": 400.0, "in": 100.0})
        self.assertEqual(res["sections"]["paybox"], {"rows": 1, "out": 150.0, "in": 0})
        self.assertEqual(res["flagged"], ["SYN-PBX"])
        wb = openpyxl.load_workbook(self.books["standard"])
        self.assertEqual(wb.sheetnames.index(common.SHEET_TRANSFERS), wb.sheetnames.index(common.SHEET_ANALYSIS) + 1)
        ws = wb[common.SHEET_TRANSFERS]
        subtot = [r for r in range(1, ws.max_row + 1) if ws.cell(r, 1).value == 'סה"כ']
        self.assertEqual(len(subtot), 3)
        f = formula_text(ws.cell(subtot[0], 6).value)
        self.assertRegex(f, r'^=SUMIFS\(F\d+:F\d+,C\d+:C\d+,"יוצא"\)$')
        self.assertTrue(formula_text(ws.cell(subtot[0], 3).value).startswith("=COUNTA("))
        titles = [ws.cell(r, 1).value for r in range(1, ws.max_row + 1) if str(ws.cell(r, 1).value or "").startswith(("1.", "2.", "3."))]
        self.assertEqual(len(titles), 3)
        for t in titles:
            self.assertNotIn(synth_db.P1, t)
        # rerun replaces the sheet in place (same count, same position)
        rc, res2, err = run(TRANSFERS, "--config", self.cfg_path, "--level", "standard", "--workbook", self.books["standard"])
        self.assertEqual(rc, 0)
        self.assertEqual(openpyxl.load_workbook(self.books["standard"]).sheetnames.count(common.SHEET_TRANSFERS), 1)

    # ---- apply_user_labels / run_pipeline -----------------------------------------------
    def test_apply_user_labels_merges_and_persists(self):
        path = os.path.join(self.root, "outputs", "labels.xlsx")
        rc, res, err = run(BUILD, "--config", self.cfg_path, "--level", "standard", "--out", path)
        self.assertEqual(rc, 0, err[-500:])
        wb = openpyxl.load_workbook(path)
        wm = wb[common.SHEET_MANUAL]
        hdrs = [wm.cell(3, c).value for c in range(1, wm.max_column + 1)]
        ut, uc = hdrs.index(common.USER_TEXT_COL) + 1, hdrs.index(common.USER_CAT_COL) + 1
        rm = next(r for r in range(4, wm.max_row + 1) if wm.cell(r, 1).value == "SYN-PBX")
        wm.cell(rm, uc, "מסעדות")
        wb.save(path)
        labels = os.path.join(self.root, "rules", "user_labels.csv")
        common.write_csv(labels, [{"id": "OLD-1", "user_text": "ישן", "user_cat": "", "timestamp": "2025-01-01T00:00:00"}],
                         ["id", "user_text", "user_cat", "timestamp"])
        rc, res, err = run(APPLY, "--config", self.cfg_path, "--level", "standard", "--workbook", path, "--no-run")
        self.assertEqual(rc, 0, err[-500:])
        self.assertEqual(res["labels"]["read"], 1)
        self.assertEqual(res["labels"]["merged"], 2)
        rows = {r["id"]: r for r in common.read_csv(labels)}
        self.assertEqual(rows["SYN-PBX"]["user_cat"], "מסעדות")
        self.assertEqual(rows["OLD-1"]["user_text"], "ישן")                    # earlier label kept
        self.assertTrue(rows["SYN-PBX"]["timestamp"])
        self.assertIn("reminder", res)

    def test_run_pipeline_plan_and_range(self):
        rc, res, err = run(PIPE, "--config", self.cfg_path, "--mode", "A", "--dry-run")
        self.assertEqual(rc, 0, err[-500:])
        names = [s["name"] for s in res["steps"]]
        self.assertEqual(names[0], "parse_bank")
        self.assertEqual(names[-1], "build_dashboard")
        status = {s["name"]: s["status"] for s in res["steps"]}
        self.assertEqual(status["parse_benefits"], "skipped")                    # no benefit_programs in config
        self.assertEqual(status["parse_card_blocks"], "skipped")                     # only cycle-format cards configured
        self.assertIn(status["build_excel"], ("planned",))
        rc, res, err = run(PIPE, "--config", self.cfg_path, "--mode", "B", "--dry-run")
        self.assertEqual([s["name"] for s in res["steps"]][0], "classify")
        rc, res, err = run(PIPE, "--config", self.cfg_path, "--mode", "B", "--from", "build_excel", "--to", "verify", "--level", "overview")
        self.assertEqual(rc, 0, err[-800:])
        self.assertEqual([s["name"] for s in res["steps"]], ["build_excel", "add_transfers_sheet", "verify"])
        self.assertTrue(all(s["status"] == "ok" for s in res["steps"]), res["steps"])
        rc, res, err = run(PIPE, "--config", self.cfg_path, "--from", "verify", "--to", "classify")
        self.assertEqual(rc, 1)
        self.assertFalse(res["ok"])


if __name__ == "__main__":
    unittest.main()
