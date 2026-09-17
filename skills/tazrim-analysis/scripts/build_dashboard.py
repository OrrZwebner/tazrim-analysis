#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_dashboard.py — single-file interactive Hebrew RTL dashboard (P17).

Purpose : build outputs/dashboard.html from the `database` sheet of the workbook (so the dashboard
          reflects the Excel; work/database.csv is the fallback, and fills columns the sheet lacks)
          plus the transfers sheet's per-row section / direction / payee / note when present.
Inputs  : tazrim.config.json (title = household.label, people order, window, chartjs mode, level);
          outputs/תזרים.xlsx (sheet `database`, Hebrew headers = common.DB_HEADERS_HE; optional sheet
          `העברות, BIT ו-PAYBOX`); work/database.csv.
Outputs : outputs/dashboard.html; ONE JSON object on stdout:
          {"ok", "dashboard", "bytes", "rows", "cols", "source", "chartjs", "tabs", "level"}.
Exit    : 0 ok; 1 neither workbook nor database.csv readable; 2 config error.

Tabs: סקירה (DEFAULT — KPIs incl. no-zero average, drill-down group → category → merchant → rows, charts
with value labels) · עסקאות (every row × every column: free-text search, type chips, window/summed toggles,
date/amount ranges, per-column filters, column chooser, sortable headers, row modal with linked row
and same-merchant list, pivot, CSV export) · העברות, BIT ו-PAYBOX · הכנסות · לסיווג (points at the
Excel sheet) · עזרה (3-line usage box). overview level drops the transfers and לסיווג tabs.
Chart.js 4.4.1: `dashboard.chartjs == "cdn"` loads it from cdnjs (guarded — tables work offline);
`"inline"` embeds a copy cached under work/chartjs/ (downloaded once when the network allows;
otherwise falls back to cdn and says so in the JSON). No personal text lives in this generator.
Python 3.8 compatible. Needs openpyxl for the workbook path (CSV fallback is stdlib).
"""
import csv
import datetime as _dt
import json
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (  # noqa: E402
    DB_COLUMNS, DB_HEADERS_HE, SHARED, SHEET_DB, SHEET_MANUAL, SHEET_TRANSFERS, TYPE_DUPLICATE,
    TYPE_EXPENSE, TYPE_INCOME, TYPE_INTERNAL, TYPE_REIMBURSABLE, TYPE_REIMBURSEMENT, TYPE_SAVINGS,
    TYPE_CARD_DEBIT, UNKNOWN_CAT, USER_CAT_COL, USER_TEXT_COL, YES, emit, fail, log, month_label,
    months, owner_label, parse_args, project_path)

CHARTJS_URL = "https://cdnjs.cloudflare.com/ajax/libs/Chart.js/4.4.1/chart.umd.min.js"
CHARTJS_CACHE = "work/chartjs/chart.umd.min.js"
HEB2KEY = {v: k for k, v in DB_HEADERS_HE.items()}
BOOL_KEYS = ("in_window", "summed")
NUM_KEYS = ("amount", "orig_amount")
DATE_KEYS = ("txn_date", "charge_date")
KIND = {"txn_date": "date", "charge_date": "date", "amount": "num", "orig_amount": "num",
        "in_window": "bool", "summed": "bool"}
for _k in ("source", "card", "pay", "person", "type", "group_tz", "cat_tz", "group_new", "cat_new",
           "month", "month_name", "orig_currency", "trip", "source_file", "name_clean"):
    KIND[_k] = "cat"
EXTRA_HEADERS = {"tr_section": "סעיף (העברות)", "tr_dir": "כיוון", "tr_payee": "מוטב / צד שני", "tr_note": "הערה מצילום"}
TABS_BY_LEVEL = {
    "overview": ["overview", "explore", "income", "help"],
    "standard": ["overview", "explore", "transfers", "income", "unc", "help"],
    "deep": ["overview", "explore", "transfers", "income", "unc", "help"],
}


# ----------------------------------------------------------------------------- normalisation
def norm_bool(v):
    if isinstance(v, bool):
        return v
    s = str(v if v is not None else "").strip().lower()
    return s in (YES, "true", "1", "yes", "y")


def norm_num(v):
    if v is None or v == "":
        return None
    try:
        return float(str(v).replace(",", ""))
    except ValueError:
        return None


def norm_str(v):
    if v is None:
        return ""
    if isinstance(v, (_dt.datetime, _dt.date)):
        return v.strftime("%Y-%m-%d")
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v)


def clean(d):
    out = {}
    for k in DB_COLUMNS:
        v = d.get(k)
        if k in BOOL_KEYS:
            out[k] = norm_bool(v)
        elif k in NUM_KEYS:
            out[k] = norm_num(v)
        elif k in DATE_KEYS:
            out[k] = norm_str(v)[:10]
        else:
            s = norm_str(v).strip()
            out[k] = "" if s.lower() in ("nan", "none") else s
    return out


# ----------------------------------------------------------------------------- readers
def read_csv_rows(path):
    if not os.path.isfile(path):
        return []
    with open(path, encoding="utf-8-sig", newline="") as fh:
        return [dict(r) for r in csv.DictReader(fh)]


def read_workbook(path):
    """(rows, order, transfers or None) from the workbook; rows are raw dicts keyed by DB column."""
    import openpyxl  # noqa: WPS433 (optional dependency)
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if SHEET_DB not in wb.sheetnames:
        wb.close()
        raise ValueError("sheet '%s' not found in %s (sheets: %s)" % (SHEET_DB, os.path.basename(path), ", ".join(wb.sheetnames)))
    ws = wb[SHEET_DB]
    it = ws.iter_rows(values_only=True)
    header = [norm_str(h).strip() for h in next(it)]
    keys = [HEB2KEY.get(h) or (h if h in DB_COLUMNS else None) for h in header]
    unknown = [h for h, k in zip(header, keys) if k is None and h]
    if unknown:
        log("note: unmapped workbook headers ignored: %s" % ", ".join(unknown))
    rows = []
    for r in it:
        if r is None or all(v is None or v == "" for v in r):
            continue
        d = {k: v for k, v in zip(keys, r) if k}
        if not norm_str(d.get("id")).strip():
            continue
        rows.append(d)
    order = [k for k in keys if k]
    tr = read_transfers_sheet(wb)
    wb.close()
    return rows, order, tr


def read_transfers_sheet(wb):
    """Per-id annotations from the transfers sheet: section 1..3, direction, payee, note."""
    if SHEET_TRANSFERS not in wb.sheetnames:
        return None
    ws = wb[SHEET_TRANSFERS]
    sec, hdr, out = 0, None, {}
    for r in ws.iter_rows(values_only=True):
        a = r[0] if r else None
        if isinstance(a, str) and len(a) > 2 and a[0] in "123" and a[1] == ".":
            sec = int(a[0])
            continue
        if a == DB_HEADERS_HE["id"]:
            hdr = [norm_str(h) for h in r]
            continue
        if hdr and isinstance(a, str) and a.strip() and not a.startswith("סה\"כ"):
            d = dict(zip(hdr, r))

            def pick(sub):
                for h, v in d.items():
                    if sub in h:
                        return norm_str(v)
                return ""
            out[a] = {"tr_section": sec, "tr_dir": pick("כיוון"), "tr_payee": pick("מוטב"), "tr_note": pick("הערה")}
    return out


def derive_transfers(cfg, data):
    """Fallback annotations when the transfers sheet is missing (F17 direction; sections by type/source)."""
    accounts = [(a.get("label", ""), a["id"]) for a in cfg.get("accounts", [])]
    p2p = cfg.get("p2p", [])
    p2p_labels = [p.get("label", "") for p in p2p]
    markers = [str(p.get("card_marker") or "").upper() for p in p2p if p.get("card_marker")]
    out = {}
    for d in data:
        src = d["source"]
        is_bank = any(src == lbl or src == aid for lbl, aid in accounts)
        is_p2p = any(src == lbl for lbl in p2p_labels)
        name = (d["original_name"] or "").upper()
        sec = 0
        if is_bank and d["type"] in (TYPE_INTERNAL, TYPE_SAVINGS, TYPE_REIMBURSEMENT, TYPE_REIMBURSABLE):
            sec = 1
        elif is_p2p or (markers and any(m in name for m in markers) and not is_bank):
            sec = 2
        elif "PAYBOX" in name:
            sec = 3
        if not sec:
            continue
        amt = d["amount"] or 0.0
        inbound = (d["type"] == TYPE_INCOME) != (amt < 0)
        out[d["id"]] = {"tr_section": sec, "tr_dir": "נכנס" if inbound else "יוצא", "tr_payee": "", "tr_note": ""}
    return out


# ----------------------------------------------------------------------------- chart.js
def chartjs_tag(cfg):
    """(script tag, mode string). inline embeds the cached file; downloads once when possible."""
    mode = cfg.get("dashboard", {}).get("chartjs", "cdn")
    if mode != "inline":
        return '<script src="%s"></script>' % CHARTJS_URL, "cdn"
    cache = project_path(cfg, CHARTJS_CACHE)
    if not os.path.isfile(cache):
        try:
            import urllib.request
            os.makedirs(os.path.dirname(cache), exist_ok=True)
            with urllib.request.urlopen(CHARTJS_URL, timeout=20) as resp:  # noqa: S310
                body = resp.read()
            if b"Chart" not in body[:4000]:
                raise ValueError("unexpected content")
            with open(cache, "wb") as fh:
                fh.write(body)
            log("chart.js cached at %s" % cache)
        except Exception as e:  # noqa: BLE001
            log("WARNING: could not download Chart.js for inline mode (%s); falling back to cdn" % e)
            return '<script src="%s"></script>' % CHARTJS_URL, "cdn (fallback: download failed)"
    with open(cache, encoding="utf-8") as fh:
        js = fh.read().replace("</script", "<\\/script")
    return "<script>%s</script>" % js, "inline"


# ----------------------------------------------------------------------------- build
def load_data(cfg, workbook=None, database=None):
    """Return (data rows, column order, transfers dict, source description)."""
    wb_path = workbook or project_path(cfg, "outputs.workbook")
    csv_path = database or project_path(cfg, "work.database")
    rows, order, tr, source = [], [], None, ""
    if os.path.isfile(wb_path):
        try:
            rows, order, tr = read_workbook(wb_path)
            source = os.path.basename(wb_path)
        except ImportError:
            log("WARNING: openpyxl not installed — reading work/database.csv instead")
        except ValueError as e:
            log("WARNING: %s — reading work/database.csv instead" % e)
    if not rows:
        rows = read_csv_rows(csv_path)
        order = list(DB_COLUMNS)
        source = os.path.basename(csv_path)
        if not rows:
            raise FileNotFoundError("neither %s nor %s is readable" % (wb_path, csv_path))
    else:
        missing = [k for k in DB_COLUMNS if k not in order]
        if missing:
            by_id = {r.get("id"): r for r in read_csv_rows(csv_path)}
            log("columns missing from the workbook, taken from CSV: %s" % ", ".join(missing))
            for d in rows:
                c = by_id.get(norm_str(d.get("id")))
                if c:
                    for k in missing:
                        d[k] = c.get(k)
            order = order + missing
    data = [clean(d) for d in rows]
    if tr is None:
        tr = derive_transfers(cfg, data)
    n_tr = {1: 0, 2: 0, 3: 0}
    for d in data:
        a = tr.get(d["id"])
        d["tr_section"] = a["tr_section"] if a else 0
        d["tr_dir"] = a["tr_dir"] if a else ""
        d["tr_payee"] = a["tr_payee"] if a else ""
        d["tr_note"] = a["tr_note"] if a else ""
        if a:
            n_tr[a["tr_section"]] = n_tr.get(a["tr_section"], 0) + 1
    return data, order, n_tr, source


def build_dashboard(cfg, workbook=None, database=None, out_path=None, level=None):
    data, order, n_tr, source = load_data(cfg, workbook, database)
    level = level or cfg.level
    win = months(cfg)
    all_months = sorted(set(d["month"] for d in data if d["month"]) | set(win))
    month_names = {m: month_label(m) for m in all_months}
    for d in data:
        if d["month"] and d["month_name"]:
            month_names.setdefault(d["month"], d["month_name"])
    by_type = {}
    for d in data:
        by_type[d["type"]] = by_type.get(d["type"], 0) + 1
    default_rows = [d for d in data if d["type"] in (TYPE_EXPENSE, TYPE_INCOME) and d["summed"] and d["in_window"]]
    people = [owner_label(cfg, p["id"]) for p in cfg["household"]["people"]] + [owner_label(cfg, SHARED)]
    cols = [{"k": k, "h": DB_HEADERS_HE.get(k, k), "t": KIND.get(k, "text")} for k in order if k in DB_COLUMNS]
    cols += [{"k": k, "h": h, "t": "cat"} for k, h in EXTRA_HEADERS.items()]
    built = _dt.datetime.now().strftime("%Y-%m-%d %H:%M")
    tabs = TABS_BY_LEVEL.get(level, TABS_BY_LEVEL["standard"])
    meta = {
        "title": cfg["household"].get("label", ""), "built": built, "source_file": source,
        "sheet": SHEET_DB, "tr_sheet": SHEET_TRANSFERS, "manual_sheet": SHEET_MANUAL,
        "user_text_col": USER_TEXT_COL, "user_cat_col": USER_CAT_COL, "unknown_cat": UNKNOWN_CAT,
        "cols": cols, "n_total": len(data), "n_by_type": by_type, "n_default": len(default_rows),
        "n_default_exp": sum(1 for d in default_rows if d["type"] == TYPE_EXPENSE),
        "n_default_inc": sum(1 for d in default_rows if d["type"] == TYPE_INCOME),
        "n_tr": n_tr, "months": all_months, "month_names": month_names, "window": win,
        "window_label": "%s–%s" % (month_label(win[0]), month_label(win[-1])) if win else "",
        "people": people, "has_trip": any(d["trip"] for d in data), "level": level, "tabs": tabs,
        "types": {"exp": TYPE_EXPENSE, "inc": TYPE_INCOME, "internal": TYPE_INTERNAL, "savings": TYPE_SAVINGS,
                  "card_debit": TYPE_CARD_DEBIT, "dup": TYPE_DUPLICATE, "reimb_paid": TYPE_REIMBURSABLE,
                  "reimb_recv": TYPE_REIMBURSEMENT},
        "refresh_cmd": "python3 %s --config tazrim.config.json" % os.path.basename(__file__),
    }
    script, chart_mode = chartjs_tag(cfg)

    def js_json(o):
        return json.dumps(o, ensure_ascii=False, separators=(",", ":")).replace("</", "<\\/")

    html = (TEMPLATE.replace("__DATA__", js_json(data)).replace("__META__", js_json(meta))
            .replace("__CHARTJS__", script).replace("__TITLE__", _esc(meta["title"]))
            .replace("__SRC__", _esc(source)).replace("__BUILT__", built))
    out_path = out_path or project_path(cfg, "outputs.dashboard")
    os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
    with open(out_path, "w", encoding="utf-8") as fh:
        fh.write(html)
    return {"ok": True, "dashboard": out_path, "bytes": len(html.encode("utf-8")), "rows": len(data),
            "cols": len(cols), "source": source, "chartjs": chart_mode, "tabs": tabs, "level": level,
            "rows_by_type": by_type, "transfers": n_tr}


def _esc(s):
    return (str(s).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace('"', "&quot;"))


# ----------------------------------------------------------------------------- template
TEMPLATE = r"""<!DOCTYPE html>
<html dir="rtl" lang="he">
<head>
<meta charset="utf-8">
<meta name="viewport" content="width=device-width, initial-scale=1">
<title>דשבורד תזרים — __TITLE__</title>
<style>
:root{--bg:#f6f6f3;--surface:#fcfcfb;--ink:#0b0b0b;--ink2:#52514e;--muted:#898781;--grid:#e1e0d9;--line:#c3c2b7;--ring:rgba(11,11,11,.10);--accent:#2a78d6;--accent-soft:#e3eefb;--income:#1baf7a;--radius:10px}
*{box-sizing:border-box}
html,body{margin:0;background:var(--bg);color:var(--ink);font-family:system-ui,-apple-system,"Segoe UI",Roboto,"Helvetica Neue",Arial,"Noto Sans Hebrew",sans-serif;font-size:14px;line-height:1.45}
body{padding:0 16px 32px}
h1{font-size:22px;margin:0;font-weight:700}
h2{font-size:15px;margin:0 0 10px;font-weight:600}
h3{font-size:14px;margin:14px 0 6px;font-weight:600}
a{color:var(--accent)}
.wrap{max-width:1500px;margin:0 auto}
header.top{display:flex;flex-wrap:wrap;align-items:baseline;gap:8px 16px;padding:16px 0 8px}
header.top .sub{color:var(--ink2);font-size:13px}
.card{background:var(--surface);border:1px solid var(--ring);border-radius:var(--radius);padding:14px 16px;box-shadow:0 1px 2px rgba(0,0,0,.03)}
nav.tabs{display:flex;flex-wrap:wrap;gap:4px;border-bottom:2px solid var(--grid);margin-bottom:12px;position:sticky;top:0;background:var(--bg);z-index:30;padding-top:4px}
nav.tabs button{border:0;background:none;font:inherit;font-size:14px;padding:8px 14px;cursor:pointer;color:var(--ink2);border-bottom:2px solid transparent;margin-bottom:-2px;border-radius:6px 6px 0 0}
nav.tabs button:hover{background:#eeede8}
nav.tabs button.on{color:var(--accent);border-bottom-color:var(--accent);font-weight:600}
nav.tabs button .n{color:var(--muted);font-size:11.5px;margin-right:4px}
nav.tabs button[hidden]{display:none}
.tab[hidden]{display:none}
.filters{display:flex;flex-wrap:wrap;gap:10px 12px;align-items:flex-end;margin-bottom:12px}
.f{display:flex;flex-direction:column;gap:4px;min-width:110px}
.f label{font-size:11.5px;color:var(--ink2)}
.f input[type=text],.f input[type=number],.f input[type=date],.f select{height:32px;border:1px solid var(--line);border-radius:7px;padding:0 8px;font:inherit;background:#fff;color:var(--ink);min-width:0}
.f input[type=number]{width:96px}
.f input[type=date]{width:140px}
.f input[type=text]{width:190px}
.f input.wide{width:260px}
.seg{display:inline-flex;border:1px solid var(--line);border-radius:7px;overflow:hidden;height:32px}
.seg button{border:0;background:#fff;padding:0 10px;font:inherit;cursor:pointer;color:var(--ink2);white-space:nowrap}
.seg button.on{background:var(--accent);color:#fff}
.tchips{display:flex;flex-wrap:wrap;gap:6px;align-items:center}
.tchip{border:1px solid var(--line);border-radius:999px;background:#fff;padding:4px 11px;font:inherit;font-size:12.5px;cursor:pointer;color:var(--ink2)}
.tchip.on{background:var(--accent);border-color:var(--accent);color:#fff}
.tchip .n{opacity:.7;font-size:11px;margin-right:4px}
.ms{position:relative}
.ms>button{height:32px;border:1px solid var(--line);border-radius:7px;background:#fff;padding:0 10px;font:inherit;cursor:pointer;color:var(--ink);min-width:120px;text-align:right;white-space:nowrap;overflow:hidden;text-overflow:ellipsis;max-width:240px}
.ms>button.act{border-color:var(--accent);background:var(--accent-soft)}
.ms>button::after{content:"▾";float:left;color:var(--muted)}
.ms .pop{display:none;position:absolute;top:36px;right:0;background:#fff;border:1px solid var(--line);border-radius:8px;box-shadow:0 8px 24px rgba(0,0,0,.12);min-width:240px;max-width:360px;max-height:340px;overflow:auto;padding:6px;z-index:50}
.ms.open .pop{display:block}
.ms .pop .tools{display:flex;gap:6px;padding:2px 4px 6px;border-bottom:1px solid var(--grid);margin-bottom:4px;align-items:center}
.ms .pop .tools button{border:0;background:none;color:var(--accent);cursor:pointer;font:inherit;font-size:12px;padding:2px 4px}
.ms .pop .tools input{flex:1;height:26px;border:1px solid var(--line);border-radius:5px;padding:0 6px;font:inherit;font-size:12px;min-width:60px}
.ms .pop label{display:flex;align-items:center;gap:8px;padding:4px 6px;border-radius:5px;cursor:pointer;font-size:13px}
.ms .pop label:hover{background:var(--bg)}
.ms .pop label span.t{overflow:hidden;text-overflow:ellipsis;white-space:nowrap}
.ms .pop label span.n{color:var(--muted);margin-right:auto;font-size:11.5px;padding-right:8px}
.btn{height:32px;border:1px solid var(--line);border-radius:7px;background:#fff;padding:0 12px;font:inherit;cursor:pointer;color:var(--ink);white-space:nowrap}
.btn:hover{background:var(--bg)}
.btn.primary{background:var(--accent);border-color:var(--accent);color:#fff}
.btn.sm{height:26px;padding:0 9px;font-size:12px}
.toggle{display:flex;align-items:center;gap:6px;height:32px;font-size:12.5px;color:var(--ink2);cursor:pointer}
.chips{display:flex;flex-wrap:wrap;gap:6px;margin:0 0 12px;min-height:6px}
.chip{display:inline-flex;align-items:center;gap:6px;background:var(--accent-soft);color:#174a86;border-radius:999px;padding:3px 10px;font-size:12.5px;max-width:100%}
.chip b{font-weight:600;overflow:hidden;text-overflow:ellipsis;white-space:nowrap;max-width:320px;display:inline-block;vertical-align:bottom}
.chip button{border:0;background:none;cursor:pointer;color:#174a86;font-size:14px;line-height:1;padding:0}
.chip.drill{background:#fdebd9;color:#7a3a10}
.chip.drill button{color:#7a3a10}
.kpis{display:grid;grid-template-columns:repeat(auto-fit,minmax(150px,1fr));gap:10px;margin-bottom:12px}
.kpi{background:var(--surface);border:1px solid var(--ring);border-radius:var(--radius);padding:12px 14px}
.kpi .l{font-size:12px;color:var(--ink2)}
.kpi .v{font-size:22px;font-weight:700;margin-top:2px}
.kpi .h{font-size:11px;color:var(--muted);margin-top:2px}
.kpi.click{cursor:pointer}
.kpi.click:hover{border-color:var(--accent)}
.kpi.sel{border-color:var(--accent);background:var(--accent-soft)}
.grid2{display:grid;grid-template-columns:1fr 1fr;gap:12px;margin-bottom:12px}
@media (max-width:900px){.grid2{grid-template-columns:1fr}}
.chartbox{position:relative;height:300px;width:100%}
.chartbox.tall{height:520px}
.nochart{color:var(--muted);font-size:13px;padding:24px 0;text-align:center}
.crumbs{display:flex;flex-wrap:wrap;align-items:center;gap:4px;margin-bottom:10px;font-size:13px}
.crumbs a{cursor:pointer;text-decoration:none;padding:3px 8px;border-radius:6px;background:var(--accent-soft)}
.crumbs .cur{padding:3px 8px;font-weight:600}
.crumbs .sep{color:var(--muted)}
.hint{font-size:12px;color:var(--muted);margin:0 0 8px}
.tw{overflow-x:auto;max-height:78vh;overflow-y:auto}
table{border-collapse:collapse;width:100%;font-size:13px}
th,td{padding:5px 8px;border-bottom:1px solid var(--grid);white-space:nowrap;text-align:right;vertical-align:top}
th{position:sticky;top:0;background:var(--surface);color:var(--ink2);font-weight:600;cursor:pointer;user-select:none;font-size:12px;z-index:2;box-shadow:0 1px 0 var(--grid)}
th.num,td.num{text-align:left;font-variant-numeric:tabular-nums;direction:ltr}
th .arr{color:var(--accent);font-size:10px;margin-right:3px}
tbody tr.click{cursor:pointer}
tbody tr.click:hover td{background:#eef3fb !important}
tbody tr.zebra:nth-child(even) td{background:#f7f7f4}
tbody tr.dim td{color:#9a9891}
td.name{max-width:280px;overflow:hidden;text-overflow:ellipsis}
td.txt{max-width:260px;overflow:hidden;text-overflow:ellipsis}
td.heat{color:#1f2937}
tfoot td{font-weight:600;background:#faf9f6;position:sticky;bottom:0}
.neg{color:#b42318}
.count{font-size:12px;color:var(--muted)}
.badge{display:inline-block;font-size:10.5px;padding:0 6px;border-radius:999px;background:#fde7e7;color:#8a1c1c;margin-right:6px;vertical-align:middle}
.badge.out{background:#eeede8;color:#6b6a65}
.badge.sec{background:#e3eefb;color:#174a86}
.note{background:#fff7ed;border:1px solid #fed7aa;border-radius:8px;padding:8px 12px;font-size:13px;margin:10px 0}
.howto{background:#eef5ff;border:1px solid #cfe0fa;border-radius:8px;padding:6px 12px;font-size:12.5px;margin:0 0 10px;color:#1c3d6b}
.howto summary{cursor:pointer;font-weight:600;list-style:none}
.howto ol{margin:6px 0 2px;padding-right:18px}
footer{margin-top:20px;color:var(--ink2);font-size:12.5px;line-height:1.7}
code{background:#eeede8;padding:1px 6px;border-radius:4px;direction:ltr;unicode-bidi:embed;font-size:12px}
.small{font-size:12px;color:var(--muted)}
.stats{display:flex;flex-wrap:wrap;gap:6px 18px;align-items:center;margin:6px 0 8px;font-size:13px}
.pager{display:flex;flex-wrap:wrap;align-items:center;gap:8px;margin:8px 0 0;font-size:13px}
.toolbar{display:flex;flex-wrap:wrap;gap:8px;align-items:center;margin:0 0 8px}
.subsecs{display:grid;grid-template-columns:repeat(auto-fit,minmax(220px,1fr));gap:10px;margin-bottom:12px}
#modal{position:fixed;inset:0;background:rgba(0,0,0,.35);z-index:100;display:none;align-items:flex-start;justify-content:center;padding:24px 12px;overflow:auto}
#modal.open{display:flex}
#modal .sheet{background:#fff;border-radius:12px;max-width:900px;width:100%;box-shadow:0 20px 60px rgba(0,0,0,.25);padding:16px 20px 20px;position:relative}
#modal .close{position:absolute;left:12px;top:10px;border:0;background:none;font-size:22px;cursor:pointer;color:var(--muted)}
#modal h2{font-size:17px;margin:0 0 4px;padding-left:30px}
#modal .amt{font-size:22px;font-weight:700;margin-bottom:8px}
.kv{display:grid;grid-template-columns:1fr 1fr;gap:0 24px}
@media (max-width:700px){.kv{grid-template-columns:1fr}}
.kv div{display:flex;gap:8px;padding:4px 0;border-bottom:1px solid var(--grid);font-size:13px}
.kv div span.k{color:var(--ink2);min-width:150px;flex:0 0 150px}
.kv div span.v{word-break:break-word;white-space:pre-wrap}
.kv div.full{grid-column:1/-1}
.actions{display:flex;flex-wrap:wrap;gap:8px;margin:10px 0}
.pivot td.cell{cursor:pointer}
.pivot td.cell:hover{outline:2px solid var(--accent);outline-offset:-2px}
.pivot th.rowh,.pivot td.rowh{position:sticky;right:0;background:var(--surface);z-index:3}
.cols-pop label{display:flex;gap:8px;align-items:center;padding:3px 6px;font-size:13px;cursor:pointer}
.help p{margin:6px 0}
.help li{margin:3px 0}
.help kbd{background:#eeede8;border-radius:4px;padding:0 5px;font-family:inherit}
</style>
</head>
<body>
<div class="wrap">
<header class="top">
  <h1>דשבורד תזרים — __TITLE__</h1>
  <span class="sub">מקור: <b><bdi dir="ltr">__SRC__</bdi></b> · נבנה __BUILT__</span>
</header>

<nav class="tabs" id="tabs">
  <button data-tab="overview">סקירה</button>
  <button data-tab="explore">עסקאות <span class="n" id="nExplore"></span></button>
  <button data-tab="transfers">העברות, BIT ו-PAYBOX <span class="n" id="nTransfers"></span></button>
  <button data-tab="income">הכנסות <span class="n" id="nIncome"></span></button>
  <button data-tab="unc">לסיווג <span class="n" id="nUnc"></span></button>
  <button data-tab="help">עזרה</button>
</nav>

<section class="tab" id="tab-overview" hidden>
<section class="card filters" id="filters">
  <div class="f"><label>סכימה</label><div class="seg" id="segScheme"><button data-v="tz" class="on">תזרים</button><button data-v="new">מוצעת</button></div></div>
  <div class="f"><label>סוג</label><div class="seg" id="segKind"><button data-v="exp" class="on">הוצאות</button><button data-v="inc">הכנסות</button></div></div>
  <div class="f"><label>קבוצה</label><div class="ms" id="msGroups"></div></div>
  <div class="f"><label>תת-קטגוריה</label><div class="ms" id="msCats"></div></div>
  <div class="f"><label>חודש</label><div class="ms" id="msMonths"></div></div>
  <div class="f"><label>אמצעי תשלום</label><div class="ms" id="msPays"></div></div>
  <div class="f"><label>אדם</label><div class="ms" id="msPersons"></div></div>
  <div class="f" id="fTrip"><label>נסיעה</label><div class="ms" id="msTrips"></div></div>
  <div class="f"><label>חיפוש טקסט</label><input type="text" id="q" placeholder="שם, שם מקורי, פרטים…"></div>
  <div class="f"><label>סכום מ־</label><input type="number" id="amin" placeholder="מינ׳"></div>
  <div class="f"><label>עד</label><input type="number" id="amax" placeholder="מקס׳"></div>
  <div class="f"><label>&nbsp;</label><label class="toggle"><input type="checkbox" id="showAll"> הצג גם מחוץ לחלון / לא נסכם</label></div>
  <div class="f"><label>&nbsp;</label><button class="btn" id="clear">נקה מסננים</button></div>
</section>
<div class="chips" id="chips"></div>
<section class="kpis" id="kpis"></section>
<section class="grid2">
  <div class="card"><h2 id="tMonthly">סכום חודשי</h2><div class="chartbox"><canvas id="cMonthly"></canvas></div></div>
  <div class="card"><h2 id="tStack">פילוח חודשי לפי קבוצה</h2><div class="chartbox"><canvas id="cStack"></canvas></div></div>
</section>
<section class="grid2">
  <div class="card"><h2 id="tLevel">לפי ממוצע חודשי</h2><div class="chartbox tall"><canvas id="cLevel"></canvas></div></div>
  <div class="card"><h2>לפי אמצעי תשלום</h2><div class="chartbox tall"><canvas id="cPay"></canvas></div></div>
</section>
<section class="card" id="drill">
  <div class="crumbs" id="crumbs"></div>
  <p class="hint" id="hint"></p>
  <div class="tw" id="table"></div>
</section>
<p class="small" style="margin:10px 2px 0">הסקירה תואמת לאקסל: הוצאה+הכנסה, נסכם=כן, בחלון=כן. לכל השורות (כולל העברות, כפילויות ומחוץ לחלון) — לשונית <a href="#explore" data-go="explore">עסקאות</a>.</p>
</section>

<section class="tab" id="tab-explore" hidden>
  <details class="howto" open>
    <summary>איך משתמשים (3 שורות)</summary>
    <ol>
      <li><b>סינון:</b> חיפוש חופשי בכל העמודות, צ'יפים לסוג תנועה, חלון/נסכם, טווח תאריכים וסכום, ומסננים לפי עמודה. לחיצה על כותרת = מיון; "עמודות" = בחירת עמודות.</li>
      <li><b>שורה:</b> לחיצה פותחת פרטי עסקה עם כל השדות, השורה המקושרת ורשימת אותו בית עסק. טבלת ציר מתחת לרשימה; "הורד CSV" = השורות המסוננות.</li>
      <li><b>עדכון מהאקסל:</b> שנה בקובץ → <code id="cmd1"></code> → רענן את הדף.</li>
    </ol>
  </details>
  <div id="exMain"></div>
</section>

<section class="tab" id="tab-transfers" hidden>
  <p class="hint">מקביל ללשונית האקסל "העברות, BIT ו-PAYBOX": העברות בעו"ש, תשלומי אפליקציה ו-PAYBOX, בתוך החלון ומחוצה לו. כיוון: נכנס=זיכוי, יוצא=חיוב; הסכומים בערך מוחלט. לחיצה על כרטיס סעיף מסננת אליו.</p>
  <div class="subsecs" id="trSecs"></div>
  <div id="exTransfers"></div>
</section>

<section class="tab" id="tab-income" hidden>
  <section class="card" style="margin-bottom:12px">
    <div class="toolbar"><h2 style="margin:0">הכנסות לפי מקור × חודש</h2>
      <label class="toggle"><input type="checkbox" id="incAll"> כולל מחוץ לחלון / לא נסכם</label>
      <span class="small">מקור = קטגוריה (תזרים). ממוצע = סה"כ ÷ מספר החודשים; ללא אפס = ÷ חודשים עם תנועה.</span></div>
    <div class="tw" id="incTable"></div>
  </section>
  <div id="exIncome"></div>
</section>

<section class="tab" id="tab-unc" hidden>
  <section class="card">
    <h2>לסיווג <span class="count" id="uncCount"></span></h2>
    <div class="note" id="uncNote"></div>
    <div class="tw" id="uncTable"></div>
  </section>
</section>

<section class="tab" id="tab-help" hidden>
  <section class="card help">
    <h2>עזרה</h2>
    <p><b>עסקאות</b> — כל השורות וכל העמודות של לשונית database, מכל הסוגים.</p>
    <ul>
      <li><b>סינון:</b> חיפוש חופשי · צ'יפים לסוג תנועה · חלון (בחלון/מחוץ/הכל) · נסכם (כן/לא/הכל) · טווח תאריך עסקה · סכום מ/עד · מסננים לפי עמודה.</li>
      <li><b>מיון:</b> לחיצה על כותרת עמודה (לחיצה נוספת הופכת את הכיוון). ברירת מחדל: סכום יורד.</li>
      <li><b>עמודות:</b> כפתור <kbd>עמודות</kbd> — סימון/ביטול עמודות, "הצג הכל" / "ברירת מחדל".</li>
      <li><b>פרטי עסקה:</b> לחיצה על שורה — כל השדות, השורה המקושרת ורשימת אותו "שם מובן", וכפתורי סינון מהיר.</li>
      <li><b>טבלת ציר:</b> ממד שורות × ממד עמודות (חודש / אמצעי תשלום / אדם / אין) × מדד (סכום / מס' עסקאות / ממוצע); לחיצה על תא מסננת את הרשימה.</li>
      <li><b>ייצוא:</b> <kbd>הורד CSV</kbd> — השורות המסוננות, כל העמודות, UTF-8 עם BOM.</li>
    </ul>
    <p><b>סקירה</b> — תואמת לאקסל: הוצאה+הכנסה, נסכם=כן, בחלון=כן. פירוט קבוצה → תת-קטגוריה → בית עסק → עסקאות; KPI כולל "ממוצע חודשי ללא אפס".</p>
    <p><b>העברות, BIT ו-PAYBOX</b> — אותו סייר, מסונן כמו לשונית האקסל. <b>הכנסות</b> — מקור × חודש + רשימת שורות ההכנסה. <b>לסיווג</b> — שורות לבדיקה; הסיווג נעשה באקסל (לשונית "לסיווג ידני").</p>
    <p><b>עדכון:</b> שנה באקסל → <code id="cmd2"></code> → רענן את הדף. הגרפים דורשים את Chart.js (מהרשת או מוטמע); הטבלאות עובדות גם בלי.</p>
  </section>
</section>

<footer id="foot"></footer>
</div>
<div id="modal"><div class="sheet" id="modalSheet"></div></div>

__CHARTJS__
<script>
const DATA = __DATA__;
const META = __META__;
const WINDOW = META.window;
const ALL_MONTHS = META.months.slice().sort();
const MONTH_NAMES = META.month_names;
const PALETTE = ['#2a78d6','#eb6834','#1baf7a','#eda100','#e87ba4','#008300','#4a3aa7','#e34948'];
const OTHER_COLOR = '#a3a29c';
const OTHER = 'אחר';
const EXP = META.types.exp, INC = META.types.inc;
const PERSON_ORDER = META.people;
const TYPE_ORDER = [META.types.exp, META.types.inc, META.types.reimb_paid, META.types.reimb_recv, META.types.internal, META.types.savings, META.types.card_debit, META.types.dup];
const SEC_NAME = {1:'העברות בנקאיות', 2:'תשלומי אפליקציה (BIT)', 3:'PAYBOX'};
const SEC_LONG = {1:'1. העברות בנקאיות (עו"ש)', 2:'2. תשלומי אפליקציה — שורות אפליקציה, שורות כרטיס ומשיכות לבנק', 3:'3. PAYBOX'};
const COLS = META.cols.slice();
const COL = {};
COLS.forEach(c => { COL[c.k] = c; });
const SHEET_KEYS = COLS.map(c => c.k).filter(k => !k.startsWith('tr_'));
const nf0 = new Intl.NumberFormat('he-IL', {maximumFractionDigits:0});
const nf1 = new Intl.NumberFormat('he-IL', {maximumFractionDigits:1, minimumFractionDigits:0});
const nf2 = new Intl.NumberFormat('he-IL', {maximumFractionDigits:2, minimumFractionDigits:2});
const fmt = v => (v==null||isNaN(v)) ? '—' : nf0.format(Math.round(v)) + ' ₪';
const fmtS = v => (v==null||isNaN(v)) ? '—' : (Math.abs(v)>=1000 ? nf0.format(Math.round(v)) : nf1.format(v)) + ' ₪';
const fmt2 = v => (v==null||isNaN(v)) ? '' : nf2.format(v);
const pct = v => (v==null||isNaN(v)||!isFinite(v)) ? '—' : nf1.format(v*100) + '%';
const esc = s => String(s==null?'':s).replace(/[&<>"']/g, c => ({'&':'&amp;','<':'&lt;','>':'&gt;','"':'&quot;',"'":'&#39;'}[c]));
const mName = m => MONTH_NAMES[m] || m;
const fmtDate = s => { if(!s) return ''; const p=String(s).split('-'); return p.length===3 ? p[2]+'.'+p[1]+'.'+p[0].slice(2) : s; };
const BY_ID = {};
const MS = {};
const TAB_INIT = {};
const EXPLORERS = {};
const charts = {};
const HAS_CHART = typeof window.Chart !== 'undefined';
const FILTER_COLS = ['source','pay','person','group_tz','cat_tz','group_new','cat_new','month','orig_currency','trip'];
const PIVOT_ROW_COLS = ['name_clean','cat_tz','group_tz','cat_new','group_new','pay','person','source','card','type','month','orig_currency','trip','in_window','summed','source_file','tr_section','tr_dir','tr_payee'];
const PIVOT_COL_COLS = [['','אין'],['month','חודש'],['pay','אמצעי תשלום'],['person','אדם']];
const DEFAULT_COLS = ['txn_date','name_clean','amount','cat_tz','pay','person','month_name','rule_note','details'];
const LEVEL_LABEL = {1:'קבוצה',2:'תת-קטגוריה',3:'בית עסק',4:'עסקאות'};
const state = {
  scheme:'tz', kind:EXP, showAll:false,
  groups:new Set(), cats:new Set(), months:new Set(WINDOW), pays:new Set(), persons:new Set(), trips:new Set(),
  q:'', amin:null, amax:null, drill:{group:null, cat:null, merchant:null},
  sort:{1:{key:'total',dir:-1},2:{key:'total',dir:-1},3:{key:'total',dir:-1},4:{key:'amount',dir:-1}}
};

function cellText(r, k){
  const v = r[k]; const c = COL[k] || {};
  if (v==null || v==='') return (c.t==='bool') ? 'לא' : '';
  if (c.t==='bool') return v ? 'כן' : 'לא';
  if (k==='tr_section') return SEC_NAME[v] || '';
  if (c.t==='num') return fmt2(v);
  if (c.t==='date') return fmtDate(v);
  return String(v);
}
DATA.forEach((r,i) => { r._i = i; BY_ID[r.id] = r; r._s = SHEET_KEYS.map(k => cellText(r,k)).concat([r.tr_payee||'', r.tr_note||'']).join(' | ').toLowerCase(); });
function debounce(fn, ms){ let t; return (...a)=>{ clearTimeout(t); t=setTimeout(()=>fn(...a), ms); }; }
function sumBy(rs, keyFn){ const o={}; rs.forEach(r=>{ const k=keyFn(r); o[k]=(o[k]||0)+(r.amount||0); }); return o; }
function countBy(rows, key){ const m = new Map(); for (const r of rows){ const v=r[key]||''; m.set(v,(m.get(v)||0)+1); } return m; }
function optList(rows, key, order){
  const m = countBy(rows, key);
  const arr = Array.from(m.entries()).map(([v,n])=>({v,n}));
  if (order) arr.sort((a,b)=> (order.indexOf(a.v)===-1?99:order.indexOf(a.v)) - (order.indexOf(b.v)===-1?99:order.indexOf(b.v)) || String(a.v).localeCompare(String(b.v),'he'));
  else arr.sort((a,b)=> b.n-a.n || String(a.v).localeCompare(String(b.v),'he'));
  return arr;
}

/* ---------- multi-select ---------- */
function closeAll(){ document.querySelectorAll('.ms.open').forEach(e=>e.classList.remove('open')); }
document.addEventListener('click', closeAll);
document.addEventListener('keydown', e => { if(e.key==='Escape'){ closeAll(); closeModal(); } });
function makeMultiEl(el, set, optsFn, opts){
  opts = opts || {};
  const btn = document.createElement('button'); btn.type='button';
  const pop = document.createElement('div'); pop.className='pop';
  el.classList.add('ms'); el.innerHTML=''; el.appendChild(btn); el.appendChild(pop);
  const comp = {el, btn, pop, set, optsFn, options:[], onChange:opts.onChange||(()=>{}), labelFn:opts.labelFn||(v=>v||'(ריק)')};
  btn.addEventListener('click', e => { e.stopPropagation(); const open = el.classList.contains('open'); closeAll(); if(!open){ comp.render(); el.classList.add('open'); const s=pop.querySelector('input.s'); if(s) s.focus(); } });
  pop.addEventListener('click', e => e.stopPropagation());
  comp.render = function(filterText){
    comp.options = optsFn();
    const avail = new Set(comp.options.map(o=>o.v));
    for (const v of Array.from(comp.set)) if (!avail.has(v)) comp.set.delete(v);
    const ft = (filterText||'').toLowerCase();
    const shown = ft ? comp.options.filter(o => comp.labelFn(o.v).toLowerCase().includes(ft)) : comp.options;
    pop.innerHTML = `<div class="tools"><button data-a="all">בחר הכל</button><button data-a="none">נקה</button>${comp.options.length>8?'<input class="s" placeholder="חפש…">':''}</div>` +
      shown.map(o => `<label><input type="checkbox" value="${esc(o.v)}" ${comp.set.has(o.v)?'checked':''}> <span class="t">${esc(comp.labelFn(o.v))}</span><span class="n">${o.n}</span></label>`).join('') +
      (shown.length? '' : '<div class="small" style="padding:6px">אין ערכים</div>');
    pop.querySelectorAll('input[type=checkbox]').forEach(inp => inp.addEventListener('change', () => { if(inp.checked) comp.set.add(inp.value); else comp.set.delete(inp.value); comp.sync(); comp.onChange(); }));
    pop.querySelector('[data-a=all]').addEventListener('click', () => { shown.forEach(o=>comp.set.add(o.v)); comp.render(ft); comp.onChange(); });
    pop.querySelector('[data-a=none]').addEventListener('click', () => { comp.set.clear(); comp.render(ft); comp.onChange(); });
    const s = pop.querySelector('input.s');
    if (s){ s.value = filterText||''; s.addEventListener('input', () => { comp.render(s.value); const s2=pop.querySelector('input.s'); if(s2){ s2.focus(); s2.setSelectionRange(s2.value.length, s2.value.length);} }); }
    comp.sync();
  };
  comp.sync = function(){
    const n = comp.set.size, total = comp.options.length;
    let t;
    if (n===0) t = opts.emptyMeansAll===false ? 'לא נבחר' : 'הכל';
    else if (n===total && opts.emptyMeansAll===false) t = 'הכל ('+n+')';
    else if (n<=2) t = Array.from(comp.set).map(comp.labelFn).join(', ');
    else t = n + ' נבחרו';
    btn.textContent = t;
    btn.classList.toggle('act', n>0 && !(opts.emptyMeansAll===false && n===total));
  };
  comp.render();
  return comp;
}
function makeMulti(id, set, optsFn, opts){ const c = makeMultiEl(document.getElementById(id), set, optsFn, opts); MS[id]=c; return c; }

/* ---------- tabs ---------- */
function showTab(name){
  if (!META.tabs.includes(name)) name = 'overview';
  document.querySelectorAll('#tabs button').forEach(b=>b.classList.toggle('on', b.dataset.tab===name));
  document.querySelectorAll('.tab').forEach(t=>{ t.hidden = (t.id !== 'tab-'+name); });
  if (name==='overview') refresh();
  if (TAB_INIT[name]) TAB_INIT[name]();
  try { history.replaceState(null, '', '#'+name); } catch(e){}
}

/* ---------- overview (Excel-consistent) ---------- */
function gKey(){ return state.scheme==='tz' ? 'group_tz' : 'group_new'; }
function cKey(){ return state.scheme==='tz' ? 'cat_tz' : 'cat_new'; }
function level(){ return state.drill.merchant ? 4 : state.drill.cat ? 3 : state.drill.group ? 2 : 1; }
function levelKey(){ return ({1:gKey(),2:cKey(),3:'name_clean'})[Math.min(level(),3)]; }
function baseRows(){ return DATA.filter(r => (r.type===EXP || r.type===INC) && (state.showAll || (r.in_window && r.summed))); }
function kindBase(){ return baseRows().filter(r => r.type===state.kind); }
function kindRows(){ return kindBase().filter(r => state.months.has(r.month)); }
function matchesFilters(r){
  const gk=gKey(), ck=cKey();
  if (state.groups.size && !state.groups.has(r[gk])) return false;
  if (state.cats.size && !state.cats.has(r[ck])) return false;
  if (state.pays.size && !state.pays.has(r.pay)) return false;
  if (state.persons.size && !state.persons.has(r.person)) return false;
  if (state.trips.size && !state.trips.has(r.trip||'')) return false;
  if (state.amin!=null && !(r.amount>=state.amin)) return false;
  if (state.amax!=null && !(r.amount<=state.amax)) return false;
  if (state.q){ const q=state.q; if (!((r.name_clean||'').toLowerCase().includes(q) || (r.original_name||'').toLowerCase().includes(q) || (r.details||'').toLowerCase().includes(q))) return false; }
  return true;
}
function filteredRows(){ return kindRows().filter(matchesFilters); }
function drillRows(rows){ const d=state.drill, gk=gKey(), ck=cKey(); return rows.filter(r => (!d.group || r[gk]===d.group) && (!d.cat || r[ck]===d.cat) && (!d.merchant || r.name_clean===d.merchant)); }
function selMonths(){ return ALL_MONTHS.filter(m => state.months.has(m)); }
function aggregate(rows, key, months){
  months = months || selMonths();
  const map = new Map();
  for (const r of rows){
    const k = r[key] || '(ריק)';
    let a = map.get(k);
    if (!a){ a={name:k,total:0,count:0,months:{}}; map.set(k,a); }
    a.total += r.amount||0; a.count++;
    a.months[r.month] = (a.months[r.month]||0) + (r.amount||0);
  }
  const nM = Math.max(months.length,1);
  const all = Array.from(map.values());
  const sumAll = all.reduce((s,a)=>s+a.total,0);
  for (const a of all){
    a.avg = a.total/nM;
    const nz = months.filter(m => Math.abs(a.months[m]||0) > 0.005).length;
    a.avgNZ = nz ? a.total/nz : 0; a.nz = nz;
    a.pct = sumAll ? a.total/sumAll : 0;
    a.avgTxn = a.count ? a.total/a.count : 0;
  }
  return all;
}
function groupOpts(){ return optList(kindBase().filter(r=>state.months.has(r.month)), gKey()); }
function catOpts(){ const gk=gKey(); return optList(kindBase().filter(r=>state.months.has(r.month) && (!state.groups.size || state.groups.has(r[gk]))), cKey()); }
function monthOpts(){ const ms = state.showAll ? ALL_MONTHS : WINDOW; return ms.map(m => ({v:m, n:kindBase().filter(r=>r.month===m).length})); }
function payOpts(){ return optList(kindBase(), 'pay'); }
function personOpts(){ return optList(kindBase(), 'person', PERSON_ORDER); }
function tripOpts(){ return optList(kindBase().filter(r=>r.trip), 'trip'); }
function setSeg(id, v){ document.querySelectorAll('#'+id+' button').forEach(b=>b.classList.toggle('on', b.dataset.v===v)); }
function kindCode(){ return state.kind===EXP ? 'exp' : 'inc'; }
function resetDrill(){ state.drill={group:null,cat:null,merchant:null}; }

function renderChips(){
  const c = document.getElementById('chips'); const items=[];
  const chip = (label, val, cls, fn) => items.push({label,val,cls,fn});
  if (state.scheme!=='tz') chip('סכימה','מוצעת','',()=>{state.scheme='tz'; setSeg('segScheme','tz'); resetDrill();});
  if (state.kind!==EXP) chip('סוג','הכנסות','',()=>{state.kind=EXP; setSeg('segKind','exp'); resetDrill();});
  if (state.showAll) chip('טווח','כולל מחוץ לחלון / לא נסכם','',()=>{state.showAll=false; document.getElementById('showAll').checked=false; onToggle();});
  const ms = selMonths();
  if (ms.length !== (state.showAll?ALL_MONTHS:WINDOW).length || ms.length===0) chip('חודשים', ms.length? ms.map(mName).join(', ') : 'אין', '', ()=>{state.months=new Set(state.showAll?ALL_MONTHS:WINDOW); MS.msMonths.set=state.months; MS.msMonths.render();});
  for (const [id,set,label] of [['msGroups',state.groups,'קבוצה'],['msCats',state.cats,'תת-קטגוריה'],['msPays',state.pays,'אמצעי תשלום'],['msPersons',state.persons,'אדם'],['msTrips',state.trips,'נסיעה']])
    if (set.size) chip(label, Array.from(set).join(', '), '', ()=>{ set.clear(); if(MS[id]) MS[id].render(); });
  if (state.q) chip('חיפוש', state.q, '', ()=>{state.q=''; document.getElementById('q').value='';});
  if (state.amin!=null) chip('סכום מ־', fmt(state.amin), '', ()=>{state.amin=null; document.getElementById('amin').value='';});
  if (state.amax!=null) chip('סכום עד', fmt(state.amax), '', ()=>{state.amax=null; document.getElementById('amax').value='';});
  const d=state.drill;
  if (d.group) chip('קבוצה ▸', d.group, 'drill', ()=>{resetDrill();});
  if (d.cat) chip('תת-קטגוריה ▸', d.cat, 'drill', ()=>{d.cat=null; d.merchant=null;});
  if (d.merchant) chip('בית עסק ▸', d.merchant, 'drill', ()=>{d.merchant=null;});
  c.innerHTML = items.map((it,i)=>`<span class="chip ${it.cls}"><span>${esc(it.label)}: <b>${esc(it.val)}</b></span><button title="הסר" data-i="${i}">×</button></span>`).join('');
  c.querySelectorAll('button').forEach(b => b.addEventListener('click', ()=>{ items[+b.dataset.i].fn(); refresh(); }));
}
function renderKPIs(rows, denomRows){
  const months = selMonths(), nM = Math.max(months.length,1);
  const total = rows.reduce((s,r)=>s+(r.amount||0),0);
  const byM = {}; rows.forEach(r=>{byM[r.month]=(byM[r.month]||0)+(r.amount||0);});
  const nz = months.filter(m=>Math.abs(byM[m]||0)>0.005).length;
  const denom = denomRows.reduce((s,r)=>s+(r.amount||0),0);
  const tiles = [
    ['סה"כ', fmt(total), months.length+' חודשים נבחרים'],
    ['ממוצע חודשי', fmt(total/nM), '÷ '+nM],
    ['ממוצע חודשי ללא אפס', fmt(nz?total/nz:0), '÷ '+nz+' חודשים עם תנועה'],
    ["מס' עסקאות", nf0.format(rows.length), ''],
    ['ממוצע לעסקה', fmtS(rows.length?total/rows.length:0), ''],
    [state.kind===EXP?'% מכלל ההוצאות':'% מכלל ההכנסות', pct(denom?total/denom:0), 'מתוך '+fmt(denom)]
  ];
  document.getElementById('kpis').innerHTML = tiles.map(t=>`<div class="kpi"><div class="l">${t[0]}</div><div class="v">${t[1]}</div><div class="h">${t[2]}</div></div>`).join('');
}
function renderCrumbs(){
  const d=state.drill, parts=[];
  const root = state.kind===EXP?'כל ההוצאות':'כל ההכנסות';
  parts.push(d.group ? `<a data-l="1">${root}</a>` : `<span class="cur">${root} — לפי קבוצה</span>`);
  if (d.group) parts.push('<span class="sep">›</span>', d.cat ? `<a data-l="2">${esc(d.group)}</a>` : `<span class="cur">${esc(d.group)} — לפי תת-קטגוריה</span>`);
  if (d.cat) parts.push('<span class="sep">›</span>', d.merchant ? `<a data-l="3">${esc(d.cat)}</a>` : `<span class="cur">${esc(d.cat)} — לפי בית עסק</span>`);
  if (d.merchant) parts.push('<span class="sep">›</span>', `<span class="cur">${esc(d.merchant)} — רשימת עסקאות</span>`);
  const el = document.getElementById('crumbs');
  el.innerHTML = parts.join('');
  el.querySelectorAll('a').forEach(a => a.addEventListener('click', ()=>{ const l=+a.dataset.l; if(l<=1){d.group=null;} if(l<=2){d.cat=null;} if(l<=3){d.merchant=null;} refresh(); }));
  const hint = {1:'לחץ על שורה כדי לראות את תת-הקטגוריות של הקבוצה.',2:'לחץ על שורה כדי לראות את בתי העסק בתת-הקטגוריה.',3:'לחץ על שורה כדי לראות את רשימת העסקאות.',4:'לחץ על עסקה לפרטים מלאים.'}[level()];
  document.getElementById('hint').textContent = hint + ' לחיצה על כותרת עמודה ממיינת.';
}
function heatStyle(v, max){
  if (!(v>0) || !(max>0)) return '';
  const t = Math.min(1, v/max); const c1=[255,255,255], c2=[253,186,116];
  const c = c1.map((a,i)=>Math.round(a+(c2[i]-a)*t));
  return `background:rgb(${c[0]},${c[1]},${c[2]})`;
}
function sortArrow(lvl, key){ const s=state.sort[lvl]; return s.key===key ? `<span class="arr">${s.dir<0?'▼':'▲'}</span>` : ''; }
function headerClick(lvl, key, numericDefaultDesc){ const s=state.sort[lvl]; if (s.key===key) s.dir=-s.dir; else { s.key=key; s.dir = numericDefaultDesc ? -1 : 1; } }
function renderAggTable(rows){
  const lvl = level(), key = levelKey(), months = selMonths();
  const agg = aggregate(rows, key);
  const s = state.sort[lvl];
  agg.sort((a,b)=>{
    const va = s.key.startsWith('m:') ? (a.months[s.key.slice(2)]||0) : a[s.key];
    const vb = s.key.startsWith('m:') ? (b.months[s.key.slice(2)]||0) : b[s.key];
    if (typeof va==='string') return va.localeCompare(vb,'he')*s.dir;
    return ((va||0)-(vb||0))*s.dir;
  });
  const heatMax = Math.max(0, ...agg.flatMap(a => months.map(m => a.months[m]||0)));
  const cols = [['name',LEVEL_LABEL[lvl],false],['total','סה"כ',true],['avg','ממוצע חודשי',true],['avgNZ','ממוצע ללא אפס',true],['count',"מס' עסקאות",true],['pct','%',true]]
    .concat(months.map(m=>['m:'+m, mName(m), true]));
  const th = cols.map(c=>`<th class="${c[2]?'num':''}" data-k="${c[0]}" data-n="${c[2]?1:0}">${sortArrow(lvl,c[0])}${esc(c[1])}</th>`).join('');
  const tot = {total:0,count:0,months:{}};
  const body = agg.map(a=>{
    tot.total+=a.total; tot.count+=a.count; months.forEach(m=>{tot.months[m]=(tot.months[m]||0)+(a.months[m]||0);});
    return `<tr class="click" data-v="${esc(a.name)}"><td class="name" title="${esc(a.name)}">${esc(a.name)}</td><td class="num${a.total<0?' neg':''}">${fmt(a.total)}</td><td class="num">${fmt(a.avg)}</td><td class="num">${fmt(a.avgNZ)}</td><td class="num">${a.count}</td><td class="num">${pct(a.pct)}</td>` +
      months.map(m=>{const v=a.months[m]||0; return `<td class="num heat${v<0?' neg':''}" style="${heatStyle(v,heatMax)}">${v?fmt(v):'·'}</td>`;}).join('') + '</tr>';
  }).join('');
  const nM=Math.max(months.length,1), nzT = months.filter(m=>Math.abs(tot.months[m]||0)>0.005).length;
  const foot = `<tr><td>סה"כ (${agg.length})</td><td class="num">${fmt(tot.total)}</td><td class="num">${fmt(tot.total/nM)}</td><td class="num">${fmt(nzT?tot.total/nzT:0)}</td><td class="num">${tot.count}</td><td class="num">100%</td>` + months.map(m=>`<td class="num">${fmt(tot.months[m]||0)}</td>`).join('') + '</tr>';
  const el = document.getElementById('table');
  el.innerHTML = `<table><thead><tr>${th}</tr></thead><tbody>${body || '<tr><td colspan="'+cols.length+'" class="small">אין נתונים בבחירה הנוכחית</td></tr>'}</tbody><tfoot>${foot}</tfoot></table>`;
  el.querySelectorAll('th').forEach(h => h.addEventListener('click', ()=>{ headerClick(lvl, h.dataset.k, h.dataset.n==='1'); refresh(); }));
  el.querySelectorAll('tbody tr.click').forEach(tr => tr.addEventListener('click', ()=>{
    const v = tr.dataset.v;
    if (lvl===1) state.drill.group=v; else if (lvl===2) state.drill.cat=v; else state.drill.merchant=v;
    refresh();
    document.getElementById('drill').scrollIntoView({behavior:'smooth',block:'start'});
  }));
}
function renderTxnTable(rows){
  const s = state.sort[4], ck=cKey();
  const arr = rows.slice().sort((a,b)=>{ const va=a[s.key], vb=b[s.key]; if (typeof va==='number' || typeof vb==='number') return ((va||0)-(vb||0))*s.dir; return String(va||'').localeCompare(String(vb||''),'he')*s.dir; });
  const cols = [['txn_date','תאריך',false],['name_clean','שם',false],['amount','סכום',true],[ck,'קטגוריה',false],['pay','אמצעי תשלום',false],['person','אדם',false],['details','פרטים',false]];
  const th = cols.map(c=>`<th class="${c[2]?'num':''}" data-k="${c[0]}" data-n="${c[2]?1:0}">${sortArrow(4,c[0])}${esc(c[1])}</th>`).join('');
  const body = arr.map(r=>`<tr class="click" data-id="${esc(r.id)}"><td>${fmtDate(r.txn_date)}${r.charge_date&&r.charge_date!==r.txn_date?` <span class="small" title="תאריך חיוב">(${fmtDate(r.charge_date)})</span>`:''}</td><td class="name" title="${esc(r.original_name)}">${esc(r.name_clean)}</td><td class="num${r.amount<0?' neg':''}">${fmt2(r.amount)} ₪</td><td>${esc(r[ck])}</td><td>${esc(r.pay)}</td><td>${esc(r.person)}</td><td class="txt" title="${esc(r.details)}">${esc(r.details)}</td></tr>`).join('');
  const total = arr.reduce((s2,r)=>s2+(r.amount||0),0);
  const foot = `<tr><td colspan="2">סה"כ (${arr.length} עסקאות)</td><td class="num">${fmt(total)}</td><td colspan="4"></td></tr>`;
  const el = document.getElementById('table');
  el.innerHTML = `<table><thead><tr>${th}</tr></thead><tbody>${body}</tbody><tfoot>${foot}</tfoot></table>`;
  el.querySelectorAll('th').forEach(h => h.addEventListener('click', ()=>{ headerClick(4, h.dataset.k, h.dataset.n==='1'); refresh(); }));
  el.querySelectorAll('tbody tr.click').forEach(tr => tr.addEventListener('click', ()=> openRow(BY_ID[tr.dataset.id], null)));
}

/* ---------- charts ---------- */
/* Inline Chart.js plugin (no external dependency; works in cdn and inline modes): values above bars
   (horizontal bars: right of the bar end; stacked bars: stack total + per-segment when tall enough) and
   percentages inside pie/doughnut slices (slices < minPct are unlabeled; amount stays in the tooltip).
   Per chart: options.plugins.valueLabels = {enabled:false} | {segments:false} | {minPct:0.04}. */
const VALUE_LABELS = {
  id: 'valueLabels',
  afterDatasetsDraw(chart, _args, opts){
    const o = Object.assign({enabled:true, segments:true, minPct:0.04, fontSize:12, color:'#3a3935'}, opts||{});
    if (!o.enabled) return;
    const ctx = chart.ctx, horiz = chart.options.indexAxis==='y', dss = chart.data.datasets;
    const sc = chart.options.scales||{}, stacked = !!((sc.y&&sc.y.stacked)||(sc.x&&sc.x.stacked)||dss.some(d=>d.stack));
    const num = v => nf0.format(Math.round(v));
    ctx.save();
    ctx.font = '600 '+o.fontSize+'px '+(Chart.defaults.font.family||'sans-serif');
    ctx.fillStyle = o.color; ctx.strokeStyle = 'rgba(255,255,255,.92)'; ctx.lineWidth = 3; ctx.lineJoin = 'round';
    const put = (t,x,y,align,base) => { ctx.textAlign=align; ctx.textBaseline=base; ctx.strokeText(t,x,y); ctx.fillText(t,x,y); };
    const stackTop = {}, stackSum = {};
    dss.forEach((ds,i)=>{
      if (!chart.isDatasetVisible(i)) return;
      const meta = chart.getDatasetMeta(i);
      if (meta.type==='doughnut' || meta.type==='pie'){
        const total = meta.data.reduce((s,el,j)=> s + (chart.getDataVisibility(j) ? Math.abs(ds.data[j]||0) : 0), 0);
        meta.data.forEach((el,j)=>{
          if (!total || !chart.getDataVisibility(j)) return;
          const p = Math.abs(ds.data[j]||0)/total; if (p < o.minPct) return;
          const pos = el.tooltipPosition(); put(Math.round(p*100)+'%', pos.x, pos.y, 'center', 'middle');
        });
        return;
      }
      if (meta.type!=='bar') return;
      meta.data.forEach((el,j)=>{
        const v = ds.data[j]; if (v==null || isNaN(v) || v===0) return;
        if (stacked){
          stackSum[j] = (stackSum[j]||0) + v;
          const top = v>=0 ? Math.min(el.y, el.base) : Math.max(el.y, el.base);
          const t = stackTop[j]; stackTop[j] = t ? {x:el.x, y: v>=0 ? Math.min(t.y, top) : Math.max(t.y, top)} : {x:el.x, y:top};
          if (o.segments && Math.abs(el.y-el.base) >= o.fontSize+6) put(num(v), el.x, (el.y+el.base)/2, 'center', 'middle');
          return;
        }
        if (horiz){ if (v>=0) put(num(v), el.x+4, el.y, 'left', 'middle'); else put(num(v), el.x-4, el.y, 'right', 'middle'); }
        else { if (v>=0) put(num(v), el.x, el.y-4, 'center', 'bottom'); else put(num(v), el.x, el.y+4, 'center', 'top'); }
      });
    });
    if (stacked && !horiz) Object.keys(stackTop).forEach(j=>{ const t=stackTop[j], s=stackSum[j]; if (!s) return; if (s>=0) put(num(s), t.x, t.y-4, 'center', 'bottom'); else put(num(s), t.x, t.y+4, 'center', 'top'); });
    ctx.restore();
  }
};
function chartCommon(){
  return { responsive:true, maintainAspectRatio:false, layout:{padding:{top:18}},
    plugins:{ valueLabels:{enabled:true}, legend:{rtl:true, textDirection:'rtl', labels:{boxWidth:10, boxHeight:10, usePointStyle:true}},
              tooltip:{rtl:true, textDirection:'rtl', callbacks:{label:(c)=> (c.dataset.label?c.dataset.label+': ':'') + fmt(c.chart.config.options.indexAxis==='y' ? c.parsed.x : (c.parsed.y!==undefined ? c.parsed.y : c.parsed)) }} },
    scales:{} };
}
function makeChart(id, cfg){
  if (!HAS_CHART) return;
  const cv = document.getElementById(id);
  if (charts[id]) { charts[id].destroy(); delete charts[id]; }
  cfg.plugins = (cfg.plugins||[]).concat(VALUE_LABELS);
  try { charts[id] = new Chart(cv.getContext('2d'), cfg); } catch(e){ console.error('chart error', id, e); }
}
function axisMoney(){ return { grace:'8%', grid:{color:'#e1e0d9'}, border:{display:false}, ticks:{color:'#898781', callback:v=>nf0.format(v)} }; }
function axisCat(){ return { grid:{display:false}, border:{color:'#c3c2b7'}, ticks:{color:'#52514e'} }; }
function renderCharts(rows){
  if (!HAS_CHART){ document.querySelectorAll('#tab-overview .chartbox').forEach(b=>{ b.innerHTML='<div class="nochart">Chart.js לא נטען (אין חיבור לאינטרנט) — הטבלאות עובדות כרגיל.</div>'; }); return; }
  const months = selMonths(), labels = months.map(mName);
  const lvl = level(), isExp = state.kind===EXP;
  const noCatFilter = !state.groups.size && !state.cats.size && !state.drill.group;
  const byM = sumBy(rows, r=>r.month);
  const ds = [{type:'bar', label: isExp?'הוצאות':'הכנסות', data: months.map(m=>Math.round(byM[m]||0)), backgroundColor: isExp?PALETTE[0]:PALETTE[2], borderRadius:4, borderSkipped:'start', maxBarThickness:48}];
  let title = 'סכום חודשי — הבחירה הנוכחית';
  if (isExp && noCatFilter){
    const inc = baseRows().filter(r=>r.type===INC && state.months.has(r.month));
    const byI = sumBy(inc, r=>r.month);
    ds.push({type:'bar', label:'הכנסות (כל ההכנסות)', data: months.map(m=>Math.round(byI[m]||0)), backgroundColor:PALETTE[2], borderRadius:4, borderSkipped:'start', maxBarThickness:48});
    title = 'הוצאות מול הכנסות לפי חודש';
  }
  document.getElementById('tMonthly').textContent = title;
  const c1 = chartCommon(); c1.scales = {x:axisCat(), y:Object.assign(axisMoney(),{beginAtZero:true})}; c1.plugins.legend.display = ds.length>1;
  makeChart('cMonthly', {data:{labels, datasets:ds}, options:c1});
  const key = levelKey();
  const agg = aggregate(rows, key).sort((a,b)=>b.avg-a.avg).slice(0,20);
  document.getElementById('tLevel').textContent = LEVEL_LABEL[Math.min(lvl,3)] + ' — לפי ממוצע חודשי (20 המובילים)';
  const c2 = chartCommon(); c2.indexAxis='y'; c2.scales={x:Object.assign(axisMoney(),{beginAtZero:true, position:'top'}), y:Object.assign(axisCat(),{ticks:{color:'#0b0b0b', autoSkip:false, font:{size:12}}})};
  c2.layout.padding = {top:4, right:48}; c2.plugins.legend.display=false; c2.plugins.tooltip.callbacks.label = c => 'ממוצע חודשי: '+fmt(c.parsed.x);
  makeChart('cLevel', {type:'bar', data:{labels:agg.map(a=>a.name), datasets:[{data:agg.map(a=>Math.round(a.avg)), backgroundColor:PALETTE[0], borderRadius:4, borderSkipped:'start', maxBarThickness:18}]}, options:c2});
  const byP = sumBy(rows, r=>r.pay||'(ריק)');
  const pays = Object.entries(byP).filter(e=>e[1]>0).sort((a,b)=>b[1]-a[1]);
  const c3 = chartCommon(); delete c3.scales; c3.layout.padding = {top:4}; c3.cutout='58%'; c3.plugins.legend.position='bottom';
  const pTotal = pays.reduce((s,e)=>s+e[1],0);
  c3.plugins.tooltip.callbacks.label = c => c.label+': '+fmt(c.parsed)+' ('+pct(pTotal?c.parsed/pTotal:0)+')';
  makeChart('cPay', {type:'doughnut', data:{labels:pays.map(e=>e[0]), datasets:[{data:pays.map(e=>Math.round(e[1])), backgroundColor:pays.map((e,i)=> i<8?PALETTE[i]:OTHER_COLOR), borderColor:'#fcfcfb', borderWidth:2}]}, options:c3});
  const sKey = lvl===1 ? gKey() : (lvl===2 ? cKey() : 'name_clean');
  const sName = {1:'קבוצה',2:'תת-קטגוריה',3:'בית עסק',4:'בית עסק'}[lvl];
  document.getElementById('tStack').textContent = 'פילוח חודשי לפי '+sName+' (8 המובילים + אחר)';
  const totals = sumBy(rows, r=>r[sKey]||'(ריק)');
  const top = Object.entries(totals).sort((a,b)=>b[1]-a[1]).slice(0,8).map(e=>e[0]);
  const series = {}; top.forEach(t=>series[t]=months.map(()=>0)); series[OTHER]=months.map(()=>0);
  const mi = {}; months.forEach((m,i)=>mi[m]=i);
  let hasOther=false;
  rows.forEach(r=>{ const k=r[sKey]||'(ריק)'; const i=mi[r.month]; if(i==null) return; if(series[k]) series[k][i]+=r.amount||0; else { series[OTHER][i]+=r.amount||0; hasOther=true; } });
  const sds = top.map((t,i)=>({label:t, data:series[t].map(Math.round), backgroundColor:PALETTE[i], stack:'s', borderColor:'#fcfcfb', borderWidth:{top:2,bottom:0,left:0,right:0}, maxBarThickness:48}));
  if (hasOther) sds.push({label:OTHER, data:series[OTHER].map(Math.round), backgroundColor:OTHER_COLOR, stack:'s', borderColor:'#fcfcfb', borderWidth:{top:2,bottom:0,left:0,right:0}, maxBarThickness:48});
  const c4 = chartCommon(); c4.scales={x:Object.assign(axisCat(),{stacked:true}), y:Object.assign(axisMoney(),{stacked:true, beginAtZero:true})}; c4.plugins.tooltip.mode='index'; c4.plugins.valueLabels = {enabled:true, segments:false};
  makeChart('cStack', {type:'bar', data:{labels, datasets:sds}, options:c4});
}
function onToggle(){
  if (!state.showAll) state.months = new Set(Array.from(state.months).filter(m=>WINDOW.includes(m)));
  if (state.months.size===0) state.months = new Set(state.showAll?ALL_MONTHS:WINDOW);
  MS.msMonths.set = state.months; MS.msMonths.render();
}
function refresh(){
  const rows = filteredRows();
  const drows = drillRows(rows);
  renderChips(); renderKPIs(drows, kindRows()); renderCrumbs();
  if (level()===4) renderTxnTable(drows); else renderAggTable(drows);
  renderCharts(drows);
}
function refreshOptions(){ ['msGroups','msCats','msMonths','msPays','msPersons','msTrips'].forEach(id=>{ if(MS[id]) MS[id].render(); }); }
function initOverview(){
  makeMulti('msGroups', state.groups, groupOpts, {onChange:()=>{ MS.msCats.render(); refresh(); }});
  makeMulti('msCats', state.cats, catOpts, {onChange:refresh});
  makeMulti('msMonths', state.months, monthOpts, {onChange:()=>{ refreshOptions(); refresh(); }, emptyMeansAll:false, labelFn:mName});
  const mm = MS.msMonths;
  mm.sync = function(){ const n=mm.set.size, tot=mm.options.length; const arr=ALL_MONTHS.filter(m=>mm.set.has(m)).map(mName); mm.btn.textContent = n===0?'לא נבחר': (n===tot?'הכל ('+n+')': (n<=3?arr.join(', '):n+' נבחרו')); mm.btn.classList.toggle('act', n>0 && n!==tot); };
  mm.sync();
  makeMulti('msPays', state.pays, payOpts, {onChange:refresh});
  makeMulti('msPersons', state.persons, personOpts, {onChange:refresh});
  if (META.has_trip) makeMulti('msTrips', state.trips, tripOpts, {onChange:refresh}); else document.getElementById('fTrip').style.display='none';
  document.querySelectorAll('#segScheme button').forEach(b=>b.addEventListener('click',()=>{ state.scheme=b.dataset.v; setSeg('segScheme',b.dataset.v); state.groups.clear(); state.cats.clear(); resetDrill(); refreshOptions(); refresh(); }));
  document.querySelectorAll('#segKind button').forEach(b=>b.addEventListener('click',()=>{ state.kind = b.dataset.v==='exp' ? EXP : INC; setSeg('segKind',b.dataset.v); state.groups.clear(); state.cats.clear(); resetDrill(); refreshOptions(); refresh(); }));
  document.getElementById('showAll').addEventListener('change', e=>{ state.showAll=e.target.checked; onToggle(); refreshOptions(); refresh(); });
  document.getElementById('q').addEventListener('input', debounce(e=>{ state.q=e.target.value.trim().toLowerCase(); refresh(); },200));
  document.getElementById('amin').addEventListener('input', debounce(e=>{ state.amin = e.target.value===''?null:+e.target.value; refresh(); },250));
  document.getElementById('amax').addEventListener('input', debounce(e=>{ state.amax = e.target.value===''?null:+e.target.value; refresh(); },250));
  document.getElementById('clear').addEventListener('click', ()=>{
    state.scheme='tz'; state.kind=EXP; state.showAll=false; setSeg('segScheme','tz'); setSeg('segKind','exp'); document.getElementById('showAll').checked=false;
    ['groups','cats','pays','persons','trips'].forEach(k=>state[k].clear());
    state.months=new Set(WINDOW); MS.msMonths.set=state.months;
    state.q=''; state.amin=null; state.amax=null; ['q','amin','amax'].forEach(id=>document.getElementById(id).value='');
    resetDrill(); refreshOptions(); refresh();
  });
}

/* ---------- Explorer (row level) ---------- */
function Explorer(rootId, opts){
  const self = this;
  this.id = rootId; this.opts = opts; EXPLORERS[rootId] = this;
  this.root = document.getElementById(rootId);
  this.base = opts.baseRows;
  this.st = { q:'', types:new Set(opts.types||[]), win:'all', summed:'all', from:'', to:'', amin:null, amax:null,
              cf:{}, merchant:null, cols:new Set(opts.cols||DEFAULT_COLS), sort:{key:'amount',dir:-1}, page:0, pageSize:100,
              pivot:{row:opts.pivotRow||'cat_tz', col:'month', m:'sum'} };
  FILTER_COLS.concat(['type','tr_section','tr_dir','tr_payee','name_clean']).forEach(k=>{ self.st.cf[k]=new Set(); });
  this.ms = {};
  this.build();
}
Explorer.prototype.filtered = function(){
  const st = this.st, rows = this.base();
  const cfKeys = Object.keys(st.cf).filter(k=>st.cf[k].size);
  return rows.filter(r => {
    if (st.types.size && !st.types.has(r.type)) return false;
    if (st.win==='in' && !r.in_window) return false;
    if (st.win==='out' && r.in_window) return false;
    if (st.summed==='yes' && !r.summed) return false;
    if (st.summed==='no' && r.summed) return false;
    if (st.from && (r.txn_date||'') < st.from) return false;
    if (st.to && (r.txn_date||'') > st.to) return false;
    if (st.amin!=null && !(r.amount>=st.amin)) return false;
    if (st.amax!=null && !(r.amount<=st.amax)) return false;
    for (const k of cfKeys){ const v = (k==='tr_section') ? String(r[k]||'') : (r[k]||''); if (!st.cf[k].has(v)) return false; }
    if (st.merchant && r.name_clean!==st.merchant) return false;
    if (st.q && !r._s.includes(st.q)) return false;
    return true;
  });
};
Explorer.prototype.build = function(){
  const self = this, st = this.st, o = this.opts;
  const showTypes = o.showTypes !== false;
  this.root.innerHTML = `
  <section class="card" style="margin-bottom:10px">
    <div class="filters">
      <div class="f"><label>חיפוש חופשי (כל העמודות)</label><input type="text" class="wide q" placeholder="שם, שם מקורי, פרטים, הערה, מזהה…"></div>
      ${showTypes ? '<div class="f"><label>סוג תנועה</label><div class="tchips types"></div></div>' : ''}
      <div class="f"><label>חלון</label><div class="seg win"><button data-v="all" class="on">הכל</button><button data-v="in">בחלון</button><button data-v="out">מחוץ</button></div></div>
      <div class="f"><label>נסכם</label><div class="seg summed"><button data-v="all" class="on">הכל</button><button data-v="yes">כן</button><button data-v="no">לא</button></div></div>
      <div class="f"><label>תאריך עסקה מ־</label><input type="date" class="from"></div>
      <div class="f"><label>עד</label><input type="date" class="to"></div>
      <div class="f"><label>סכום מ־</label><input type="number" class="amin" placeholder="מינ׳"></div>
      <div class="f"><label>עד</label><input type="number" class="amax" placeholder="מקס׳"></div>
    </div>
    <div class="filters colf"></div>
    <div class="chips"></div>
    <div class="toolbar">
      <span class="stats"></span>
      <span style="flex:1"></span>
      <div class="ms colsms"></div>
      <button class="btn export">הורד CSV</button>
      <button class="btn clear">נקה</button>
    </div>
    <div class="tw grid"></div>
    <div class="pager"></div>
  </section>
  <section class="card pivot" style="margin-bottom:10px">
    <div class="toolbar"><h2 style="margin:0">טבלת ציר מהירה</h2>
      <div class="f"><label>שורות</label><select class="prow"></select></div>
      <div class="f"><label>עמודות</label><select class="pcol"></select></div>
      <div class="f"><label>מדד</label><select class="pm"><option value="sum">סכום</option><option value="count">מס' עסקאות</option><option value="avg">ממוצע לעסקה</option></select></div>
      <span class="small">מחושב על השורות המסוננות; לחיצה על תא מסננת את הרשימה.</span></div>
    <div class="tw ptable"></div>
  </section>`;
  const q = (s) => this.root.querySelector(s);
  q('.q').addEventListener('input', debounce(e=>{ st.q=e.target.value.trim().toLowerCase(); self.update(); },200));
  if (showTypes){
    const tc = q('.types');
    tc.innerHTML = TYPE_ORDER.map(t=>`<button class="tchip" data-v="${esc(t)}">${esc(t)} <span class="n"></span></button>`).join('');
    tc.querySelectorAll('button').forEach(b => b.addEventListener('click', ()=>{ const v=b.dataset.v; if(st.types.has(v)) st.types.delete(v); else st.types.add(v); self.update(); }));
  }
  q('.win').querySelectorAll('button').forEach(b => b.addEventListener('click', ()=>{ st.win=b.dataset.v; self.update(); }));
  q('.summed').querySelectorAll('button').forEach(b => b.addEventListener('click', ()=>{ st.summed=b.dataset.v; self.update(); }));
  q('.from').addEventListener('change', e=>{ st.from=e.target.value; self.update(); });
  q('.to').addEventListener('change', e=>{ st.to=e.target.value; self.update(); });
  q('.amin').addEventListener('input', debounce(e=>{ st.amin = e.target.value===''?null:+e.target.value; self.update(); },250));
  q('.amax').addEventListener('input', debounce(e=>{ st.amax = e.target.value===''?null:+e.target.value; self.update(); },250));
  const colf = q('.colf');
  const fcols = (o.filterCols || FILTER_COLS).filter(k => COL[k] && (k!=='trip' || META.has_trip));
  fcols.forEach(k => {
    const f = document.createElement('div'); f.className='f';
    f.innerHTML = `<label>${esc(COL[k].h)}</label><div class="ms"></div>`;
    colf.appendChild(f);
    const order = k==='person' ? PERSON_ORDER : (k==='type' ? TYPE_ORDER : null);
    self.ms[k] = makeMultiEl(f.querySelector('.ms'), st.cf[k], () => {
      const base = self.base().filter(r => k!=='trip' || r.trip);
      const arr = optList(base, k, order);
      if (k==='month') arr.sort((a,b)=>a.v.localeCompare(b.v));
      if (k==='tr_section') return arr.map(a=>({v:String(a.v), n:a.n})).sort((a,b)=>a.v.localeCompare(b.v));
      return arr;
    }, {onChange:()=>self.update(), labelFn: k==='month' ? mName : (k==='tr_section' ? (v=>SEC_NAME[v]||v) : undefined)});
  });
  const cms = q('.colsms');
  cms.innerHTML = '<button type="button">עמודות</button><div class="pop cols-pop"></div>';
  const cbtn = cms.querySelector('button'), cpop = cms.querySelector('.pop');
  cbtn.addEventListener('click', e => { e.stopPropagation(); const open=cms.classList.contains('open'); closeAll(); if(!open){ self.renderColChooser(); cms.classList.add('open'); } });
  cpop.addEventListener('click', e => e.stopPropagation());
  q('.export').addEventListener('click', ()=> self.exportCSV());
  q('.clear').addEventListener('click', ()=> self.clear());
  const prow = q('.prow'), pcol = q('.pcol'), pm = q('.pm');
  const rowCols = (o.pivotRows || PIVOT_ROW_COLS).filter(k => COL[k] && (k!=='trip' || META.has_trip));
  prow.innerHTML = rowCols.map(k=>`<option value="${k}">${esc(COL[k].h)}</option>`).join('');
  pcol.innerHTML = PIVOT_COL_COLS.map(c=>`<option value="${c[0]}">${c[1]}</option>`).join('');
  prow.value = st.pivot.row; pcol.value = st.pivot.col; pm.value = st.pivot.m;
  prow.addEventListener('change', ()=>{ st.pivot.row=prow.value; self.renderPivot(); });
  pcol.addEventListener('change', ()=>{ st.pivot.col=pcol.value; self.renderPivot(); });
  pm.addEventListener('change', ()=>{ st.pivot.m=pm.value; self.renderPivot(); });
  this.update();
};
Explorer.prototype.colOrder = function(){
  const avail = this.opts.availCols || COLS.map(c=>c.k);
  const first = (this.opts.cols||DEFAULT_COLS).filter(k=>avail.includes(k));
  return first.concat(avail.filter(k=>!first.includes(k)));
};
Explorer.prototype.renderColChooser = function(){
  const self = this, st = this.st, pop = this.root.querySelector('.cols-pop');
  const avail = this.colOrder();
  pop.innerHTML = `<div class="tools"><button data-a="all">הצג הכל</button><button data-a="def">ברירת מחדל</button></div>` +
    avail.map(k=>`<label><input type="checkbox" value="${k}" ${st.cols.has(k)?'checked':''}> ${esc(COL[k].h)}</label>`).join('');
  pop.querySelectorAll('input').forEach(inp => inp.addEventListener('change', ()=>{ if(inp.checked) st.cols.add(inp.value); else st.cols.delete(inp.value); self.renderGrid(); }));
  pop.querySelector('[data-a=all]').addEventListener('click', ()=>{ avail.forEach(k=>st.cols.add(k)); self.renderColChooser(); self.renderGrid(); });
  pop.querySelector('[data-a=def]').addEventListener('click', ()=>{ st.cols = new Set(self.opts.cols||DEFAULT_COLS); self.renderColChooser(); self.renderGrid(); });
};
Explorer.prototype.clear = function(){
  const st = this.st, root = this.root;
  st.q=''; st.types=new Set(this.opts.types||[]); st.win='all'; st.summed='all'; st.from=''; st.to=''; st.amin=null; st.amax=null; st.merchant=null; st.page=0;
  Object.values(st.cf).forEach(s=>s.clear());
  ['.q','.from','.to','.amin','.amax'].forEach(s => { root.querySelector(s).value=''; });
  Object.values(this.ms).forEach(m=>m.render());
  this.update();
};
Explorer.prototype.update = function(){
  const st = this.st, root = this.root;
  st.page = 0;
  this.rows = this.filtered();
  root.querySelectorAll('.win button').forEach(b=>b.classList.toggle('on', b.dataset.v===st.win));
  root.querySelectorAll('.summed button').forEach(b=>b.classList.toggle('on', b.dataset.v===st.summed));
  const tc = root.querySelector('.types');
  if (tc){ const cnt = countBy(this.base(), 'type'); tc.querySelectorAll('button').forEach(b=>{ b.classList.toggle('on', st.types.has(b.dataset.v)); b.querySelector('.n').textContent = cnt.get(b.dataset.v)||0; }); }
  Object.values(this.ms).forEach(m=>m.sync());
  this.renderChips(); this.renderGrid(); this.renderPivot();
  if (this.opts.onUpdate) this.opts.onUpdate(this);
};
Explorer.prototype.renderChips = function(){
  const self=this, st=this.st, c=this.root.querySelector('.chips'); const items=[];
  const chip=(l,v,fn)=>items.push({l,v,fn});
  if (st.q) chip('חיפוש', st.q, ()=>{st.q=''; self.root.querySelector('.q').value='';});
  if (st.types.size && !(this.opts.types && st.types.size===this.opts.types.length && this.opts.types.every(t=>st.types.has(t)))) chip('סוג', Array.from(st.types).join(', '), ()=>{st.types=new Set(self.opts.types||[]);});
  if (st.win!=='all') chip('חלון', st.win==='in'?'בחלון':'מחוץ לחלון', ()=>{st.win='all';});
  if (st.summed!=='all') chip('נסכם', st.summed==='yes'?'כן':'לא', ()=>{st.summed='all';});
  if (st.from) chip('מתאריך', fmtDate(st.from), ()=>{st.from=''; self.root.querySelector('.from').value='';});
  if (st.to) chip('עד תאריך', fmtDate(st.to), ()=>{st.to=''; self.root.querySelector('.to').value='';});
  if (st.amin!=null) chip('סכום מ־', fmt(st.amin), ()=>{st.amin=null; self.root.querySelector('.amin').value='';});
  if (st.amax!=null) chip('סכום עד', fmt(st.amax), ()=>{st.amax=null; self.root.querySelector('.amax').value='';});
  if (st.merchant) chip('בית עסק', st.merchant, ()=>{st.merchant=null;});
  Object.keys(st.cf).forEach(k=>{ if (st.cf[k].size){ const lf = k==='month'?mName:(k==='tr_section'?(v=>SEC_NAME[v]||v):(v=>v)); chip(COL[k]?COL[k].h:k, Array.from(st.cf[k]).map(lf).join(', '), ()=>{ st.cf[k].clear(); if(self.ms[k]) self.ms[k].render(); }); } });
  c.innerHTML = items.map((it,i)=>`<span class="chip"><span>${esc(it.l)}: <b title="${esc(it.v)}">${esc(it.v)}</b></span><button title="הסר" data-i="${i}">×</button></span>`).join('');
  c.querySelectorAll('button').forEach(b=>b.addEventListener('click', ()=>{ items[+b.dataset.i].fn(); self.update(); }));
};
Explorer.prototype.sorted = function(){
  const s = this.st.sort, k = s.key, t = (COL[k]||{}).t;
  const arr = this.rows.slice();
  arr.sort((a,b)=>{
    let va=a[k], vb=b[k];
    if (t==='num' || t==='bool' || k==='tr_section'){ va = va==null?-Infinity:(+va||0); vb = vb==null?-Infinity:(+vb||0); return (va-vb)*s.dir; }
    va = String(va==null?'':va); vb = String(vb==null?'':vb);
    if (va==='' && vb!=='') return 1; if (vb==='' && va!=='') return -1;
    return va.localeCompare(vb,'he')*s.dir;
  });
  return arr;
};
Explorer.prototype.renderGrid = function(){
  const self=this, st=this.st, root=this.root;
  const rows = this.sorted();
  const n = rows.length, sum = rows.reduce((s,r)=>s+(r.amount||0),0);
  const pos = rows.reduce((s,r)=>s+(r.amount>0?r.amount:0),0), neg = rows.reduce((s,r)=>s+(r.amount<0?r.amount:0),0);
  root.querySelector('.stats').innerHTML = `<span><b>${nf0.format(n)}</b> שורות</span><span>סה"כ <b>${fmt2(sum)} ₪</b></span><span>ממוצע <b>${n?fmt2(sum/n):'0'} ₪</b></span>` + (neg<0 ? `<span class="small">(חיובים ${fmt2(pos)} ₪ · זיכויים ${fmt2(neg)} ₪)</span>` : '');
  const cols = this.colOrder().filter(k=>st.cols.has(k));
  const th = cols.map(k=>{ const c=COL[k]; const num = c.t==='num'; return `<th class="${num?'num':''}" data-k="${k}" title="${esc(c.h)}">${st.sort.key===k?`<span class="arr">${st.sort.dir<0?'▼':'▲'}</span>`:''}${esc(c.h)}</th>`; }).join('');
  const start = st.pageSize===Infinity ? 0 : st.page*st.pageSize;
  const end = st.pageSize===Infinity ? n : Math.min(n, start+st.pageSize);
  const body = rows.slice(start,end).map(r=>{
    const cls = ['click','zebra']; if (!r.in_window) cls.push('dim');
    return `<tr class="${cls.join(' ')}" data-id="${esc(r.id)}">` + cols.map(k=>{
      const c=COL[k]; const v=r[k];
      if (c.t==='num'){ return `<td class="num${(v<0)?' neg':''}">${v==null?'':(k==='amount'? fmt2(v)+' ₪' : fmt2(v))}</td>`; }
      const txt = cellText(r,k);
      let extra = '';
      if (k==='name_clean'){ if (!r.summed) extra += '<span class="badge">לא נסכם</span>'; if (!r.in_window) extra += '<span class="badge out">מחוץ לחלון</span>'; }
      if (k==='tr_section' && v) return `<td><span class="badge sec">${esc(SEC_NAME[v])}</span></td>`;
      const long = (c.t==='text' && k!=='id');
      return `<td class="${long?'txt':''}${k==='name_clean'?' name':''}" title="${esc(txt)}">${extra}${esc(txt)}</td>`;
    }).join('') + '</tr>';
  }).join('');
  root.querySelector('.grid').innerHTML = `<table><thead><tr>${th}</tr></thead><tbody>${body || `<tr><td colspan="${cols.length}" class="small">אין שורות בבחירה הנוכחית</td></tr>`}</tbody></table>`;
  root.querySelectorAll('.grid th').forEach(h=>h.addEventListener('click', ()=>{ const k=h.dataset.k; if (st.sort.key===k) st.sort.dir=-st.sort.dir; else { st.sort.key=k; st.sort.dir = (COL[k].t==='num'||COL[k].t==='date') ? -1 : 1; } self.renderGrid(); }));
  root.querySelectorAll('.grid tbody tr.click').forEach(tr=>tr.addEventListener('click', ()=> openRow(BY_ID[tr.dataset.id], self)));
  const pg = root.querySelector('.pager');
  if (st.pageSize===Infinity){ pg.innerHTML = `<span class="small">מוצגות כל ${n} השורות</span> <button class="btn sm" data-a="page">הצג 100 בעמוד</button>`; }
  else {
    const pages = Math.max(1, Math.ceil(n/st.pageSize));
    pg.innerHTML = `<button class="btn sm" data-a="prev" ${st.page<=0?'disabled':''}>‹ הקודם</button><span>עמוד ${st.page+1} מתוך ${pages} (שורות ${n?start+1:0}–${end})</span><button class="btn sm" data-a="next" ${st.page>=pages-1?'disabled':''}>הבא ›</button><button class="btn sm" data-a="all">הצג הכל (${n})</button>`;
  }
  pg.querySelectorAll('button').forEach(b=>b.addEventListener('click', ()=>{ const a=b.dataset.a; if(a==='prev') st.page--; else if(a==='next') st.page++; else if(a==='all') st.pageSize=Infinity; else if(a==='page'){ st.pageSize=100; st.page=0; } self.renderGrid(); }));
};
Explorer.prototype.renderPivot = function(){
  const self=this, st=this.st, p=st.pivot, rows=this.rows;
  const rk=p.row, ck=p.col;
  const rval = r => rk==='tr_section' ? (SEC_NAME[r[rk]]||'(ריק)') : (COL[rk].t==='bool' ? (r[rk]?'כן':'לא') : (r[rk]||'(ריק)'));
  const cval = r => ck ? (r[ck]||'(ריק)') : '';
  const cells = new Map(); const rtot = new Map(); const ctot = new Map(); const colSet = new Set();
  const acc = (map,k,r)=>{ let a=map.get(k); if(!a){a={s:0,n:0}; map.set(k,a);} a.s+=r.amount||0; a.n++; };
  rows.forEach(r=>{ const rv=rval(r), cv=cval(r); colSet.add(cv); acc(cells, rv+''+cv, r); acc(rtot, rv, r); acc(ctot, cv, r); });
  const measure = a => !a ? null : (p.m==='sum' ? a.s : p.m==='count' ? a.n : (a.n? a.s/a.n : 0));
  const fmtM = v => v==null ? '' : (p.m==='count' ? nf0.format(v) : fmt(v));
  const colsArr = Array.from(colSet);
  if (ck==='month') colsArr.sort(); else colsArr.sort((a,b)=> (measure(ctot.get(b))||0)-(measure(ctot.get(a))||0));
  const rowsArr = Array.from(rtot.keys()).sort((a,b)=> (measure(rtot.get(b))||0)-(measure(rtot.get(a))||0));
  const all = rows.reduce((a,r)=>{a.s+=r.amount||0; a.n++; return a;},{s:0,n:0});
  const clabel = c => ck==='month' ? mName(c) : c;
  const MAXR = 60;
  const head = `<tr><th class="rowh">${esc(COL[rk].h)}</th>` + (ck ? colsArr.map(c=>`<th class="num" title="${esc(clabel(c))}">${esc(clabel(c))}</th>`).join('') : '') + `<th class="num">סה"כ</th><th class="num">מס'</th></tr>`;
  const body = rowsArr.slice(0,MAXR).map(rv => `<tr><td class="rowh name" title="${esc(rv)}">${esc(rv)}</td>` +
      (ck ? colsArr.map(c=>{ const a=cells.get(rv+''+c); const v=measure(a); return `<td class="num cell${(v<0)?' neg':''}" data-r="${esc(rv)}" data-c="${esc(c)}">${fmtM(v)}</td>`; }).join('') : '') +
      `<td class="num cell" data-r="${esc(rv)}"><b>${fmtM(measure(rtot.get(rv)))}</b></td><td class="num">${rtot.get(rv).n}</td></tr>`).join('');
  const foot = `<tr><td class="rowh">סה"כ (${rowsArr.length}${rowsArr.length>MAXR?', מוצגות '+MAXR:''})</td>` + (ck ? colsArr.map(c=>`<td class="num">${fmtM(measure(ctot.get(c)))}</td>`).join('') : '') + `<td class="num">${fmtM(measure(all))}</td><td class="num">${all.n}</td></tr>`;
  const el = this.root.querySelector('.ptable');
  el.innerHTML = `<table><thead>${head}</thead><tbody>${body || '<tr><td class="small">אין נתונים</td></tr>'}</tbody><tfoot>${foot}</tfoot></table>`;
  el.querySelectorAll('td.cell').forEach(td => td.addEventListener('click', ()=>{
    const rv = td.dataset.r, cv = td.dataset.c;
    self.applyDim(rk, rv);
    if (ck && cv!=null) self.applyDim(ck, cv);
    self.update();
    self.root.querySelector('.grid').scrollIntoView({behavior:'smooth', block:'start'});
  }));
};
Explorer.prototype.applyDim = function(k, v){
  const st = this.st;
  if (v==='(ריק)') v = '';
  if (k==='name_clean'){ st.merchant = v; return; }
  if (k==='tr_section'){ const num = Object.keys(SEC_NAME).find(n=>SEC_NAME[n]===v); if (num){ st.cf.tr_section.clear(); st.cf.tr_section.add(String(num)); if (this.ms.tr_section) this.ms.tr_section.render(); } return; }
  if (k==='type'){ st.types.clear(); st.types.add(v); return; }
  if (k==='in_window'){ st.win = v==='כן' ? 'in' : 'out'; return; }
  if (k==='summed'){ st.summed = v==='כן' ? 'yes' : 'no'; return; }
  if (!st.cf[k]) st.cf[k] = new Set();
  st.cf[k].clear(); st.cf[k].add(v);
  if (this.ms[k]) this.ms[k].render();
};
Explorer.prototype.exportCSV = function(){
  const keys = SHEET_KEYS.concat(this.opts.exportExtra||[]);
  const hdr = keys.map(k=>COL[k].h);
  const esc1 = v => { const s = String(v==null?'':v); return /[",\n\r]/.test(s) ? '"'+s.replace(/"/g,'""')+'"' : s; };
  const rows = this.sorted();
  const lines = [hdr.map(esc1).join(',')];
  rows.forEach(r => lines.push(keys.map(k=>{ const c=COL[k]; const v=r[k]; if (c.t==='bool') return v?'כן':'לא'; if (c.t==='num') return v==null?'':v; if (k==='tr_section') return SEC_NAME[v]||''; return esc1(v); }).join(',')));
  const csv = '\uFEFF' + lines.join('\r\n');
  const blob = new Blob([csv], {type:'text/csv;charset=utf-8'});
  const url = URL.createObjectURL(blob);
  const a = document.createElement('a'); a.href = url; a.download = 'transactions_' + (new Date()).toISOString().slice(0,10) + '.csv';
  document.body.appendChild(a); a.click(); document.body.removeChild(a);
  setTimeout(()=>URL.revokeObjectURL(url), 2000);
  return {rows: rows.length, bytes: csv.length};
};

/* ---------- row modal ---------- */
let modalExplorer = null;
function closeModal(){ document.getElementById('modal').classList.remove('open'); }
document.getElementById('modal').addEventListener('click', e => { if (e.target.id==='modal') closeModal(); });
function openRow(r, explorer){
  if (!r) return;
  modalExplorer = explorer || EXPLORERS.exMain;
  const sheet = document.getElementById('modalSheet');
  const kv = SHEET_KEYS.map(k => { const c=COL[k]; const txt = cellText(r,k); const full = (c.t==='text' && txt.length>40); return `<div class="${full?'full':''}"><span class="k">${esc(c.h)}</span><span class="v${(k==='amount'&&r.amount<0)?' neg':''}">${esc(txt)}</span></div>`; }).join('');
  const extra = r.tr_section ? `<div><span class="k">סעיף (העברות)</span><span class="v">${esc(SEC_NAME[r.tr_section])}</span></div><div><span class="k">כיוון</span><span class="v">${esc(r.tr_dir)}</span></div><div class="full"><span class="k">מוטב / צד שני</span><span class="v">${esc(r.tr_payee)}</span></div><div class="full"><span class="k">הערה מצילום</span><span class="v">${esc(r.tr_note)}</span></div>` : '';
  const same = DATA.filter(x => x.name_clean===r.name_clean && x.id!==r.id).sort((a,b)=>(b.txn_date||'').localeCompare(a.txn_date||''));
  const linked = r.linked_id ? BY_ID[r.linked_id] : null;
  const backlinks = DATA.filter(x => x.linked_id===r.id);
  const relRow = x => `<tr class="click" data-id="${esc(x.id)}"><td>${esc(x.id)}</td><td>${fmtDate(x.txn_date)}</td><td class="num${x.amount<0?' neg':''}">${fmt2(x.amount)} ₪</td><td>${esc(x.type)}</td><td>${esc(x.cat_tz)}</td><td>${esc(x.pay)}</td><td>${x.in_window?'כן':'לא'} / ${x.summed?'כן':'לא'}</td><td class="txt" title="${esc(x.details)}">${esc(x.details)}</td></tr>`;
  const relHead = `<tr><th>מזהה</th><th>תאריך</th><th class="num">סכום</th><th>סוג</th><th>קטגוריה (תזרים)</th><th>אמצעי תשלום</th><th>בחלון / נסכם</th><th>פרטים</th></tr>`;
  const sameSum = same.reduce((s,x)=>s+(x.amount||0),0);
  const flags = (r.summed?'':'<span class="badge">לא נסכם</span>') + (r.in_window?'':'<span class="badge out">מחוץ לחלון</span>');
  sheet.innerHTML = `<button class="close" title="סגור (Esc)">×</button>
    <h2>${esc(r.name_clean)} <span class="small">· ${esc(r.id)} · ${esc(r.type)}</span></h2>
    <div class="amt${r.amount<0?' neg':''}">${fmt2(r.amount)} ₪ <span class="small" style="font-weight:400">${fmtDate(r.txn_date)} · ${esc(r.pay)} · ${esc(r.person)}</span> ${flags}</div>
    <div class="actions"><button class="btn sm primary" data-a="merchant">סנן לפי בית עסק זה</button><button class="btn sm" data-a="cat">סנן לפי קטגוריה זו (${esc(r.cat_tz)})</button><button class="btn sm" data-a="catnew">לפי קטגוריה מוצעת (${esc(r.cat_new)})</button></div>
    <div class="kv">${kv}${extra}</div>
    <div class="rel">
      ${(linked||backlinks.length) ? `<h3>שורה מקושרת (מזהה מקושר)</h3><div class="tw" style="max-height:200px"><table><thead>${relHead}</thead><tbody>${(linked?[linked]:[]).concat(backlinks).map(relRow).join('')}</tbody></table></div>` : ''}
      <h3>עסקאות קשורות — אותו "שם מובן" (${same.length}${same.length?', סה"כ '+fmt2(sameSum)+' ₪':''})</h3>
      ${same.length ? `<div class="tw" style="max-height:260px"><table><thead>${relHead}</thead><tbody>${same.slice(0,200).map(relRow).join('')}</tbody></table></div>${same.length>200?'<div class="small">מוצגות 200 הראשונות</div>':''}` : '<div class="small">אין עסקאות נוספות באותו שם.</div>'}
    </div>`;
  sheet.querySelector('.close').addEventListener('click', closeModal);
  sheet.querySelectorAll('tr.click').forEach(tr => tr.addEventListener('click', ()=> openRow(BY_ID[tr.dataset.id], modalExplorer)));
  sheet.querySelectorAll('.actions button').forEach(b => b.addEventListener('click', ()=>{
    const ex = modalExplorer || EXPLORERS.exMain; const a = b.dataset.a;
    if (a==='merchant') ex.st.merchant = r.name_clean;
    else if (a==='cat') ex.applyDim('cat_tz', r.cat_tz);
    else ex.applyDim('cat_new', r.cat_new);
    ex.update(); closeModal();
    const tab = ex.id==='exTransfers' ? 'transfers' : ex.id==='exIncome' ? 'income' : 'explore';
    showTab(tab); ex.root.scrollIntoView({behavior:'smooth', block:'start'});
  }));
  document.getElementById('modal').classList.add('open');
  document.getElementById('modal').scrollTop = 0;
}

/* ---------- transfers / income / unclassified ---------- */
function initTransfers(){
  const secs = document.getElementById('trSecs');
  const ex = new Explorer('exTransfers', {
    baseRows: () => DATA.filter(r => r.tr_section),
    cols: ['tr_section','tr_dir','txn_date','amount','tr_payee','name_clean','type','cat_tz','pay','person','summed','tr_note','rule_note','details','linked_id'],
    filterCols: ['tr_section','tr_dir','source','pay','person','cat_tz','month'],
    pivotRow: 'tr_payee', exportExtra: ['tr_section','tr_dir','tr_payee','tr_note'],
    onUpdate: (ex2) => {
      const rows = ex2.rows;
      secs.innerHTML = [1,2,3].map(s => {
        const rs = rows.filter(r=>r.tr_section===s);
        const out = rs.filter(r=>r.tr_dir==='יוצא').reduce((a,r)=>a+Math.abs(r.amount||0),0);
        const inn = rs.filter(r=>r.tr_dir==='נכנס').reduce((a,r)=>a+Math.abs(r.amount||0),0);
        const sel = ex2.st.cf.tr_section.size===1 && ex2.st.cf.tr_section.has(String(s));
        return `<div class="kpi click${sel?' sel':''}" data-s="${s}"><div class="l">${esc(SEC_LONG[s])}</div><div class="v" style="font-size:16px">${rs.length} שורות</div><div class="h">יוצא <b>${fmt(out)}</b> · נכנס <b>${fmt(inn)}</b></div></div>`;
      }).join('');
      secs.querySelectorAll('.kpi').forEach(k => k.addEventListener('click', ()=>{ const s=k.dataset.s; if (ex2.st.cf.tr_section.has(s) && ex2.st.cf.tr_section.size===1) ex2.st.cf.tr_section.clear(); else { ex2.st.cf.tr_section.clear(); ex2.st.cf.tr_section.add(s); } if (ex2.ms.tr_section) ex2.ms.tr_section.render(); ex2.update(); }));
    }
  });
  ex.st.sort = {key:'amount', dir:-1};
  ex.update();
}
function renderIncomeTable(){
  const all = document.getElementById('incAll').checked;
  const rows = DATA.filter(r => r.type===INC && (all || (r.in_window && r.summed)));
  const months = all ? ALL_MONTHS : WINDOW;
  const agg = aggregate(rows, 'cat_tz', months).sort((a,b)=>b.total-a.total);
  const heatMax = Math.max(0, ...agg.flatMap(a=>months.map(m=>a.months[m]||0)));
  const nM = Math.max(months.length,1);
  const th = `<tr><th>מקור (קטגוריה תזרים)</th><th class="num">סה"כ</th><th class="num">ממוצע</th><th class="num">ממוצע ללא אפס</th><th class="num">מס'</th><th class="num">%</th>${months.map(m=>`<th class="num">${mName(m)}</th>`).join('')}</tr>`;
  const tot = {total:0,count:0,months:{}};
  const body = agg.map(a=>{ tot.total+=a.total; tot.count+=a.count; months.forEach(m=>tot.months[m]=(tot.months[m]||0)+(a.months[m]||0));
    return `<tr class="click" data-v="${esc(a.name)}"><td class="name">${esc(a.name)}</td><td class="num${a.total<0?' neg':''}">${fmt(a.total)}</td><td class="num">${fmt(a.avg)}</td><td class="num">${fmt(a.avgNZ)}</td><td class="num">${a.count}</td><td class="num">${pct(a.pct)}</td>${months.map(m=>{const v=a.months[m]||0; return `<td class="num heat" style="${heatStyle(v,heatMax)}">${v?fmt(v):'·'}</td>`;}).join('')}</tr>`; }).join('');
  const nz = months.filter(m=>Math.abs(tot.months[m]||0)>0.005).length;
  const foot = `<tr><td>סה"כ (${agg.length})</td><td class="num">${fmt(tot.total)}</td><td class="num">${fmt(tot.total/nM)}</td><td class="num">${fmt(nz?tot.total/nz:0)}</td><td class="num">${tot.count}</td><td class="num">100%</td>${months.map(m=>`<td class="num">${fmt(tot.months[m]||0)}</td>`).join('')}</tr>`;
  const el = document.getElementById('incTable');
  el.innerHTML = `<table><thead>${th}</thead><tbody>${body}</tbody><tfoot>${foot}</tfoot></table>`;
  el.querySelectorAll('tbody tr.click').forEach(tr => tr.addEventListener('click', ()=>{ const ex=EXPLORERS.exIncome; ex.applyDim('cat_tz', tr.dataset.v); ex.update(); ex.root.scrollIntoView({behavior:'smooth', block:'start'}); }));
}
function initIncome(){
  document.getElementById('incAll').addEventListener('change', renderIncomeTable);
  renderIncomeTable();
  new Explorer('exIncome', { baseRows: () => DATA.filter(r => r.type===INC), showTypes:false,
    cols: ['txn_date','name_clean','amount','cat_tz','pay','person','month_name','in_window','summed','rule_note','details'],
    filterCols: ['source','pay','person','cat_tz','month'], pivotRow:'cat_tz' });
}
function initUnclassified(){
  const rows = DATA.filter(r => (r.type===EXP||r.type===INC) && (r.cat_tz===META.unknown_cat || /לסיווג ידני|לבדיקה|נדרש:/.test(r.rule_note||'')))
    .sort((a,b)=> Math.abs(b.amount||0)-Math.abs(a.amount||0));
  document.getElementById('uncCount').textContent = '('+rows.length+' שורות)';
  document.getElementById('nUnc').textContent = rows.length;
  document.getElementById('uncNote').innerHTML = `שורות שקטגוריית התזרים שלהן היא "${esc(META.unknown_cat)}" או שההערה מכילה "לסיווג ידני" / "לבדיקה" / "נדרש:". הסיווג עצמו נעשה באקסל, בלשונית <b>${esc(META.manual_sheet)}</b> (עמודות "${esc(META.user_text_col)}" ו-"${esc(META.user_cat_col)}"), ואז מריצים את הצינור מחדש ובונים את הדשבורד.`;
  const body = rows.map(r=>`<tr class="click zebra" data-id="${esc(r.id)}"><td>${esc(r.id)}</td><td>${fmtDate(r.txn_date)}</td><td class="name" title="${esc(r.original_name)}">${esc(r.name_clean)}</td><td class="txt" title="${esc(r.original_name)}">${esc(r.original_name)}</td><td class="num${r.amount<0?' neg':''}">${fmt2(r.amount)} ₪</td><td>${esc(r.type)}</td><td>${esc(r.cat_tz)}</td><td>${esc(r.pay)}</td><td>${r.in_window?'כן':'לא'} / ${r.summed?'כן':'לא'}</td><td class="txt" title="${esc(r.rule_note+' | '+r.details)}">${esc(r.rule_note)}${r.details?' · '+esc(r.details):''}</td></tr>`).join('');
  const el = document.getElementById('uncTable');
  el.innerHTML = `<table><thead><tr><th>מזהה</th><th>תאריך</th><th>שם מובן</th><th>שם כפי שמופיע בקובץ</th><th class="num">סכום</th><th>סוג תנועה</th><th>קטגוריה (תזרים)</th><th>אמצעי תשלום</th><th>בחלון / נסכם</th><th>הערה / פרטים</th></tr></thead><tbody>${body || '<tr><td colspan="10" class="small">אין שורות לסיווג</td></tr>'}</tbody></table>`;
  el.querySelectorAll('tr.click').forEach(tr => tr.addEventListener('click', ()=> openRow(BY_ID[tr.dataset.id], EXPLORERS.exMain)));
}
function renderFooter(){
  const t = META.n_by_type;
  document.getElementById('foot').innerHTML =
    `<div>נבנה: <b>${esc(META.built)}</b> · קובץ מקור: <b><bdi dir="ltr">${esc(META.source_file)}</bdi></b> (לשונית ${esc(META.sheet)}) · ${META.n_total} שורות × ${META.cols.length} עמודות` +
    ` (${Object.keys(t).map(k=>esc(k)+': '+t[k]).join(', ')}) · בסקירה (נסכם ובחלון, ${esc(META.window_label)}): ${META.n_default} שורות — ${META.n_default_exp} הוצאות, ${META.n_default_inc} הכנסות · העברות/אפליקציה/PAYBOX: ${META.n_tr[1]||0}/${META.n_tr[2]||0}/${META.n_tr[3]||0}.</div>` +
    `<div>לעדכון: שנה באקסל → <code>${esc(META.refresh_cmd)}</code> → רענן את הדף.</div>` +
    `<div class="small">סקירה: ממוצע חודשי = סה"כ ÷ מספר החודשים הנבחרים; "ללא אפס" = ÷ חודשים עם תנועה; % = מתוך כלל ההוצאות/ההכנסות בחודשים הנבחרים.</div>`;
}

/* ---------- init ---------- */
function init(){
  if (HAS_CHART){ Chart.defaults.font.family = 'system-ui, -apple-system, "Segoe UI", Roboto, Arial, sans-serif'; Chart.defaults.color = '#52514e'; }
  document.querySelectorAll('#tabs button').forEach(b => { b.hidden = !META.tabs.includes(b.dataset.tab); });
  ['cmd1','cmd2'].forEach(id => { const el=document.getElementById(id); if (el) el.textContent = META.refresh_cmd; });
  initOverview();
  new Explorer('exMain', { baseRows: () => DATA });
  document.getElementById('nExplore').textContent = DATA.length;
  document.getElementById('nTransfers').textContent = DATA.filter(r=>r.tr_section).length;
  document.getElementById('nIncome').textContent = DATA.filter(r=>r.type===INC).length;
  TAB_INIT.transfers = () => { if (!EXPLORERS.exTransfers) initTransfers(); };
  TAB_INIT.income = () => { if (!EXPLORERS.exIncome) initIncome(); };
  if (META.tabs.includes('unc')) initUnclassified();
  renderFooter();
  document.querySelectorAll('#tabs button').forEach(b => b.addEventListener('click', ()=> showTab(b.dataset.tab)));
  document.querySelectorAll('[data-go]').forEach(a => a.addEventListener('click', e => { e.preventDefault(); showTab(a.dataset.go); }));
  const h = (location.hash||'').replace('#','');
  showTab(META.tabs.includes(h) ? h : 'overview');
}
document.addEventListener('DOMContentLoaded', init);
</script>
</body>
</html>
"""


def main(argv=None):
    def extra(p):
        p.add_argument("--workbook", default=None, help="override outputs/תזרים.xlsx")
        p.add_argument("--database", default=None, help="override work/database.csv (fallback source)")
        p.add_argument("--out", default=None, help="override outputs/dashboard.html")
    args, cfg = parse_args("build the single-file RTL dashboard from the workbook's database sheet (P17)", extra, argv)
    try:
        res = build_dashboard(cfg, args.workbook, args.database, args.out)
    except FileNotFoundError as e:
        fail(str(e), "run build_excel.py (or classify.py for the CSV fallback) first")
    log("dashboard: %d rows x %d cols from %s -> %s (%.0f KB, chart.js %s)"
        % (res["rows"], res["cols"], res["source"], res["dashboard"], res["bytes"] / 1024.0, res["chartjs"]))
    emit(res)


if __name__ == "__main__":
    main()
