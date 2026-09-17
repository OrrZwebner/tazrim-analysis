#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""build_excel.py — build the formula-driven analysis workbook (P11) from work/database.csv.

Purpose : write outputs/תזרים.xlsx: a `database` Excel Table with named ranges over every
          column, the `עזר_חודשי` SUMIFS helper, the `הוצאות` / `הכנסות` statistic sheets
          (N month columns, totals, averages incl. "ממוצע ללא אפס", MAX/MIN month, median,
          STDEV.S, group subtotals), `ניתוח נתונים` (sorted tables + native
          charts) and — from level `standard` — the interactive `פילוח` sheet (dynamic arrays),
          `פירוט לפי קטגוריה`, `פירוט עסקאות`, `הוצאות משתנות לפי בית עסק`, `לסיווג ידני`
          (Excel Table with two user columns) and the hidden `רשימות`; at level `deep` also
          `קטגוריות מוצעות` (secondary scheme). Every aggregate is an Excel formula over
          `database`. Sheet set = common.SHEETS_BY_LEVEL[level]; order = common.SHEET_ORDER.
Inputs  : tazrim.config.json (window -> N months, household, cards, categories, thresholds,
          outputs.workbook), work/database.csv (exactly common.DB_COLUMNS), optional template
          workbook (categories.from_template — ONLY its category list is read; the workbook is
          always created from scratch and the template file is never copied or written),
          optional primary scheme CSV (group,cat,fixed), optional secondary scheme CSV, and the
          PREVIOUS output workbook (its "לסיווג ידני" user columns are carried forward).
Outputs : the workbook; one JSON object on stdout {ok, workbook, level, n_months, months,
          sheets, rows_per_sheet, preserved_labels, manual_rows, dropped_categories, totals}.
Exit    : 0 ok; 1 on failure ({"ok": false, "error", "hint"}), e.g. missing database.csv,
          a category name containing a comma (FM5), wrong columns; 2 on config error.

Formula conventions (failure modes FM1-FM4): the full row order is decided before writing
(no insert_rows); Excel-365 functions carry the `_xlfn.` / `_xlfn._xlws.` / `_xlpm.` prefixes;
SUMPRODUCT / LET tables are written as ArrayFormula; every string literal inside a formula
doubles its quotes (q()). Python 3.8 compatible. Needs openpyxl and pandas.
"""
import datetime as _dt
import os
import re
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (  # noqa: E402
    DB_COLUMNS, DB_HEADERS_HE, HIDDEN_SHEETS, NO, SHEETS_BY_LEVEL, SHEET_ANALYSIS,
    SHEET_BY_CAT, SHEET_DB, SHEET_EXP, SHEET_HELPER, SHEET_INC, SHEET_LISTS, SHEET_MANUAL,
    SHEET_ORDER, SHEET_PIVOT, SHEET_PROPOSED, SHEET_TXNS, SHEET_VARIABLE, TYPE_CARD_DEBIT,
    TYPE_EXPENSE, TYPE_INCOME, TYPE_INTERNAL, TYPE_REIMBURSABLE, TYPE_REIMBURSEMENT,
    TYPE_SAVINGS, GENERATED_CATEGORIES, UNKNOWN_CAT, USER_CAT_COL, USER_TEXT_COL, YES, emit, fail, log, month_label,
    months, parse_args, project_path, read_csv, scheme_path,
)

try:
    import pandas as pd
    import openpyxl
    from openpyxl.chart import BarChart, LineChart, PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import ColorScaleRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.formula import ArrayFormula
    from openpyxl.worksheet.table import Table, TableStyleInfo
except ImportError as _e:  # pragma: no cover
    fail("missing dependency: %s" % _e, "pip install -r requirements.txt (openpyxl, pandas)")

# ----------------------------------------------------------------------------- constants
#: Named range per database column (name -> column key). Every SUMIFS in the workbook uses these.
NAMED_RANGES = {
    "DBid": "id", "DBsource": "source", "DBcard": "card", "DBpay": "pay", "DBperson": "person",
    "DBorig": "original_name", "DBname": "name_clean", "DBtype": "type", "DBgroup": "group_tz",
    "DBcat": "cat_tz", "DBgroup2": "group_new", "DBcat2": "cat_new", "DBdate": "txn_date",
    "DBcharge": "charge_date", "DBmonth": "month", "DBmname": "month_name", "DBamt": "amount",
    "DBcur": "orig_currency", "DBin": "in_window", "DBsummed": "summed", "DBnote": "rule_note",
    "DBdet": "details", "DBlink": "linked_id", "DBtrip": "trip",
}
TOTAL_EXP_LABEL = 'סה"כ הוצאות שוטפות (נסכם ישירות מ-database)'
TOTAL_INC_LABEL = 'סה"כ הכנסות (נסכם ישירות מ-database)'
NEEDED_COL = "נדרש:"
DISCOUNT_CAT = "הנחת מועדון (זיכוי מחושב)"
#: rule_note markers that send an in-window row to "לסיווג ידני" (any substring match).
MANUAL_MARKERS = ("לסיווג ידני", "לבדיקה", "לאימות", "נדרש:", "לבדוק")
MANUAL_CAT_MARKERS = ("לא מזוהה", "לא ידוע", "לסיווג", "לא מפורט")
SCHEME_TZ, SCHEME_NEW = "תזרים", "מוצעת"
ALL_MONTHS, ALL = "כל החודשים", "הכל"

# ----------------------------------------------------------------------------- styles
HDR = PatternFill("solid", fgColor="1F4E78")
HDRF = Font(bold=True, color="FFFFFF")
GRP = PatternFill("solid", fgColor="DDEBF7")
TOT = PatternFill("solid", fgColor="FFF2CC")
SUBTOT = PatternFill("solid", fgColor="FCE4D6")
INPUT = PatternFill("solid", fgColor="E2EFDA")
GREY = Font(italic=True, color="808080")
BOLD = Font(bold=True)
TITLE = Font(bold=True, size=13)
RED = Font(bold=True, color="C00000")
LINK = Font(color="0563C1", underline="single")
NUM = "#,##0;[Red]-#,##0"
NUM2 = "#,##0.00;[Red]-#,##0.00"
PCT = "0.0%"
_thin = Side(style="thin", color="BFBFBF")
_med = Side(style="medium", color="1F4E78")
BORDER = Border(top=_thin, bottom=_thin, left=_thin, right=_thin)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
RIGHT = Alignment(horizontal="right")


# ----------------------------------------------------------------------------- helpers
def q(text):
    """Excel string literal with doubled inner quotes (FM: a quote inside a category name)."""
    return '"' + str(text).replace('"', '""') + '"'


def sref(sheet):
    """Sheet reference prefix for formulas: 'name'! (always quoted; safe for spaces/Hebrew)."""
    return "'%s'!" % sheet.replace("'", "''")


def nz_avg(vals):
    v = [x for x in vals if abs(x) > 0.005]
    return sum(v) / len(v) if v else 0.0


def hdr(ws, row, cols, start=1):
    for i, c in enumerate(cols):
        cell = ws.cell(row, start + i, c)
        cell.fill, cell.font, cell.alignment, cell.border = HDR, HDRF, CENTER, BORDER


def setw(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col if isinstance(col, str) else L(col)].width = w


def colour_scale(ws, rng, end="9DC3E6"):
    """Full-cell colour scale white -> end colour (dark text stays readable; no data bars)."""
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                      end_type="max", end_color=end))


def heat(ws, rng):
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type="num", start_value=0, start_color="FFFFFF",
                                                      mid_type="percentile", mid_value=70, mid_color="FFF2CC",
                                                      end_type="max", end_color="F8CBAD"))


def put_text(ws, r, c, v):
    """Write a data value; strings that look like formulas are forced to text."""
    cell = ws.cell(r, c, v)
    if isinstance(v, str) and v.startswith("="):
        cell.data_type = "s"
    return cell


def rtl(ws):
    ws.sheet_view.rightToLeft = True
    return ws


class Layout(object):
    """Column layout of a statistic sheet for N months:
    קבוצה | תת-קטגוריה | months×N | סה"כ | ממוצע | ממוצע ללא אפס | מקסימום | חודש מקס |
    מינימום | חודש מין | חציון | סטיית תקן | חודשים עם חיוב | הערות."""

    def __init__(self, n):
        self.n = n
        self.c_group, self.c_cat, self.m0 = 1, 2, 3
        self.m1 = self.m0 + n - 1
        self.c_total = self.m1 + 1
        self.c_avg = self.c_total + 1
        self.c_nz = self.c_avg + 1
        s = self.c_nz
        (self.c_max, self.c_maxm, self.c_min, self.c_minm, self.c_med, self.c_std,
         self.c_cnt, self.c_note) = (s + 1, s + 2, s + 3, s + 4, s + 5, s + 6, s + 7, s + 8)
        self.ncols = self.c_note

    def headers(self, mlabels):
        return (["קבוצה", "תת-קטגוריה"] + list(mlabels) +
                ['סה"כ %d חודשים' % self.n, "ממוצע חודשי", "ממוצע ללא אפס"] +
                ["מקסימום", "חודש מקס", "מינימום", "חודש מין", "חציון", "סטיית תקן", "חודשים עם חיוב", "הערות"])

    def widths(self):
        return [14, 40] + [11] * self.n + [13, 13, 13] + [11, 11, 11, 11, 11, 11, 9, 60]

    def mrng(self, r):
        return "%s%d:%s%d" % (L(self.m0), r, L(self.m1), r)

    def mhdr(self):
        return "$%s$2:$%s$2" % (L(self.m0), L(self.m1))


# ----------------------------------------------------------------------------- data
def load_db(path):
    if not os.path.isfile(path):
        fail("database not found: %s" % path, "run classify.py first (it writes work/database.csv)")
    db = pd.read_csv(path, dtype=str, keep_default_na=False, encoding="utf-8-sig")
    missing = [c for c in DB_COLUMNS if c not in db.columns]
    if missing:
        fail("database.csv is missing columns: %s" % ", ".join(missing),
             "work/database.csv must have exactly common.DB_COLUMNS (written by classify.py)")
    db = db[DB_COLUMNS].copy()
    db["amount"] = pd.to_numeric(db["amount"], errors="coerce").fillna(0.0)
    db["orig_amount_f"] = pd.to_numeric(db["orig_amount"], errors="coerce")
    db["_in"] = db["in_window"].str.strip() == YES
    db["_summed"] = db["summed"].str.strip() == YES
    if db.empty:
        fail("database.csv has no rows", "check the parsers' output and the window in tazrim.config.json")
    return db


def read_scheme_csv(path):
    """[(group, cat, fixed_bool)] from a group,cat[,fixed] CSV; [] when the file is absent."""
    if not path or not os.path.isfile(path):
        return []
    out = []
    for r in read_csv(path, skip_comments=True):
        g, c = (r.get("group") or "").strip(), (r.get("cat") or "").strip()
        if not c:
            continue
        out.append((g, c, str(r.get("fixed", "")).strip().lower() in ("yes", "true", "1", "כן")))
    return out


def read_template_scheme(cfg):
    """[(row, group, cat)] from the template sheet (groups forward-filled). None if no template.
    The template is opened read-only: only the category scheme is taken from it."""
    ft = cfg["categories"].get("from_template")
    if not ft:
        return None, None
    path = project_path(cfg, ft["file"])
    if not os.path.isfile(path):
        log("WARNING: template workbook not found: %s — building without a template" % path)
        return None, None
    wb = openpyxl.load_workbook(path, read_only=True)
    if ft["sheet"] not in wb.sheetnames:
        wb.close()
        fail("template sheet %r not found in %s" % (ft["sheet"], path),
             "check categories.from_template.sheet (sheets: %s)" % ", ".join(wb.sheetnames))
    ws = wb[ft["sheet"]]
    gcol = openpyxl.utils.column_index_from_string(ft["groups_col"])
    ccol = openpyxl.utils.column_index_from_string(ft["cats_col"])
    rows, g = [], None
    for r in range(int(ft["first_row"]), int(ft["last_row"]) + 1):
        gv = ws.cell(r, gcol).value
        cv = ws.cell(r, ccol).value
        if gv not in (None, ""):
            g = str(gv).strip()
        if cv not in (None, ""):
            rows.append((r, g or "", str(cv).strip()))
    wb.close()
    return path, rows


def read_previous(path):
    """User text/category from the previous workbook's "לסיווג ידני" sheet ({id: (text, cat)})."""
    labels = {}
    if not os.path.isfile(path):
        return labels
    try:
        prev = openpyxl.load_workbook(path, read_only=True)
    except Exception as e:  # noqa: BLE001
        log("WARNING: could not open the previous workbook (%s); nothing preserved" % e)
        return labels
    try:
        if SHEET_MANUAL in prev.sheetnames:
            hd = None
            for row in prev[SHEET_MANUAL].iter_rows(values_only=True):
                if hd is None:
                    if row and row[0] == DB_HEADERS_HE["id"]:
                        hd = list(row)
                    continue
                if row and row[0]:
                    d = dict(zip(hd, row))
                    t = str(d.get(USER_TEXT_COL) or "").strip()
                    c = str(d.get(USER_CAT_COL) or "").strip()
                    if t or c:
                        labels[str(row[0])] = (t, c)
    except Exception as e:  # noqa: BLE001
        log("WARNING: could not read the previous workbook fully (%s)" % e)
    finally:
        prev.close()
    return labels


# ----------------------------------------------------------------------------- builder
class Builder(object):
    def __init__(self, cfg, db, template):
        self.cfg, self.db = cfg, db
        self.level = cfg.level
        self.sheets = set(SHEETS_BY_LEVEL[self.level])
        self.months = months(cfg)
        self.n = len(self.months)
        # month label shown in headers / dropdowns: the data's month_name when present
        self.mlabel = {}
        for m in self.months:
            names = db.loc[db["month"] == m, "month_name"]
            names = [x for x in names if x]
            self.mlabel[m] = names[0] if names else month_label(m)
        self.mlist = [self.mlabel[m] for m in self.months]
        th = cfg["thresholds"]
        self.hl_exp, self.hl_inc = th["highlight_expense_avg"], th["highlight_income_avg"]
        self.dyn_m = int(th.get("dynamic_merchant_rows", 250))
        self.dyn_t = int(th.get("dynamic_txn_rows", 500))
        self.top_n = int(th.get("top_merchants", 25))
        self.tall = int(th.get("hyperlink_rows", 45))
        self.template_path, self.template_rows = template
        self.wb = None
        self.E = db[(db["type"] == TYPE_EXPENSE) & db["_summed"]]
        self.I = db[(db["type"] == TYPE_INCOME) & db["_summed"]]
        self.rows_written = {}
        self.cat_stat_row = {}   # cat -> row in הוצאות
        self.inc_row = {}        # cat -> row in הכנסות
        self.cat_anchor = {}     # cat -> row in פירוט לפי קטגוריה
        self.merch_anchor = {}   # (cat, merchant) -> row in פירוט עסקאות
        self.merch_first = {}    # merchant -> first row in פירוט עסקאות
        self.merch_cell = {}     # (cat, merchant) -> row in פירוט לפי קטגוריה
        self.lists = {}
        self.newgrp_rows = {}
        self.pivot_meta = {}

    # ---- pandas helpers -------------------------------------------------------------
    def mv(self, mask, flag="_summed"):
        s = self.db[mask & self.db[flag]].groupby("month")["amount"].sum()
        return [float(s.get(m, 0.0)) for m in self.months]

    def cat_month_values(self, typ, cat, col="cat_tz"):
        return self.mv((self.db["type"] == typ) & (self.db[col] == cat))

    # ---- scheme -----------------------------------------------------------------------
    def build_scheme(self):
        cfg, db = self.cfg, self.db
        scheme = read_scheme_csv(scheme_path(cfg))
        if self.template_rows:
            tpl = [(g, c, False) for _, g, c in self.template_rows]
            known = {c for _, c, _ in tpl}
            scheme = tpl + [x for x in scheme if x[1] not in known]
        self.fixed = set(cfg.get("fixed_categories") or []) | {c for _, c, f in scheme if f}
        # every category name in play must be comma-free (FM5)
        names = {c for _, c, _ in scheme} | set(db.loc[db["_summed"], "cat_tz"]) | set(db.loc[db["_summed"], "cat_new"])
        bad = sorted(c for c in names if "," in c)
        if bad:
            fail("category names must not contain commas: %s" % "; ".join(bad),
                 "rename the category in the scheme / rules CSV (a comma breaks the rules file, FM5)")
        # expense categories present in the data, group taken from the data (classify canonicalised it)
        E = self.E
        grp_of = E.groupby("cat_tz")["group_tz"].agg(lambda s: s.mode().iloc[0] if len(s.mode()) else "")
        cat_tot = E.groupby("cat_tz")["amount"].sum()
        self.exp_cats = [(str(grp_of[c]), c) for c in cat_tot.index]
        self.dropped = [c for _, c, _ in scheme if c not in cat_tot.index and c not in set(self.I["cat_tz"])]
        gtot = {}
        for g, c in self.exp_cats:
            gtot[g] = gtot.get(g, 0.0) + float(cat_tot[c])
        self.group_order = sorted(gtot, key=lambda g: -gtot[g])
        self.cat_avg = {c: float(cat_tot[c]) / self.n for _, c in self.exp_cats}
        self.exp_rows = []
        for g in self.group_order:
            self.exp_rows += sorted([x for x in self.exp_cats if x[0] == g], key=lambda x: -self.cat_avg[x[1]])
        self.cat_list = [c for _, c in self.exp_rows]
        self.all_cats = []
        for _, c, _ in scheme:
            if c not in self.all_cats:
                self.all_cats.append(c)
        for c in self.cat_list:
            if c not in self.all_cats:
                self.all_cats.append(c)
        for c in list(self.I["cat_tz"].unique()) + list(GENERATED_CATEGORIES):
            if c not in self.all_cats:
                self.all_cats.append(c)
        # incomes: groups from the data, sorted by total desc / avg desc
        I = self.I
        igrp = I.groupby("cat_tz")["group_tz"].agg(lambda s: s.mode().iloc[0] if len(s.mode()) else "")
        itot = I.groupby("cat_tz")["amount"].sum()
        ig = {}
        for c in itot.index:
            ig[str(igrp[c])] = ig.get(str(igrp[c]), 0.0) + float(itot[c])
        self.inc_group_order = sorted(ig, key=lambda g: -ig[g])
        self.inc_avg = {c: float(itot[c]) / self.n for c in itot.index}
        self.inc_group = {c: str(igrp[c]) for c in itot.index}
        self.inc_rows = []
        for g in self.inc_group_order:
            self.inc_rows += sorted([(g, c) for c in itot.index if self.inc_group[c] == g], key=lambda x: -self.inc_avg[x[1]])
        # secondary scheme (deep): from the data's cat_new, universe optionally from a CSV
        sec = read_scheme_csv(project_path(cfg, cfg["categories"]["secondary_scheme"])) if cfg["categories"].get("secondary_scheme") else []
        E2 = E[E["cat_new"] != ""]
        g2 = E2.groupby("cat_new")["group_new"].agg(lambda s: s.mode().iloc[0] if len(s.mode()) else "")
        t2 = E2.groupby("cat_new")["amount"].sum()
        self.new_avg = {c: float(t2[c]) / self.n for c in t2.index}
        groups2 = {}
        for c in t2.index:
            groups2.setdefault(str(g2[c]), []).append(c)
        self.new_scheme = sorted([(g, sorted(cs, key=lambda c: -self.new_avg[c])) for g, cs in groups2.items()],
                                 key=lambda x: -sum(self.new_avg[c] for c in x[1]))
        self.new_all = [c for _, c, _ in sec] + [c for _, cs in self.new_scheme for c in cs if c not in {x[1] for x in sec}]

    # ---- 1. database -------------------------------------------------------------------
    def build_database(self):
        ws = rtl(self.wb.create_sheet(SHEET_DB))
        hdr(ws, 1, [DB_HEADERS_HE[k] for k in DB_COLUMNS])
        self.col_idx = {k: i + 1 for i, k in enumerate(DB_COLUMNS)}
        self.CL = {k: L(i) for k, i in self.col_idx.items()}
        for r, rec in enumerate(self.db.itertuples(index=False), start=2):
            d = rec._asdict()
            for k in DB_COLUMNS:
                v = d[k]
                if k == "amount":
                    v = float(v)
                elif k == "orig_amount":
                    v = None if pd.isna(d["orig_amount_f"]) else float(d["orig_amount_f"])
                elif v == "":
                    v = None
                c = put_text(ws, r, self.col_idx[k], v)
                if k in ("amount", "orig_amount"):
                    c.number_format = NUM2
        self.N = len(self.db) + 1
        tab = Table(displayName="DB", ref="A1:%s%d" % (L(len(DB_COLUMNS)), self.N))
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(tab)
        ws.freeze_panes = "A2"
        setw(ws, {self.CL["original_name"]: 34, self.CL["name_clean"]: 36, self.CL["details"]: 50,
                  self.CL["rule_note"]: 40, self.CL["source"]: 16, self.CL["pay"]: 22, self.CL["cat_tz"]: 26,
                  self.CL["cat_new"]: 24, self.CL["group_tz"]: 14, self.CL["group_new"]: 16, self.CL["type"]: 16})
        for nm, k in NAMED_RANGES.items():
            self.define_name(nm, "%s$%s$2:$%s$%d" % (sref(SHEET_DB), self.CL[k], self.CL[k], self.N))
        self.db_row = {rid: i + 2 for i, rid in enumerate(self.db["id"])}
        self.rows_written[SHEET_DB] = len(self.db)

    def define_name(self, name, ref):
        dn = DefinedName(name, attr_text=ref)
        try:
            self.wb.defined_names[name] = dn
        except TypeError:  # openpyxl < 3.1
            self.wb.defined_names.append(dn)

    # ---- 2. helper ----------------------------------------------------------------------
    def build_helper(self):
        pv = rtl(self.wb.create_sheet(SHEET_HELPER))
        pv["A1"] = "גיליון עזר: סכומים חודשיים לפי קטגוריה בנוסחאות SUMIFS על database (חודש ללא חיוב = 0). אין להזין ידנית."
        pv["A1"].font = GREY
        hdr(pv, 2, ["סכימה", "סוג", "קבוצה", "קטגוריה"] + self.months + ['סה"כ'])
        self.PV = {}
        self.prow = 3
        for g, c in self.exp_rows:
            self.pivot_row(SCHEME_TZ, TYPE_EXPENSE, g, c, "DBcat")
        for g, c in self.inc_rows:
            self.pivot_row(SCHEME_TZ, TYPE_INCOME, g, c, "DBcat")
        self.NS = {}
        for typ in (TYPE_SAVINGS, TYPE_INTERNAL, TYPE_CARD_DEBIT):
            for c in sorted(set(self.db.loc[(self.db["type"] == typ) & self.db["_in"], "cat_tz"])):
                self.NS[(typ, c)] = self.pivot_row(SCHEME_TZ, typ, typ, c, "DBcat")
        self.RB = {}
        for typ in (TYPE_REIMBURSABLE, TYPE_REIMBURSEMENT):
            for c in sorted(set(self.db.loc[(self.db["type"] == typ) & self.db["_in"], "cat_tz"])):
                self.RB[(typ, c)] = self.pivot_row(SCHEME_TZ, typ, "הוצאות בהחזר", c, "DBcat")
        if SHEET_PROPOSED in self.sheets:
            for g, cs in self.new_scheme:
                for c in cs:
                    self.pivot_row(SCHEME_NEW, TYPE_EXPENSE, g, c, "DBcat2")
        self.TOTAL_EXP = self.total_pivot("סה\"כ הוצאות (ישירות מ-database)", TYPE_EXPENSE)
        self.TOTAL_INC = self.total_pivot("סה\"כ הכנסות (ישירות מ-database)", TYPE_INCOME)
        pv.freeze_panes = "E3"
        setw(pv, {"A": 10, "B": 18, "C": 18, "D": 34})
        self.rows_written[SHEET_HELPER] = self.prow - 3

    def pivot_row(self, scheme, typ, group, cat, catrange):
        pv, r = self.wb[SHEET_HELPER], self.prow
        pv.cell(r, 1, scheme)
        pv.cell(r, 2, typ)
        pv.cell(r, 3, group)
        put_text(pv, r, 4, cat)
        flag = "DBsummed" if typ in (TYPE_EXPENSE, TYPE_INCOME) else "DBin"
        for j in range(self.n):
            col = L(5 + j)
            pv.cell(r, 5 + j, "=SUMIFS(DBamt,%s,$D%d,DBtype,$B%d,DBmonth,%s$2,%s,%s)" % (
                catrange, r, r, col, flag, q(YES))).number_format = NUM
        pv.cell(r, 5 + self.n, "=SUM(E%d:%s%d)" % (r, L(4 + self.n), r)).number_format = NUM
        self.PV[(scheme, typ, group, cat)] = r
        self.prow += 1
        return r

    def total_pivot(self, label, typ):
        pv, r = self.wb[SHEET_HELPER], self.prow
        pv.cell(r, 1, SCHEME_TZ)
        pv.cell(r, 2, typ)
        pv.cell(r, 3, 'סה"כ')
        pv.cell(r, 4, label)
        for j in range(self.n):
            pv.cell(r, 5 + j, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s$2,DBsummed,%s)" % (
                q(typ), L(5 + j), q(YES))).number_format = NUM
        pv.cell(r, 5 + self.n, "=SUM(E%d:%s%d)" % (r, L(4 + self.n), r)).number_format = NUM
        self.prow += 1
        return r

    def pv_formulas(self, prow):
        return ["=%s%s%d" % (sref(SHEET_HELPER), L(5 + j), prow) for j in range(self.n)]

    # ---- 3. statistic sheets ----------------------------------------------------------
    def setup_stat_sheet(self, name, title, legend):
        ws = rtl(self.wb.create_sheet(name))
        lay = Layout(self.n)
        ws["A1"] = title
        ws["A1"].font = TITLE
        ws["A1"].alignment = RIGHT
        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=lay.c_nz)
        ws.cell(1, lay.c_nz + 1, legend).font = GREY
        hdr(ws, 2, lay.headers(self.mlist))
        for i, w in enumerate(lay.widths(), 1):
            ws.column_dimensions[L(i)].width = w
        ws.freeze_panes = "C3"
        return ws, lay

    def stat_formulas(self, ws, lay, r):
        rng = lay.mrng(r)
        ws.cell(r, lay.c_total, "=SUM(%s)" % rng)
        ws.cell(r, lay.c_avg, "=AVERAGE(%s)" % rng)
        ws.cell(r, lay.c_nz, '=IFERROR(AVERAGEIF(%s,"<>0"),0)' % rng)
        ws.cell(r, lay.c_max, "=MAX(%s)" % rng)
        ws.cell(r, lay.c_maxm, "=INDEX(%s,MATCH(%s%d,%s,0))" % (lay.mhdr(), L(lay.c_max), r, rng))
        ws.cell(r, lay.c_min, "=MIN(%s)" % rng)
        ws.cell(r, lay.c_minm, "=INDEX(%s,MATCH(%s%d,%s,0))" % (lay.mhdr(), L(lay.c_min), r, rng))
        ws.cell(r, lay.c_med, "=MEDIAN(%s)" % rng)
        ws.cell(r, lay.c_std, "=IFERROR(_xlfn.STDEV.S(%s),0)" % rng)
        ws.cell(r, lay.c_cnt, '=COUNTIF(%s,"<>0")' % rng)
        for c in (lay.c_total, lay.c_avg, lay.c_nz, lay.c_max, lay.c_min, lay.c_med, lay.c_std):
            ws.cell(r, c).number_format = NUM

    def stat_row(self, ws, lay, r, group, cat, month_formulas, vals, note="", grey=False):
        put_text(ws, r, lay.c_group, group)
        put_text(ws, r, lay.c_cat, cat)
        for j in range(self.n):
            ws.cell(r, lay.m0 + j, month_formulas[j]).number_format = NUM
        self.stat_formulas(ws, lay, r)
        if vals is not None and any(abs(v) > 0.005 for v in vals):
            imax = max(range(self.n), key=lambda i: vals[i])
            imin = min(range(self.n), key=lambda i: vals[i])
            ws.cell(r, lay.c_max).comment = Comment("חודש המקסימום: %s (%s ₪)" % (self.mlist[imax], "{:,.0f}".format(vals[imax])), "ניתוח")
            ws.cell(r, lay.c_min).comment = Comment("חודש המינימום: %s (%s ₪)" % (self.mlist[imin], "{:,.0f}".format(vals[imin])), "ניתוח")
        put_text(ws, r, lay.c_note, note)
        for c in range(1, lay.ncols + 1):
            ws.cell(r, c).border = BORDER
            if grey:
                ws.cell(r, c).font = GREY

    def group_subtotal(self, ws, lay, r, group, rows):
        """Subtotal row (F11): month cells = '+' of the member rows; stats over the sums."""
        put_text(ws, r, lay.c_group, group)
        put_text(ws, r, lay.c_cat, 'סה"כ %s' % group)
        for j in range(self.n):
            col = L(lay.m0 + j)
            ws.cell(r, lay.m0 + j, ("=" + "+".join("%s%d" % (col, x) for x in rows)) if rows else 0).number_format = NUM
        self.stat_formulas(ws, lay, r)
        ws.cell(r, lay.c_note, "סכימת הקבוצה (שורות נסכמות בלבד)")
        for c in range(1, lay.ncols + 1):
            ws.cell(r, c).fill, ws.cell(r, c).font, ws.cell(r, c).border = SUBTOT, BOLD, BORDER

    def merge_groups(self, ws, lay, first, last):
        r = first
        while r <= last:
            g = ws.cell(r, 1).value
            r2 = r
            while r2 + 1 <= last and ws.cell(r2 + 1, 1).value == g:
                r2 += 1
            if r2 > r:
                ws.merge_cells(start_row=r, start_column=1, end_row=r2, end_column=1)
            ws.cell(r, 1).alignment, ws.cell(r, 1).fill, ws.cell(r, 1).font = CENTER, GRP, BOLD
            for c in range(1, lay.ncols + 1):
                ws.cell(r2, c).border = Border(top=_thin, left=_thin, right=_thin, bottom=_med)
            r = r2 + 1

    def total_row(self, ws, lay, r, label, formulas, note, vals=None):
        self.stat_row(ws, lay, r, 'סה"כ', label, formulas, vals, note)
        for c in range(1, lay.ncols + 1):
            ws.cell(r, c).fill, ws.cell(r, c).font = TOT, BOLD

    def apply_cf(self, ws, lay, first, last, hl):
        colour_scale(ws, "%s%d:%s%d" % (L(lay.c_avg), first, L(lay.c_avg), last))
        colour_scale(ws, "%s%d:%s%d" % (L(lay.c_nz), first, L(lay.c_nz), last))
        colour_scale(ws, "%s%d:%s%d" % (L(lay.c_total), first, L(lay.c_total), last), "A9D18E")
        heat(ws, "%s%d:%s%d" % (L(lay.m0), first, L(lay.m1), last))
        ws.conditional_formatting.add("B%d:%s%d" % (first, L(lay.c_avg), last),
                                      FormulaRule(formula=["$%s%d>=%s" % (L(lay.c_avg), first, hl)], font=RED))

    def build_expenses(self):
        db, n = self.db, self.n
        legend = ('מקרא: כחול = ממוצע גבוה יותר · ירוק = סה"כ גבוה יותר · כתום בחודשים = גבוה יותר · אדום מודגש = ממוצע ≥ '
                  '{:,.0f} ₪ · "ממוצע ללא אפס" = ממוצע רק על חודשים עם חיוב · לחיצה על שם תת-קטגוריה → פירוט בתי העסק · '
                  'אפור = לא נסכם'.format(self.hl_exp))
        title = ("הוצאות שוטפות — %s–%s · ממוין: קבוצות לפי סה\"כ, תתי-קטגוריות לפי ממוצע (חודש = חודש החיוב בבנק; "
                 "חודש ללא חיוב = 0; הכל נוסחאות מ-database)" % (self.mlist[0], self.mlist[-1]))
        ws, lay = self.setup_stat_sheet(SHEET_EXP, title, legend)
        self.lay_exp = lay
        r, subtotal_rows, grp_rows, cur = 3, [], [], None
        for g, cat in self.exp_rows:
            if cur is not None and g != cur:
                self.group_subtotal(ws, lay, r, cur, list(grp_rows))
                subtotal_rows.append(r)
                r += 1
                grp_rows = []
            cur = g
            note = ""
            if cat == UNKNOWN_CAT:
                note = 'ראה לשונית "לסיווג ידני" – ניתן להשלים שם סיווג' if SHEET_MANUAL in self.sheets else "שורות ללא כלל סיווג"
            self.stat_row(ws, lay, r, g, cat, self.pv_formulas(self.PV[(SCHEME_TZ, TYPE_EXPENSE, g, cat)]),
                          self.cat_month_values(TYPE_EXPENSE, cat), note)
            self.cat_stat_row[cat] = r
            grp_rows.append(r)
            r += 1
        if cur is not None:
            self.group_subtotal(ws, lay, r, cur, list(grp_rows))
            subtotal_rows.append(r)
            r += 1
        self.LAST_EXP = r - 1
        if self.LAST_EXP >= 3:
            self.merge_groups(ws, lay, 3, self.LAST_EXP)
            self.apply_cf(ws, lay, 3, self.LAST_EXP, self.hl_exp)
        self.TR = r
        self.total_row(ws, lay, r, TOTAL_EXP_LABEL, self.pv_formulas(self.TOTAL_EXP),
                       "סכום כל שורות database עם סוג=הוצאה ונסכם=כן", self.mv(db["type"] == TYPE_EXPENSE))
        self.CHK = r + 1
        tc = L(lay.c_total)
        ws.cell(self.CHK, 2, "בדיקה: סכום שורות הקטגוריות (ללא שורות סה\"כ קבוצה)").font = GREY
        ws.cell(self.CHK, lay.c_total, "=SUM(%s3:%s%d)" % (tc, tc, max(self.LAST_EXP, 3)) +
                "".join("-%s%d" % (tc, x) for x in subtotal_rows)).number_format = NUM
        ws.cell(self.CHK, lay.c_note, '=IF(ABS(%s%d-%s%d)<1,"✓ תואם לסה""כ","✗ פער: "&TEXT(%s%d-%s%d,"#,##0"))' % (
            tc, self.CHK, tc, self.TR, tc, self.CHK, tc, self.TR))
        # non-summed block: savings / internal / card debits
        r += 3
        ws.cell(r, 1, "לא נסכם")
        ws.cell(r, 2, 'חיסכון, השקעות, העברות פנימיות וחיובי כרטיסי אשראי (מוצג בלבד – לא חלק מסה"כ ההוצאות)').font = BOLD
        r += 1
        hdr(ws, r, lay.headers(self.mlist))
        r += 1
        ns_first = r
        card_note = "שורה אינפורמטיבית, לא נסכמת: חיוב הכרטיס בבנק. ההוצאות עצמן נספרות לפי בתי העסק."
        order = sorted(self.NS.items(), key=lambda kv: -sum(self.mv((db["type"] == kv[0][0]) & (db["cat_tz"] == kv[0][1]), "_in")))
        for (typ, c), prw in order:
            vals = self.mv((db["type"] == typ) & (db["cat_tz"] == c), "_in")
            self.stat_row(ws, lay, r, typ, c, self.pv_formulas(prw), vals, card_note if typ == TYPE_CARD_DEBIT else "לא נסכם", grey=True)
            r += 1
        if r - 1 >= ns_first:
            self.merge_groups(ws, lay, ns_first, r - 1)
        # reimbursables with an open balance
        self.RB_BAL = None
        if self.RB:
            r += 2
            ws.cell(r, 1, "לא נסכם")
            ws.cell(r, 2, "הוצאות בהחזר — הוצאות מול החזרים ומאזן פתוח; מחוץ לממוצעי ההוצאות וההכנסות").font = BOLD
            r += 1
            hdr(ws, r, lay.headers(self.mlist))
            r += 1
            rb_first = r
            for (typ, c), prw in sorted(self.RB.items()):
                vals = self.mv((db["type"] == typ) & (db["cat_tz"] == c), "_in")
                self.stat_row(ws, lay, r, typ, c, self.pv_formulas(prw), vals,
                              "לא נסכם – " + ("יצא מהחשבון, צפוי החזר" if typ == TYPE_REIMBURSABLE else "החזר שהתקבל"), grey=True)
                r += 1
            self.merge_groups(ws, lay, rb_first, r - 1)
            ws.cell(r, 1, "מאזן")
            ws.cell(r, 2, "מאזן פתוח (הוצאות בהחזר − החזרים שהתקבלו; חיובי = עדיין מגיע לכם)")
            for j, m in enumerate(self.months):
                ws.cell(r, lay.m0 + j, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBin,%s)-SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBin,%s)" % (
                    q(TYPE_REIMBURSABLE), q(m), q(YES), q(TYPE_REIMBURSEMENT), q(m), q(YES))).number_format = NUM
            ws.cell(r, lay.c_total, "=SUM(%s)" % lay.mrng(r)).number_format = NUM
            for c in range(1, lay.ncols + 1):
                ws.cell(r, c).fill, ws.cell(r, c).font = TOT, BOLD
            self.RB_BAL = r
        self.rows_written[SHEET_EXP] = r

    def build_incomes(self):
        db = self.db
        ws, lay = self.setup_stat_sheet(SHEET_INC, "הכנסות לפי מקור — %s–%s · ממוין לפי ממוצע בתוך כל קבוצה" % (self.mlist[0], self.mlist[-1]),
                                        "מקרא: כחול = ממוצע גבוה יותר · צבע בחודשים = כהה יותר = גבוה יותר · אדום מודגש = ממוצע ≥ {:,.0f} ₪".format(self.hl_inc))
        self.lay_inc = lay
        r, grp_rows, cur = 3, [], None
        for g, c in self.inc_rows:
            if cur is not None and g != cur:
                self.group_subtotal(ws, lay, r, cur, list(grp_rows))
                r += 1
                grp_rows = []
            cur = g
            self.stat_row(ws, lay, r, g, c, self.pv_formulas(self.PV[(SCHEME_TZ, TYPE_INCOME, g, c)]),
                          self.cat_month_values(TYPE_INCOME, c), "")
            self.inc_row[c] = r
            grp_rows.append(r)
            r += 1
        if cur is not None:
            self.group_subtotal(ws, lay, r, cur, list(grp_rows))
            r += 1
        self.LAST_INC = r - 1
        if self.LAST_INC >= 3:
            self.merge_groups(ws, lay, 3, self.LAST_INC)
            self.apply_cf(ws, lay, 3, self.LAST_INC, self.hl_inc)
        self.TRI = r
        self.total_row(ws, lay, r, TOTAL_INC_LABEL, self.pv_formulas(self.TOTAL_INC), "", self.mv(db["type"] == TYPE_INCOME))
        r += 1
        ws.cell(r, 2, 'סה"כ הוצאות שוטפות (מלשונית הוצאות)')
        le = self.lay_exp
        for j in range(self.n):
            ws.cell(r, lay.m0 + j, "=%s%s%d" % (sref(SHEET_EXP), L(le.m0 + j), self.TR)).number_format = NUM
        self.simple_stats(ws, lay, r)
        r += 1
        self.BAL = r
        ws.cell(r, 2, "מאזן חודשי (הכנסות − הוצאות שוטפות)")
        for j in range(self.n):
            col = L(lay.m0 + j)
            ws.cell(r, lay.m0 + j, "=%s%d-%s%d" % (col, self.TRI, col, r - 1)).number_format = NUM
        self.simple_stats(ws, lay, r)
        for c in range(1, lay.ncols + 1):
            ws.cell(r, c).fill, ws.cell(r, c).font = TOT, BOLD
        r += 1
        ws.cell(r, 2, "החזרי הוצאות שהתקבלו – לא נסכם בהכנסות")
        for j, m in enumerate(self.months):
            ws.cell(r, lay.m0 + j, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBin,%s)" % (q(TYPE_REIMBURSEMENT), q(m), q(YES))).number_format = NUM
        self.simple_stats(ws, lay, r, grey=True)
        r += 1
        self.SAV = r
        ws.cell(r, 2, "חיסכון והשקעות בפועל – לא נסכם")
        for j, m in enumerate(self.months):
            ws.cell(r, lay.m0 + j, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBin,%s)" % (q(TYPE_SAVINGS), q(m), q(YES))).number_format = NUM
        self.simple_stats(ws, lay, r, grey=True)
        self.rows_written[SHEET_INC] = r

    def simple_stats(self, ws, lay, r, grey=False):
        rng = lay.mrng(r)
        ws.cell(r, lay.c_total, "=SUM(%s)" % rng).number_format = NUM
        ws.cell(r, lay.c_avg, "=AVERAGE(%s)" % rng).number_format = NUM
        ws.cell(r, lay.c_nz, '=IFERROR(AVERAGEIF(%s,"<>0"),0)' % rng).number_format = NUM
        if grey:
            for c in range(2, lay.ncols + 1):
                ws.cell(r, c).font = GREY

    def build_proposed(self):
        ws, lay = self.setup_stat_sheet(SHEET_PROPOSED, 'הוצאות לפי הסכימה המוצעת — ממוין · סה"כ זהה ללשונית הוצאות',
                                        'מקרא: כחול = ממוצע גבוה יותר · ירוק = סה"כ · כתום בחודשים = גבוה יותר')
        self.lay_new = lay
        r = 3
        self.newcat_row = {}
        for g, cs in self.new_scheme:
            rows = []
            for c in cs:
                self.stat_row(ws, lay, r, g, c, self.pv_formulas(self.PV[(SCHEME_NEW, TYPE_EXPENSE, g, c)]),
                              self.cat_month_values(TYPE_EXPENSE, c, "cat_new"))
                self.newcat_row[c] = r
                rows.append(r)
                r += 1
            self.group_subtotal(ws, lay, r, g, rows)
            r += 1
        last = r - 1
        if last >= 3:
            self.merge_groups(ws, lay, 3, last)
            self.apply_cf(ws, lay, 3, last, self.hl_exp)
        else:
            ws.cell(r, 2, "(אין ערכים בעמודות הסכימה המוצעת ב-database)").font = GREY
            r += 1
        self.TRN = r
        self.total_row(ws, lay, r, 'סה"כ הוצאות שוטפות (ישירות מ-database)', self.pv_formulas(self.TOTAL_EXP), "")
        r += 2
        ws.cell(r, 2, "סיכום לפי קבוצה (ממוין)").font = BOLD
        r += 1
        hdr(ws, r, ["", "קבוצה"] + self.mlist + ['סה"כ', "ממוצע חודשי", "% מסה\"כ"])
        r += 1
        first = r
        for g, _ in self.new_scheme:
            put_text(ws, r, 2, g)
            for j, m in enumerate(self.months):
                ws.cell(r, lay.m0 + j, "=SUMIFS(DBamt,DBtype,%s,DBgroup2,$B%d,DBmonth,%s,DBsummed,%s)" % (
                    q(TYPE_EXPENSE), r, q(m), q(YES))).number_format = NUM
            ws.cell(r, lay.c_total, "=SUM(%s)" % lay.mrng(r)).number_format = NUM
            ws.cell(r, lay.c_avg, "=AVERAGE(%s)" % lay.mrng(r)).number_format = NUM
            ws.cell(r, lay.c_nz, "=IFERROR(%s%d/$%s$%d,0)" % (L(lay.c_total), r, L(lay.c_total), self.TRN)).number_format = PCT
            self.newgrp_rows[g] = r
            r += 1
        if r > first:
            colour_scale(ws, "%s%d:%s%d" % (L(lay.c_avg), first, L(lay.c_avg), r - 1))
        self.rows_written[SHEET_PROPOSED] = r

    # ---- 4. lists (hidden) -------------------------------------------------------------
    def build_lists(self):
        db = self.db
        wl = rtl(self.wb.create_sheet(SHEET_LISTS))
        persons = sorted(set(self.E["person"]) - {""})
        pays = sorted(set(self.E["pay"]) - {""})
        self.lists = {
            "L_tz_g": list(self.group_order), "L_tz_c": list(self.cat_list),
            "L_new_g": [g for g, _ in self.new_scheme], "L_new_c": [c for _, cs in self.new_scheme for c in cs],
            "L_months": [ALL_MONTHS] + self.mlist, "L_pay": [ALL] + pays, "L_person": [ALL] + persons,
            "L_scheme": [SCHEME_TZ] + ([SCHEME_NEW] if SHEET_PROPOSED in self.sheets and self.new_scheme else []),
            "L_level": ["קבוצה", "תת-קטגוריה"], "L_allcats": list(self.all_cats),
        }
        self.list_col = 0
        for nm, items in self.lists.items():
            self.write_list(wl, nm, items)
        wl.sheet_state = "hidden"

    def write_list(self, wl, name, items):
        self.list_col += 1
        col = L(self.list_col)
        wl.cell(1, self.list_col, name).font = BOLD
        items = list(items) or ["—"]
        for i, v in enumerate(items):
            put_text(wl, 2 + i, self.list_col, v)
        self.define_name(name, "%s$%s$2:$%s$%d" % (sref(SHEET_LISTS), col, col, 1 + len(items)))
        wl.column_dimensions[col].width = 26

    # ---- 5. פילוח ----------------------------------------------------------------------
    def build_pivot(self):
        wd = rtl(self.wb.create_sheet(SHEET_PIVOT))
        wd["B1"] = "פילוח אינטראקטיבי — בחר בתאים הירוקים (C4:C9); הטבלאות והגרף מתעדכנים מיד (נוסחאות דינמיות על database)"
        wd["B1"].font = TITLE
        wd["B2"] = ('איך משתמשים: (1) סכימה ורמה, (2) בחר קבוצה/תת-קטגוריה מהרשימה הנפתחת (משתנה לפי הסכימה והרמה; אפשר גם להקליד), '
                    '(3) סנן חודש / אמצעי תשלום / אדם. "הכל" = ללא סינון.')
        wd["B2"].font = GREY
        default_sel = self.cat_list[0] if self.cat_list else ""
        inputs = [("סכימה", SCHEME_TZ, "L_scheme"), ("רמה", "תת-קטגוריה", "L_level"),
                  ("בחירה (קבוצה / תת-קטגוריה)", default_sel, None), ("חודש", ALL_MONTHS, "L_months"),
                  ("אמצעי תשלום", ALL, "L_pay"), ("אדם", ALL, "L_person")]
        for i, (lab, default, lst) in enumerate(inputs):
            rr = 4 + i
            wd.cell(rr, 2, lab).font = BOLD
            c = put_text(wd, rr, 3, default)
            c.fill, c.border, c.font = INPUT, BORDER, BOLD
            f1 = ("=%s" % lst) if lst else ('=INDIRECT(IF($C$4=%s,IF($C$5="קבוצה","L_tz_g","L_tz_c"),IF($C$5="קבוצה","L_new_g","L_new_c")))' % q(SCHEME_TZ))
            dv = DataValidation(type="list", formula1=f1, allow_blank=False)
            wd.add_data_validation(dv)
            dv.add("C%d" % rr)
        cond = ('((DBtype=%s)*(DBsummed=%s)*IF($C$4=%s,IF($C$5="קבוצה",DBgroup=$C$6,DBcat=$C$6),IF($C$5="קבוצה",DBgroup2=$C$6,DBcat2=$C$6))'
                '*(($C$7=%s)+(DBmname=$C$7))*(($C$8=%s)+(DBpay=$C$8))*(($C$9=%s)+(DBperson=$C$9)))'
                % (q(TYPE_EXPENSE), q(YES), q(SCHEME_TZ), q(ALL_MONTHS), q(ALL), q(ALL)))
        kpis = [('סה"כ לבחירה', "=SUMPRODUCT(%s*DBamt)" % cond, NUM),
                ("מס' עסקאות", "=SUMPRODUCT(%s)" % cond, "0"),
                ("ממוצע לעסקה", "=IFERROR(F4/F5,0)", NUM),
                ("ממוצע חודשי (÷%d כשנבחרו כל החודשים)" % self.n, "=IF($C$7=%s,F4/%d,F4)" % (q(ALL_MONTHS), self.n), NUM),
                ("% מכלל ההוצאות באותם חודשים",
                 "=IFERROR(F4/SUMPRODUCT((DBtype=%s)*(DBsummed=%s)*(($C$7=%s)+(DBmname=$C$7))*DBamt),0)" % (q(TYPE_EXPENSE), q(YES), q(ALL_MONTHS)), PCT)]
        for i, (lab, f, fmt) in enumerate(kpis):
            rr = 4 + i
            wd.cell(rr, 5, lab).font = BOLD
            c = wd.cell(rr, 6, ArrayFormula("F%d" % rr, f) if "SUMPRODUCT" in f else f)
            c.number_format, c.fill, c.border, c.font = fmt, TOT, BORDER, BOLD
        T0 = 24
        wd.cell(T0 - 2, 2, "בתי עסק בבחירה — ממוין מהגבוה לנמוך · לחיצה על ↗ פותחת את שורות העסקאות של בית העסק").font = BOLD
        hdr(wd, T0 - 1, ["בית עסק (שם מובן)", 'סה"כ ₪', "מס' עסקאות", "ממוצע לעסקה", "% מהבחירה", "↗"], start=2)
        nm, nt = self.dyn_m, self.dyn_t
        mt = ("=IFERROR(INDEX(_xlfn.LET(_xlpm.c," + cond + ",_xlpm.fn,_xlfn._xlws.FILTER(DBname,_xlpm.c>0),_xlpm.fa,_xlfn._xlws.FILTER(DBamt,_xlpm.c>0),"
              "_xlpm.u,_xlfn.UNIQUE(_xlpm.fn),_xlpm.m,--(TRANSPOSE(_xlpm.u)=_xlpm.fn),_xlpm.s,MMULT(TRANSPOSE(_xlpm.m),_xlpm.fa),_xlpm.n,MMULT(TRANSPOSE(_xlpm.m),_xlpm.fa*0+1),"
              "_xlpm.t,SUM(_xlpm.fa),_xlfn.SORTBY(CHOOSE({1,2,3,4,5},_xlpm.u,_xlpm.s,_xlpm.n,_xlpm.s/_xlpm.n,_xlpm.s/_xlpm.t),_xlpm.s,-1)),"
              "_xlfn.SEQUENCE(%d),{1,2,3,4,5}),\"\")" % nm)
        wd.cell(T0, 2, ArrayFormula("B%d:F%d" % (T0, T0 + nm - 1), mt))
        for rr in range(T0, T0 + nm):
            wd.cell(rr, 3).number_format, wd.cell(rr, 4).number_format = NUM, "0"
            wd.cell(rr, 5).number_format, wd.cell(rr, 6).number_format = NUM, PCT
        colour_scale(wd, "C%d:C%d" % (T0, T0 + nm - 1))
        wd.cell(T0 - 2, 8, "כל העסקאות בבחירה — ממוין לפי סכום").font = BOLD
        hdr(wd, T0 - 1, ["תאריך עסקה", "בית עסק", "סכום ₪", "קטגוריה (תזרים)", "אמצעי תשלום", "פרטים"], start=8)
        tt = ("=IFERROR(INDEX(_xlfn.LET(_xlpm.c," + cond + ",_xlpm.fa,_xlfn._xlws.FILTER(DBamt,_xlpm.c>0),"
              "_xlfn.SORTBY(_xlfn._xlws.FILTER(CHOOSE({1,2,3,4,5,6},DBdate,DBname,DBamt,DBcat,DBpay,DBdet),_xlpm.c>0),_xlpm.fa,-1)),"
              "_xlfn.SEQUENCE(%d),{1,2,3,4,5,6}),\"\")" % nt)
        wd.cell(T0, 8, ArrayFormula("H%d:M%d" % (T0, T0 + nt - 1), tt))
        for rr in range(T0, T0 + nt):
            wd.cell(rr, 10).number_format = NUM
        setw(wd, {"A": 3, "B": 34, "C": 16, "D": 12, "E": 36, "F": 16, "G": 4, "H": 12, "I": 36, "J": 11, "K": 24, "L": 20, "M": 60})
        ch = BarChart()
        ch.type, ch.title, ch.width, ch.height, ch.legend = "bar", "15 בתי העסק הגדולים בבחירה", 20, 9.5, None
        ch.y_axis.numFmt = "#,##0"
        ch.add_data(Reference(wd, min_col=3, min_row=T0 - 1, max_row=T0 + 14), titles_from_data=True)
        ch.set_categories(Reference(wd, min_col=2, min_row=T0, max_row=T0 + 14))
        wd.add_chart(ch, "H3")
        self.T0 = T0
        self.pivot_meta = {"inputs": "C4:C9", "kpi_total": "F4", "kpi_count": "F5", "table_row": T0}
        self.rows_written[SHEET_PIVOT] = T0 + nt - 1

    def finish_pivot_links(self):
        """HYPERLINK column (↗) from the merchant table to the merchant's transaction block."""
        wl, wd, T0 = self.wb[SHEET_LISTS], self.wb[SHEET_PIVOT], self.T0
        keys = [("%s||%s" % (c, m), a) for (c, m), a in self.merch_anchor.items()]
        firsts = list(self.merch_first.items())
        self.write_list(wl, "A_keys", [k for k, _ in keys])
        self.write_list(wl, "A_rows", [a for _, a in keys])
        self.write_list(wl, "A_first_keys", [k for k, _ in firsts])
        self.write_list(wl, "A_first_rows", [a for _, a in firsts])
        tgt = "#'%s'!A" % SHEET_TXNS
        for rr in range(T0, T0 + self.dyn_m):
            wd.cell(rr, 7, ('=IF(B{r}="","",IFERROR(HYPERLINK({t}&INDEX(A_rows,MATCH($C$6&"||"&B{r},A_keys,0))&":N"&(INDEX(A_rows,MATCH($C$6&"||"&B{r},A_keys,0))+{k}),"↗"),'
                            'IFERROR(HYPERLINK({t}&INDEX(A_first_rows,MATCH(B{r},A_first_keys,0))&":N"&(INDEX(A_first_rows,MATCH(B{r},A_first_keys,0))+{k}),"↗"),"")))'
                            ).format(r=rr, t=q(tgt), k=self.tall)).font = LINK

    # ---- 6. פירוט לפי קטגוריה -----------------------------------------------------------
    def link(self, cell, sheet, ref, text=None, tall=False):
        if tall and ":" not in ref:
            m = re.match(r"([A-Z]+)(\d+)$", ref)
            ref = "%s%s:%s%d" % (m.group(1), m.group(2), L(max(14, 7 + self.n)), int(m.group(2)) + self.tall)
        cell.hyperlink = "#'%s'!%s" % (sheet, ref)
        cell.font = LINK
        if text:
            cell.value = text

    def build_by_category(self):
        E, n = self.E, self.n
        wc = rtl(self.wb.create_sheet(SHEET_BY_CAT))
        width = max(14, 7 + n)      # block width; also the column of the ↩ link and of the tall hyperlink ranges
        wc["A1"] = "פירוט בתי עסק לכל תת-קטגוריה (סכימת תזרים) — ממוין מהגבוה לנמוך; ערכים בנוסחאות SUMIFS על database"
        wc["A1"].font = TITLE
        wc.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        wc["A2"] = "אינדקס (לחץ לקפיצה):"
        wc["A2"].font = BOLD
        merch = E.groupby(["cat_tz", "name_clean"])["amount"].sum().reset_index()
        r = 3 + (len(self.cat_list) + 3) // 4 + 2
        le = self.lay_exp
        for g, cat in self.exp_rows:
            self.cat_anchor[cat] = r
            for c in range(1, width + 1):
                wc.cell(r, c).fill = GRP
            put_text(wc, r, 1, "%s / %s" % (g, cat)).font = Font(bold=True, size=12)
            wc.cell(r, 6, 'סה"כ:')
            wc.cell(r, 7, "=%s%s%d" % (sref(SHEET_EXP), L(le.c_total), self.cat_stat_row[cat])).number_format = NUM
            wc.cell(r, 8, "ממוצע חודשי:")
            wc.cell(r, 9, "=%s%s%d" % (sref(SHEET_EXP), L(le.c_avg), self.cat_stat_row[cat])).number_format = NUM
            self.link(wc.cell(r, width), SHEET_EXP, "B%d" % self.cat_stat_row[cat], "↩ חזרה להוצאות", tall=True)
            r += 1
            hdr(wc, r, ["בית עסק (שם מובן)", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס", "מס' עסקאות", "ממוצע לעסקה", "% מהקטגוריה"] + self.mlist)
            r += 1
            sub = E[E["cat_tz"] == cat]
            trips = [t for t in sub["trip"].unique() if t]
            if trips:
                vt = sub.groupby("trip")["amount"].sum().sort_values(ascending=False)
                for trip in vt.index:
                    for c in range(1, width + 1):
                        wc.cell(r, c).fill = GRP
                        wc.cell(r, c).border = Border(top=_med)
                    put_text(wc, r, 1, trip if trip else "(ללא שיוך לנסיעה)").font = Font(bold=True, size=12)
                    if "לסיווג" in trip:
                        wc.cell(r, width, 'השלם מדינה בלשונית "לסיווג ידני"').font = RED
                    r += 1
                    tcrit = ",DBtrip,%s" % q(trip)
                    names = sub[sub["trip"] == trip].groupby("name_clean")["amount"].sum().sort_values(ascending=False).index.tolist()
                    r = self.merchant_rows(wc, r, cat, names, tcrit)
                    base = "DBcat,%s,DBtype,%s,DBsummed,%s%s" % (q(cat), q(TYPE_EXPENSE), q(YES), tcrit)
                    put_text(wc, r, 1, 'סה"כ %s' % trip)
                    wc.cell(r, 2, "=SUMIFS(DBamt,%s)" % base).number_format = NUM
                    wc.cell(r, 3, "=B%d/%d" % (r, n)).number_format = NUM
                    wc.cell(r, 4, '=IFERROR(AVERAGEIF(%s%d:%s%d,"<>0"),0)' % (L(8), r, L(7 + n), r)).number_format = NUM
                    wc.cell(r, 5, "=COUNTIFS(%s)" % base)
                    wc.cell(r, 6, "=IFERROR(B%d/E%d,0)" % (r, r)).number_format = NUM
                    wc.cell(r, 7, "=IFERROR(B%d/$G$%d,0)" % (r, self.cat_anchor[cat])).number_format = PCT
                    for j, m in enumerate(self.months):
                        wc.cell(r, 8 + j, "=SUMIFS(DBamt,%s,DBmonth,%s)" % (base, q(m))).number_format = NUM
                    for c in range(1, width + 1):
                        wc.cell(r, c).fill, wc.cell(r, c).font = TOT, BOLD
                        wc.cell(r, c).border = Border(top=_thin, bottom=_med)
                    r += 2
            else:
                names = merch[merch["cat_tz"] == cat].sort_values("amount", ascending=False)["name_clean"].tolist()
                if not names:
                    wc.cell(r, 1, "(אין עסקאות בחלון)").font = GREY
                    r += 1
                r = self.merchant_rows(wc, r, cat, names)
            r += 2
        for i, cat in enumerate(self.cat_list):
            self.link(wc.cell(3 + i // 4, 1 + 3 * (i % 4)), SHEET_BY_CAT, "A%d" % self.cat_anchor[cat], cat, tall=True)
        setw(wc, {"A": 40, "B": 14, "C": 13, "D": 13, "E": 10, "F": 12, "G": 11})
        for j in range(n):
            wc.column_dimensions[L(8 + j)].width = 11
        wc.column_dimensions[L(width)].width = 18
        we = self.wb[SHEET_EXP]
        for cat, rr in self.cat_stat_row.items():
            self.link(we.cell(rr, 2), SHEET_BY_CAT, "A%d" % self.cat_anchor[cat], tall=True)
        we.cell(2, 2).comment = Comment('לחיצה על שם תת-קטגוריה פותחת את פירוט בתי העסק שלה (לשונית "פירוט לפי קטגוריה"). '
                                        'לפילוח חופשי לפי חודש/כרטיס/אדם → לשונית "פילוח".', "ניתוח")
        self.rows_written[SHEET_BY_CAT] = r

    def merchant_rows(self, wc, r, cat, names, extra=""):
        n, f0 = self.n, r
        for nm in names:
            self.merch_cell[(cat, nm)] = r
            put_text(wc, r, 1, nm)
            base = "DBname,$A%d,DBcat,%s,DBtype,%s,DBsummed,%s%s" % (r, q(cat), q(TYPE_EXPENSE), q(YES), extra)
            wc.cell(r, 2, "=SUMIFS(DBamt,%s)" % base).number_format = NUM
            wc.cell(r, 3, "=B%d/%d" % (r, n)).number_format = NUM
            wc.cell(r, 4, '=IFERROR(AVERAGEIF(%s%d:%s%d,"<>0"),0)' % (L(8), r, L(7 + n), r)).number_format = NUM
            wc.cell(r, 5, "=COUNTIFS(%s)" % base)
            wc.cell(r, 6, "=IFERROR(B%d/E%d,0)" % (r, r)).number_format = NUM
            wc.cell(r, 7, "=IFERROR(B%d/$G$%d,0)" % (r, self.cat_anchor[cat])).number_format = PCT
            for j, m in enumerate(self.months):
                wc.cell(r, 8 + j, "=SUMIFS(DBamt,%s,DBmonth,%s)" % (base, q(m))).number_format = NUM
            for c in range(1, 8 + n):
                wc.cell(r, c).border = BORDER
            r += 1
        if names:
            colour_scale(wc, "B%d:B%d" % (f0, r - 1))
            colour_scale(wc, "C%d:D%d" % (f0, r - 1))
            heat(wc, "%s%d:%s%d" % (L(8), f0, L(7 + n), r - 1))
        return r

    # ---- 7. פירוט עסקאות -----------------------------------------------------------------
    def build_txns(self):
        E = self.E
        wt = rtl(self.wb.create_sheet(SHEET_TXNS))
        wt["A1"] = ('פירוט עסקאות לפי קטגוריה ← בית עסק — כל שורות ההוצאה שנספרו (ערכים מקושרים ללשונית database). '
                    'מגיעים לכאן בלחיצה על בית עסק ב"פירוט לפי קטגוריה", ב"הוצאות משתנות" או ב"פילוח".')
        wt["A1"].font = TITLE
        wt.merge_cells("A1:J1")
        tcols = ["txn_date", "charge_date", "original_name", "amount", "pay", "person", "month_name", "details", "rule_note", "id"]
        merch = E.groupby(["cat_tz", "name_clean"])["amount"].sum().reset_index()
        r = 3
        for g, cat in self.exp_rows:
            names = merch[merch["cat_tz"] == cat].sort_values("amount", ascending=False)["name_clean"].tolist()
            for nm in names:
                self.merch_anchor[(cat, nm)] = r
                self.merch_first.setdefault(nm, r)
                for c in range(1, 11):
                    wt.cell(r, c).fill = GRP
                    wt.cell(r, c).border = Border(top=_med)
                put_text(wt, r, 1, "%s / %s / %s" % (g, cat, nm)).font = Font(bold=True, size=12)
                self.link(wt.cell(r, 9), SHEET_BY_CAT, "A%d" % self.cat_anchor[cat], "↩ לפירוט הקטגוריה", tall=True)
                r += 1
                hdr(wt, r, [DB_HEADERS_HE[k] for k in tcols])
                r += 1
                f0 = r
                rows = E[(E["cat_tz"] == cat) & (E["name_clean"] == nm)].sort_values("amount", ascending=False)
                for t in rows.itertuples():
                    dr = self.db_row[t.id]
                    for c, k in enumerate(tcols, 1):
                        wt.cell(r, c, "=%s%s%d" % (sref(SHEET_DB), self.CL[k], dr)).border = BORDER
                    wt.cell(r, 4).number_format = NUM2
                    r += 1
                wt.cell(r, 3, 'סה"כ').font = BOLD
                wt.cell(r, 4, "=SUM(D%d:D%d)" % (f0, max(f0, r - 1))).number_format = NUM2
                wt.cell(r, 4).font = BOLD
                for c in range(1, 11):
                    wt.cell(r, c).fill = TOT
                r += 2
        setw(wt, {"A": 12, "B": 12, "C": 34, "D": 12, "E": 22, "F": 14, "G": 12, "H": 50, "I": 40, "J": 12})
        wc = self.wb[SHEET_BY_CAT]
        for (cat, nm), rr in self.merch_cell.items():
            if (cat, nm) in self.merch_anchor:
                self.link(wc.cell(rr, 1), SHEET_TXNS, "A%d" % self.merch_anchor[(cat, nm)], tall=True)
        self.rows_written[SHEET_TXNS] = r

    # ---- 8. הוצאות משתנות לפי בית עסק ------------------------------------------------------
    def build_variable(self):
        E, n = self.E, self.n
        wv = rtl(self.wb.create_sheet(SHEET_VARIABLE))
        width = 10 + n
        wv["A1"] = "הוצאות משתנות לפי בית עסק — כל בתי העסק שאינם בקטגוריות קבועות, ממוין מהגבוה לנמוך (יש מסננים בכותרת)"
        wv["A1"].font = TITLE
        wv.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
        wv["A2"] = "קטגוריות קבועות (לא כאן): " + (", ".join(sorted(self.fixed)) or "(לא הוגדרו fixed_categories)")
        wv["A2"].font = GREY
        wv.merge_cells(start_row=2, start_column=1, end_row=2, end_column=width)
        hdr(wv, 3, ["#", "בית עסק (שם מובן)", "תת-קטגוריה", "קבוצה", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס",
                    "מס' עסקאות", "ממוצע לעסקה", "% מהמשתנות"] + self.mlist)
        var = E[~E["cat_tz"].isin(self.fixed)].groupby(["name_clean", "cat_tz", "group_tz"])["amount"].sum().reset_index().sort_values("amount", ascending=False)
        r = 4
        vf = r
        for k, t in enumerate(var.itertuples(), 1):
            wv.cell(r, 1, k)
            put_text(wv, r, 2, t.name_clean)
            put_text(wv, r, 3, t.cat_tz)
            put_text(wv, r, 4, t.group_tz)
            base = "DBname,$B%d,DBcat,$C%d,DBtype,%s,DBsummed,%s" % (r, r, q(TYPE_EXPENSE), q(YES))
            wv.cell(r, 5, "=SUMIFS(DBamt,%s)" % base).number_format = NUM
            wv.cell(r, 6, "=E%d/%d" % (r, n)).number_format = NUM
            wv.cell(r, 7, '=IFERROR(AVERAGEIF(%s%d:%s%d,"<>0"),0)' % (L(11), r, L(10 + n), r)).number_format = NUM
            wv.cell(r, 8, "=COUNTIFS(%s)" % base)
            wv.cell(r, 9, "=IFERROR(E%d/H%d,0)" % (r, r)).number_format = NUM
            for j, m in enumerate(self.months):
                wv.cell(r, 11 + j, "=SUMIFS(DBamt,%s,DBmonth,%s)" % (base, q(m))).number_format = NUM
            if t.cat_tz in self.cat_anchor:
                self.link(wv.cell(r, 3), SHEET_BY_CAT, "A%d" % self.cat_anchor[t.cat_tz], tall=True)
            if (t.cat_tz, t.name_clean) in self.merch_anchor:
                self.link(wv.cell(r, 2), SHEET_TXNS, "A%d" % self.merch_anchor[(t.cat_tz, t.name_clean)], tall=True)
            r += 1
        vl = r - 1
        if vl >= vf:
            for rr in range(vf, vl + 1):
                wv.cell(rr, 10, "=IFERROR(E%d/SUM($E$%d:$E$%d),0)" % (rr, vf, vl)).number_format = PCT
            colour_scale(wv, "E%d:E%d" % (vf, vl))
            colour_scale(wv, "F%d:G%d" % (vf, vl))
            heat(wv, "%s%d:%s%d" % (L(11), vf, L(10 + n), vl))
            wv.conditional_formatting.add("B%d:B%d" % (vf, vl), FormulaRule(formula=["$F%d>=%s" % (vf, self.hl_exp / 3.0)], font=RED))
            wv.auto_filter.ref = "A3:%s%d" % (L(width), vl)
        wv.cell(r, 2, 'סה"כ הוצאות משתנות').font = BOLD
        wv.cell(r, 5, "=SUM(E%d:E%d)" % (vf, max(vf, vl))).number_format = NUM
        wv.cell(r, 6, "=E%d/%d" % (r, n)).number_format = NUM
        for c in range(1, width + 1):
            wv.cell(r, c).fill = TOT
        setw(wv, {"A": 5, "B": 40, "C": 26, "D": 14, "E": 14, "F": 13, "G": 13, "H": 10, "I": 12, "J": 11})
        wv.freeze_panes = "C4"
        self.rows_written[SHEET_VARIABLE] = r

    # ---- 9. ניתוח נתונים -----------------------------------------------------------------
    def build_analysis(self):
        db, E, n = self.db, self.E, self.n
        wa = rtl(self.wb.create_sheet(SHEET_ANALYSIS))
        wa["A1"] = "ניתוח נתונים — טבלאות ממוינות (נוסחאות על database); הגרפים משמאל מצביעים על הטבלאות; סיכומים לפי קטגוריה/מקור = ממוצע חודשי"
        wa["A1"].font = TITLE
        wa.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(12, 7 + n))
        state = {"row": 3, "minrow": 3}
        anchor1, anchor2 = max(13, 8 + n), max(13, 8 + n) + 11
        le, li = self.lay_exp, self.lay_inc

        def chart(kind, title, hrow, first, last, ncols, second=False, width=18, height=9, pie=False, stacked=False, col=2):
            ch = PieChart() if pie else (LineChart() if kind == "line" else BarChart())
            ch.title, ch.width, ch.height = title, width, height
            ch.add_data(Reference(wa, min_col=col, max_col=col - 1 + ncols, min_row=hrow, max_row=last), titles_from_data=True)
            ch.set_categories(Reference(wa, min_col=1, min_row=first, max_row=last))
            if not pie:
                ch.y_axis.numFmt = "#,##0"
                if kind == "bar":
                    ch.type = "col"
                if stacked:
                    ch.grouping, ch.overlap = "stacked", 100
            else:
                ch.dataLabels = DataLabelList()
                ch.dataLabels.showPercent = True
            state["minrow"] = max(state["minrow"], hrow + int(height * 2.1) + 2)
            wa.add_chart(ch, "%s%d" % (L(anchor2 if second else anchor1), hrow))

        def block(title, headers):
            state["row"] = max(state["row"], state["minrow"])
            wa.cell(state["row"], 1, title).font = BOLD
            state["row"] += 1
            hdr(wa, state["row"], headers)
            h = state["row"]
            state["row"] += 1
            return h

        def nxt():
            state["row"] += 1
            return state["row"] - 1

        def row_stats(row, tot_col, first_m):
            """avg / nz-avg columns for a row whose months start at column first_m."""
            wa.cell(row, tot_col + 1, "=%s%d/%d" % (L(tot_col), row, n)).number_format = NUM
            wa.cell(row, tot_col + 2, '=IFERROR(AVERAGEIF(%s%d:%s%d,"<>0"),0)' % (L(first_m), row, L(first_m + n - 1), row)).number_format = NUM

        def months_sumifs(row, first_m, cond):
            for j, m in enumerate(self.months):
                wa.cell(row, first_m + j, "=SUMIFS(DBamt,%s,DBmonth,%s)" % (cond, q(m))).number_format = NUM

        # א. monthly
        h = block("א. הוצאות, הכנסות ומאזן לפי חודש", ["חודש", "הוצאות שוטפות", "הכנסות", "מאזן", "חיסכון והשקעות"])
        f1 = state["row"]
        for j in range(n):
            row = nxt()
            wa.cell(row, 1, self.mlist[j])
            for c, f in enumerate(["=%s%s%d" % (sref(SHEET_EXP), L(le.m0 + j), self.TR), "=%s%s%d" % (sref(SHEET_INC), L(li.m0 + j), self.TRI),
                                   "=%s%s%d" % (sref(SHEET_INC), L(li.m0 + j), self.BAL), "=%s%s%d" % (sref(SHEET_INC), L(li.m0 + j), self.SAV)]):
                wa.cell(row, 2 + c, f).number_format = NUM
        l1 = state["row"] - 1
        state["row"] += 2
        chart("bar", "הוצאות מול הכנסות לפי חודש", h, f1, l1, 2)
        chart("line", "מאזן חודשי (הכנסות − הוצאות)", h, f1, l1, 3, second=True)
        # ב. by group
        h = block("ב. הוצאות לפי קבוצה (סכימת תזרים) — ממוין", ["קבוצה", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס", "% מסה\"כ"] + self.mlist)
        fB = state["row"]
        for g in self.group_order:
            row = nxt()
            put_text(wa, row, 1, g)
            cond = "DBtype,%s,DBgroup,$A%d,DBsummed,%s" % (q(TYPE_EXPENSE), row, q(YES))
            wa.cell(row, 2, "=SUMIFS(DBamt,%s)" % cond).number_format = NUM
            row_stats(row, 2, 6)
            wa.cell(row, 5, "=IFERROR(B%d/%s$%s$%d,0)" % (row, sref(SHEET_EXP), L(le.c_total), self.TR)).number_format = PCT
            months_sumifs(row, 6, cond)
        lB = state["row"] - 1
        state["row"] += 2
        if lB >= fB:
            colour_scale(wa, "C%d:C%d" % (fB, lB))
            colour_scale(wa, "D%d:D%d" % (fB, lB))
            heat(wa, "%s%d:%s%d" % (L(6), fB, L(5 + n), lB))
            chart("bar", "הוצאות לפי קבוצה – ממוצע חודשי", h, fB, lB, 1, col=3)
            chart("bar", "פילוח הוצאות לפי קבוצה", h, fB, lB, 1, second=True, pie=True, col=3)
            h2 = block("ב2. הוצאות לפי קבוצה – ממוצע חודשי ללא אפס (no zero, ממוין לפי ממוצע ללא אפס)", ["קבוצה", "ממוצע ללא אפס", "ממוצע חודשי"])
            fB2 = state["row"]
            for g in sorted(self.group_order, key=lambda x: -nz_avg(self.mv((db["type"] == TYPE_EXPENSE) & (db["group_tz"] == x)))):
                row = nxt()
                k = self.group_order.index(g)
                put_text(wa, row, 1, g)
                wa.cell(row, 2, "=D%d" % (fB + k)).number_format = NUM
                wa.cell(row, 3, "=C%d" % (fB + k)).number_format = NUM
            lB2 = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "B%d:B%d" % (fB2, lB2))
            chart("bar", "הוצאות לפי קבוצה – ממוצע חודשי no zero", h2, fB2, lB2, 1)
        # ג. by payment method
        src_tot = E.groupby("pay")["amount"].sum().sort_values(ascending=False)
        h = block("ג. הוצאות לפי אמצעי תשלום — ממוין", ["אמצעי תשלום", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס", "מס' עסקאות"] + self.mlist)
        fC = state["row"]
        for pay in src_tot.index:
            row = nxt()
            put_text(wa, row, 1, pay)
            cond = "DBtype,%s,DBpay,$A%d,DBsummed,%s" % (q(TYPE_EXPENSE), row, q(YES))
            wa.cell(row, 2, "=SUMIFS(DBamt,%s)" % cond).number_format = NUM
            row_stats(row, 2, 6)
            wa.cell(row, 5, "=COUNTIFS(%s)" % cond)
            months_sumifs(row, 6, cond)
        lC = state["row"] - 1
        state["row"] += 2
        if lC >= fC:
            colour_scale(wa, "C%d:C%d" % (fC, lC))
            colour_scale(wa, "D%d:D%d" % (fC, lC))
            heat(wa, "%s%d:%s%d" % (L(6), fC, L(5 + n), lC))
            chart("bar", "הוצאות לפי אמצעי תשלום – ממוצע חודשי", h, fC, lC, 1, col=3)
            h = block("ג-nz. הוצאות לפי אמצעי תשלום – ממוצע חודשי ללא אפס (ממוין)", ["אמצעי תשלום", "ממוצע ללא אפס", "ממוצע חודשי"])
            fCn = state["row"]
            for pay in sorted(src_tot.index, key=lambda p: -nz_avg(self.mv((db["type"] == TYPE_EXPENSE) & (db["pay"] == p)))):
                row = nxt()
                k = list(src_tot.index).index(pay)
                put_text(wa, row, 1, pay)
                wa.cell(row, 2, "=D%d" % (fC + k)).number_format = NUM
                wa.cell(row, 3, "=C%d" % (fC + k)).number_format = NUM
            lCn = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "B%d:B%d" % (fCn, lCn))
            chart("bar", "הוצאות לפי אמצעי תשלום – ממוצע חודשי no zero", h, fCn, lCn, 1)
        # ג2. by card (pie) — payment methods of rows that carry a card
        cards = E[E["card"] != ""].groupby("pay")["amount"].sum().sort_values(ascending=False)
        if len(cards):
            h = block("ג2. הוצאות לפי כרטיס אשראי — ממוצע חודשי", ["כרטיס", "ממוצע חודשי", 'סה"כ %d חודשים' % n, "% מהכרטיסים"])
            fC2 = state["row"]
            for pay in cards.index:
                row = nxt()
                put_text(wa, row, 1, pay)
                cond = "DBtype,%s,DBpay,$A%d,DBsummed,%s" % (q(TYPE_EXPENSE), row, q(YES))
                wa.cell(row, 3, "=SUMIFS(DBamt,%s)" % cond).number_format = NUM
                wa.cell(row, 2, "=C%d/%d" % (row, n)).number_format = NUM
            lC2 = state["row"] - 1
            for rr in range(fC2, lC2 + 1):
                wa.cell(rr, 4, "=IFERROR(C%d/SUM($C$%d:$C$%d),0)" % (rr, fC2, lC2)).number_format = PCT
            state["row"] += 2
            chart("bar", "פילוח הוצאות לפי כרטיס אשראי (ממוצע חודשי)", h, fC2, lC2, 1, pie=True)
        # ג3. by person
        persons = E.groupby("person")["amount"].sum().sort_values(ascending=False)
        h = block("ג3. הוצאות לפי אדם (לפי הכרטיס/המקור) — ממוצע חודשי", ["אדם", "ממוצע חודשי", 'סה"כ %d חודשים' % n, "%"] + self.mlist)
        fP = state["row"]
        for p in persons.index:
            row = nxt()
            put_text(wa, row, 1, p)
            cond = "DBtype,%s,DBperson,$A%d,DBsummed,%s" % (q(TYPE_EXPENSE), row, q(YES))
            wa.cell(row, 3, "=SUMIFS(DBamt,%s)" % cond).number_format = NUM
            wa.cell(row, 2, "=C%d/%d" % (row, n)).number_format = NUM
            months_sumifs(row, 5, cond)
        lP = state["row"] - 1
        if lP >= fP:
            for rr in range(fP, lP + 1):
                wa.cell(rr, 4, "=IFERROR(C%d/SUM($C$%d:$C$%d),0)" % (rr, fP, lP)).number_format = PCT
            state["row"] += 2
            chart("bar", "הוצאות לפי אדם – ממוצע חודשי", h, fP, lP, 1)
        # ד. income by source
        inc_sorted = [c for _, c in sorted(self.inc_rows, key=lambda x: -self.inc_avg[x[1]])]
        h = block("ד. הכנסות לפי מקור — ממוין", ["מקור", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס", "% מסה\"כ", "חודשים עם הכנסה"] + self.mlist)
        fD = state["row"]
        for c in inc_sorted:
            row = nxt()
            ir = self.inc_row[c]
            put_text(wa, row, 1, c)
            wa.cell(row, 2, "=%s%s%d" % (sref(SHEET_INC), L(li.c_total), ir)).number_format = NUM
            wa.cell(row, 3, "=B%d/%d" % (row, n)).number_format = NUM
            wa.cell(row, 4, "=%s%s%d" % (sref(SHEET_INC), L(li.c_nz), ir)).number_format = NUM
            wa.cell(row, 5, "=IFERROR(B%d/%s$%s$%d,0)" % (row, sref(SHEET_INC), L(li.c_total), self.TRI)).number_format = PCT
            wa.cell(row, 6, "=%s%s%d" % (sref(SHEET_INC), L(li.c_cnt), ir))
            for j in range(n):
                wa.cell(row, 7 + j, "=%s%s%d" % (sref(SHEET_INC), L(li.m0 + j), ir)).number_format = NUM
        lD = state["row"] - 1
        state["row"] += 2
        if lD >= fD:
            colour_scale(wa, "C%d:C%d" % (fD, lD))
            colour_scale(wa, "D%d:D%d" % (fD, lD))
            heat(wa, "%s%d:%s%d" % (L(7), fD, L(6 + n), lD))
            chart("bar", "הכנסות לפי מקור – ממוצע חודשי", h, fD, lD, 1, width=22, col=3)
            h = block("ד-nz. הכנסות לפי מקור – ממוצע חודשי ללא אפס (ממוין)", ["מקור", "ממוצע ללא אפס", "ממוצע חודשי"])
            fDn = state["row"]
            for c in sorted(inc_sorted, key=lambda x: -nz_avg(self.cat_month_values(TYPE_INCOME, x))):
                row = nxt()
                k = inc_sorted.index(c)
                put_text(wa, row, 1, c)
                wa.cell(row, 2, "=D%d" % (fD + k)).number_format = NUM
                wa.cell(row, 3, "=C%d" % (fD + k)).number_format = NUM
            lDn = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "B%d:B%d" % (fDn, lDn))
            chart("bar", "הכנסות לפי מקור – ממוצע חודשי no zero", h, fDn, lDn, 1, width=22)
            hP = block("ד2. פילוח הכנסות (עוגות): ממוצע חודשי / ממוצע ללא אפס / סה\"כ", ["מקור", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס"])
            fPi = state["row"]
            for k, c in enumerate(inc_sorted):
                row = nxt()
                put_text(wa, row, 1, c)
                wa.cell(row, 2, "=B%d" % (fD + k)).number_format = NUM
                wa.cell(row, 3, "=C%d" % (fD + k)).number_format = NUM
                wa.cell(row, 4, "=D%d" % (fD + k)).number_format = NUM
            lPi = state["row"] - 1
            state["row"] += 2
            chart("bar", "פילוח הכנסות – ממוצע חודשי", hP, fPi, lPi, 1, pie=True, col=3)
            chart("bar", "פילוח הכנסות – ממוצע חודשי ללא אפס", hP, fPi, lPi, 1, second=True, pie=True, col=4)
            state["row"] = max(state["row"], state["minrow"])
            wa.cell(state["row"], 1, 'ד3. פילוח הכנסות – סה"כ %d חודשים (עוגה משמאל)' % n).font = BOLD
            state["row"] += 1
            chart("bar", 'פילוח הכנסות – סה"כ %d חודשים' % n, state["row"] - 1, fPi, lPi, 1, pie=True, col=2)
        # ה. income by group (person) per month
        if self.inc_group_order:
            h = block("ה. הכנסות לפי קבוצה (בן/בת זוג / אחר) לפי חודש", ["חודש"] + list(self.inc_group_order))
            fE = state["row"]
            for j, m in enumerate(self.months):
                row = nxt()
                wa.cell(row, 1, self.mlist[j])
                for c, g in enumerate(self.inc_group_order):
                    wa.cell(row, 2 + c, "=SUMIFS(DBamt,DBtype,%s,DBgroup,%s,DBmonth,%s,DBsummed,%s)" % (q(TYPE_INCOME), q(g), q(m), q(YES))).number_format = NUM
            lE = state["row"] - 1
            state["row"] += 2
            chart("bar", "הכנסות לפי קבוצה", h, fE, lE, len(self.inc_group_order), stacked=True)
        # ו. top merchants
        top = E.groupby("name_clean")["amount"].sum().sort_values(ascending=False).head(self.top_n)
        if len(top):
            h = block('ו. %d בתי העסק / ההוצאות הגדולים ביותר (לפי "שם מובן") — ממוין' % len(top),
                      ["שם מובן", 'סה"כ %d חודשים' % n, "ממוצע חודשי", "ממוצע ללא אפס", "מס' עסקאות", "קטגוריה"] + self.mlist)
            fF = state["row"]
            for name in top.index:
                row = nxt()
                put_text(wa, row, 1, name)
                cond = "DBtype,%s,DBname,$A%d,DBsummed,%s" % (q(TYPE_EXPENSE), row, q(YES))
                wa.cell(row, 2, "=SUMIFS(DBamt,%s)" % cond).number_format = NUM
                wa.cell(row, 3, "=B%d/%d" % (row, n)).number_format = NUM
                wa.cell(row, 4, '=IFERROR(AVERAGEIF(%s%d:%s%d,"<>0"),0)' % (L(7), row, L(6 + n), row)).number_format = NUM
                wa.cell(row, 5, "=COUNTIFS(%s)" % cond)
                put_text(wa, row, 6, E.loc[E["name_clean"] == name, "cat_tz"].iloc[0])
                months_sumifs(row, 7, cond)
            lF = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "C%d:C%d" % (fF, lF))
            colour_scale(wa, "D%d:D%d" % (fF, lF))
            heat(wa, "%s%d:%s%d" % (L(7), fF, L(6 + n), lF))
            chart("bar", "%d ההוצאות הגדולות – ממוצע חודשי" % len(top), h, fF, lF, 1, width=22, height=13, col=3)
            topn = sorted(top.index, key=lambda x: -nz_avg(self.mv((db["type"] == TYPE_EXPENSE) & (db["name_clean"] == x))))
            h = block("ו-nz. %d ההוצאות הגדולות – ממוצע חודשי ללא אפס (ממוין לפי ממוצע ללא אפס)" % len(top), ["שם מובן", "ממוצע ללא אפס", "ממוצע חודשי", "חודשים עם חיוב"])
            fFn = state["row"]
            for name in topn:
                row = nxt()
                k = list(top.index).index(name)
                put_text(wa, row, 1, name)
                wa.cell(row, 2, "=D%d" % (fF + k)).number_format = NUM
                wa.cell(row, 3, "=C%d" % (fF + k)).number_format = NUM
                wa.cell(row, 4, '=COUNTIF(%s%d:%s%d,"<>0")' % (L(7), fF + k, L(6 + n), fF + k))
            lFn = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "B%d:B%d" % (fFn, lFn))
            chart("bar", "%d ההוצאות הגדולות – ממוצע חודשי no zero" % len(top), h, fFn, lFn, 1, width=22, height=13)
        # ז. fixed vs variable
        fixed_rows = [rr for c, rr in self.cat_stat_row.items() if c in self.fixed]
        h = block("ז. הוצאות קבועות מול משתנות לפי חודש (קבועות = fixed_categories בהגדרות)", ["חודש", "קבועות", "משתנות", "% קבועות"])
        fG = state["row"]
        for j in range(n):
            row = nxt()
            wa.cell(row, 1, self.mlist[j])
            fx = "+".join("%s%s%d" % (sref(SHEET_EXP), L(le.m0 + j), rr) for rr in fixed_rows) or "0"
            wa.cell(row, 2, "=" + fx).number_format = NUM
            wa.cell(row, 3, "=%s%s%d-B%d" % (sref(SHEET_EXP), L(le.m0 + j), self.TR, row)).number_format = NUM
            wa.cell(row, 4, "=IFERROR(B%d/(B%d+C%d),0)" % (row, row, row)).number_format = "0%"
        lG = state["row"] - 1
        state["row"] += 2
        chart("bar", "קבועות מול משתנות", h, fG, lG, 2, stacked=True)
        # ח. Israel vs abroad
        h = block('ח. הוצאות בישראל מול חו"ל (לפי מטבע העסקה) לפי חודש', ["חודש", "ישראל (₪)", 'חו"ל (מט"ח, בש"ח)'])
        fH = state["row"]
        for j, m in enumerate(self.months):
            row = nxt()
            wa.cell(row, 1, self.mlist[j])
            wa.cell(row, 2, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBcur,%s,DBsummed,%s)" % (q(TYPE_EXPENSE), q(m), q("ILS"), q(YES))).number_format = NUM
            wa.cell(row, 3, "=SUMIFS(DBamt,DBtype,%s,DBmonth,%s,DBcur,%s,DBsummed,%s)" % (q(TYPE_EXPENSE), q(m), q("<>ILS"), q(YES))).number_format = NUM
        lH = state["row"] - 1
        state["row"] += 2
        chart("bar", 'ישראל מול חו"ל', h, fH, lH, 2, stacked=True)
        # ט. benefits club (only when configured and present in the data)
        for prog in self.cfg.get("benefit_programs") or []:
            label = prog.get("label") or prog["id"]
            if not ((db["source"] == label) | (db["source"] == prog["id"])).any():
                continue
            pat = prog.get("bank_debit_pattern") or ""
            h = block("ט. %s – ערך נקוב, חיוב בפועל וחיסכון לפי חודש חיוב" % label, ["חודש", "רכישות (ערך נקוב)", "חיוב בבנק", "הנחה (חיסכון)"])
            fI = state["row"]
            for j, m in enumerate(self.months):
                row = nxt()
                wa.cell(row, 1, self.mlist[j])
                wa.cell(row, 2, "=SUMIFS(DBamt,DBsource,%s,DBtype,%s,DBmonth,%s,DBamt,\">0\")" % (q(label), q(TYPE_EXPENSE), q(m))).number_format = NUM
                wa.cell(row, 3, ("=SUMIFS(DBamt,DBtype,%s,DBorig,%s,DBmonth,%s)" % (q(TYPE_CARD_DEBIT), q("*%s*" % pat), q(m))) if pat else 0).number_format = NUM
                wa.cell(row, 4, "=-SUMIFS(DBamt,DBcat,%s,DBsource,%s,DBmonth,%s)" % (q(DISCOUNT_CAT), q(label), q(m))).number_format = NUM
            lI = state["row"] - 1
            state["row"] += 2
            chart("bar", "%s: ערך נקוב מול חיוב בפועל" % label, h, fI, lI, 3)
        # י. secondary scheme groups by month (deep)
        if SHEET_PROPOSED in self.sheets and self.newgrp_rows:
            ln = self.lay_new
            h = block("י. הוצאות לפי קבוצה (סכימה מוצעת) לפי חודש — ממוין", ["קבוצה"] + self.mlist + ['סה"כ', "ממוצע חודשי"])
            fJ = state["row"]
            for g, _ in self.new_scheme:
                row = nxt()
                put_text(wa, row, 1, g)
                for j in range(n):
                    wa.cell(row, 2 + j, "=%s%s%d" % (sref(SHEET_PROPOSED), L(ln.m0 + j), self.newgrp_rows[g])).number_format = NUM
                wa.cell(row, 2 + n, "=SUM(B%d:%s%d)" % (row, L(1 + n), row)).number_format = NUM
                wa.cell(row, 3 + n, "=%s%d/%d" % (L(2 + n), row, n)).number_format = NUM
            lJ = state["row"] - 1
            state["row"] += 2
            colour_scale(wa, "%s%d:%s%d" % (L(3 + n), fJ, L(3 + n), lJ))
            chart("bar", "הוצאות לפי קבוצה מוצעת ולפי חודש", h, fJ, lJ, n, width=24, height=12)
        setw(wa, {"A": 42, "B": 14, "C": 14, "D": 24, "E": 13, "F": 11, "G": 11, "H": 11, "I": 13, "J": 11})
        self.rows_written[SHEET_ANALYSIS] = state["row"]

    # ---- 10. לסיווג ידני ------------------------------------------------------------------
    def manual_mask(self):
        db = self.db
        note = db["rule_note"].fillna("")
        cat = db["cat_tz"].fillna("")
        m_note = note.apply(lambda s: any(k in s for k in MANUAL_MARKERS))
        m_cat = (cat == "") | (cat == UNKNOWN_CAT) | cat.apply(lambda s: any(k in s for k in MANUAL_CAT_MARKERS))
        return db["_in"] & db["type"].isin([TYPE_EXPENSE, TYPE_INCOME, TYPE_REIMBURSABLE]) & (m_cat | m_note)

    @staticmethod
    def needed_text(cat, note):
        m = re.search(r"נדרש:\s*(.+)", note or "")
        if m:
            return "נדרש: " + m.group(1).strip()
        if not cat or cat == UNKNOWN_CAT or any(k in cat for k in MANUAL_CAT_MARKERS):
            return "נדרש: בחירת קטגוריה מהרשימה (לא נמצא כלל סיווג)" + (" — " + note if note and note != "לא נמצא כלל סיווג" else "")
        return "נדרש: אימות הסיווג — " + (note or "")

    def build_manual(self):
        wm = rtl(self.wb.create_sheet(SHEET_MANUAL))
        ncols = len(DB_COLUMNS) + 3
        wm["A1"] = ('שורות לא מזוהות / לאימות — מלא "%s" ו/או בחר "%s", שמור את הקובץ, והרץ: python3 apply_user_labels.py' % (USER_TEXT_COL, USER_CAT_COL))
        wm["A1"].font = Font(bold=True, size=12)
        wm.merge_cells(start_row=1, start_column=1, end_row=1, end_column=min(ncols, 14))
        wm["A2"] = 'הטקסט והבחירה שלך נשמרים גם אחרי בנייה מחדש של הקובץ. בחירת קטגוריה מעבירה את השורה לקטגוריה שנבחרה בכל הלשוניות; טקסט חופשי נשמר כ"שם מובן" והערה. עמודת "נדרש:" מסבירה איזו החלטה חסרה.'
        wm["A2"].font = GREY
        wm.merge_cells(start_row=2, start_column=1, end_row=2, end_column=min(ncols, 14))
        hdr(wm, 3, [DB_HEADERS_HE[k] for k in DB_COLUMNS] + [NEEDED_COL, USER_TEXT_COL, USER_CAT_COL])
        todo = self.db[self.manual_mask()].sort_values("amount", ascending=False)
        dv = DataValidation(type="list", formula1="=L_allcats", allow_blank=True)
        wm.add_data_validation(dv)
        r = 4
        nN, nU, nC = len(DB_COLUMNS) + 1, len(DB_COLUMNS) + 2, len(DB_COLUMNS) + 3
        for t in todo.itertuples():
            d = t._asdict()
            for c, k in enumerate(DB_COLUMNS, 1):
                v = d[k]
                if k == "amount":
                    v = float(v)
                elif v == "":
                    v = None
                put_text(wm, r, c, v)
            wm.cell(r, self.col_idx["amount"]).number_format = NUM2
            put_text(wm, r, nN, self.needed_text(d["cat_tz"], d["rule_note"])).font = RED
            ut, uc = self.labels.get(str(t.id), ("", ""))
            put_text(wm, r, nU, ut or None).fill = INPUT
            put_text(wm, r, nC, uc or None).fill = INPUT
            dv.add("%s%d" % (L(nC), r))
            r += 1
        tab = Table(displayName="ManualTable", ref="A3:%s%d" % (L(nC), max(r - 1, 4)))
        tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        wm.add_table(tab)
        setw(wm, {self.CL["original_name"]: 30, self.CL["name_clean"]: 34, self.CL["details"]: 40, self.CL["rule_note"]: 40,
                  self.CL["cat_tz"]: 22, self.CL["pay"]: 22, L(nN): 44, L(nU): 34, L(nC): 30})
        self.manual_rows = r - 4
        self.rows_written[SHEET_MANUAL] = self.manual_rows

    # ---- run ---------------------------------------------------------------------------------
    def build(self, out_path, prev_labels):
        self.labels = prev_labels
        self.build_scheme()
        # always a fresh workbook: the template (if any) only contributed its category list
        self.wb = openpyxl.Workbook()
        del self.wb[self.wb.sheetnames[0]]
        self.build_database()
        self.build_helper()
        self.build_expenses()
        self.build_incomes()
        if SHEET_PROPOSED in self.sheets:
            self.build_proposed()
        if SHEET_LISTS in self.sheets:
            self.build_lists()
        if SHEET_PIVOT in self.sheets:
            self.build_pivot()
        if SHEET_BY_CAT in self.sheets:
            self.build_by_category()
        if SHEET_TXNS in self.sheets:
            self.build_txns()
        if SHEET_VARIABLE in self.sheets:
            self.build_variable()
        self.build_analysis()
        self.manual_rows = 0
        if SHEET_MANUAL in self.sheets:
            self.build_manual()
        if SHEET_PIVOT in self.sheets and SHEET_LISTS in self.sheets:
            self.finish_pivot_links()
        # final order = SHEET_ORDER (הוצאות first)
        order = [s for s in SHEET_ORDER if s in self.wb.sheetnames]
        order += [s for s in self.wb.sheetnames if s not in order]
        self.wb._sheets = [self.wb[s] for s in order]
        for s in HIDDEN_SHEETS:
            if s in self.wb.sheetnames:
                self.wb[s].sheet_state = "hidden"
        self.wb.active = order.index(SHEET_EXP)
        os.makedirs(os.path.dirname(out_path) or ".", exist_ok=True)
        self.wb.save(out_path)
        return {"template_file": self.template_path, "template_categories": len(self.template_rows or [])}


# ----------------------------------------------------------------------------- main
def main(argv=None):
    def extra(p):
        p.add_argument("--database", default=None, help="database.csv (default: config work.database)")
        p.add_argument("--out", default=None, help="output workbook (default: config outputs.workbook)")
    args, cfg = parse_args(__doc__.splitlines()[0], extra, argv)
    db_path = args.database or project_path(cfg, "work.database")
    out_path = args.out or project_path(cfg, "outputs.workbook")
    db = load_db(db_path)
    template = read_template_scheme(cfg)
    labels = read_previous(out_path)
    if labels:
        log("preserved %d user labels from the previous workbook" % len(labels))
    b = Builder(cfg, db, template)
    try:
        info = b.build(out_path, labels)
    except SystemExit:
        raise
    except Exception as e:  # noqa: BLE001
        import traceback
        log(traceback.format_exc())
        fail("workbook build failed: %s: %s" % (type(e).__name__, e),
             "check work/database.csv columns/values and the template settings; run with the synthetic project to isolate")
        return
    E, I = b.E, b.I
    result = {
        "ok": True, "workbook": out_path, "level": cfg.level, "n_months": b.n, "months": b.months,
        "sheets": [s for s in b.wb.sheetnames], "rows_per_sheet": b.rows_written,
        "db_rows": int(len(db)), "expense_categories": len(b.cat_list), "income_categories": len(b.inc_rows),
        "dropped_categories": b.dropped, "manual_rows": b.manual_rows,
        "preserved_labels": len(labels),
        "template": info, "built_at": _dt.datetime.now().isoformat(timespec="seconds"),
        "totals": {"expense": round(float(E["amount"].sum()), 2), "income": round(float(I["amount"].sum()), 2)},
    }
    log("saved %s | sheets %d | db rows %d | expense %.2f | income %.2f" % (
        out_path, len(result["sheets"]), len(db), result["totals"]["expense"], result["totals"]["income"]))
    emit(result)


if __name__ == "__main__":
    main()
